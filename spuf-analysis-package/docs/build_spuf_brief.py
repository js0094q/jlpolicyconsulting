from pathlib import Path
import csv
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas

root = Path('/Users/josephstewart/Documents/JLPolicyConsulting/output/exports/spuf-analysis-package')
data = root / 'data'
out_dir = root / 'docs'
out_dir.mkdir(parents=True, exist_ok=True)
docx_path = out_dir / 'spuf-policy-brief.docx'
pdf_path = out_dir / 'spuf-policy-brief.pdf'

def read_csv(name):
    with (data / name).open(newline='', encoding='utf-8') as f:
        return list(csv.DictReader(f))

rc = read_csv('repricing_churn_decomposition.csv')
ss_sum = read_csv('sponsor_segmentation_summary.csv')
bt_delta = read_csv('beneficiary_tier_linkage_delta_q4_to_q1.csv')

# DOCX

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)

doc.add_heading('SPUF Part D Analysis Brief', level=0)
doc.add_paragraph('Date: April 22, 2026')
doc.add_paragraph('Audience: Market access, payer strategy, reimbursement, and policy stakeholders.')

p = doc.add_paragraph()
p.add_run('Executive summary').bold = True

doc.add_paragraph(
    'The strongest signal across the current SPUF windows is that quarter-over-quarter movement is not uniform. '
    'The 2025Q3 to 2025Q4 pair behaves like a stable panel, while 2025Q4 to 2026Q1 shows high composition churn. '
    'This changes how pricing movement should be interpreted and supports a decomposition-first narrative. '
    'Generic tier movement, sponsor clustering, and PBM channel fingerprints provide the most publishable website story set.'
)

doc.add_heading('1) Repricing vs churn decomposition', level=1)
t = doc.add_table(rows=1, cols=5)
h = t.rows[0].cells
h[0].text='Pair'; h[1].text='Matched share'; h[2].text='Add/drop share'; h[3].text='Repriced share of matched'; h[4].text='Interpretation'
for row in rc:
    r=t.add_row().cells
    r[0].text=row['pair']
    r[1].text=f"{float(row['matched_share'])*100:.2f}%"
    r[2].text=f"{float(row['add_drop_share'])*100:.2f}%"
    r[3].text=f"{float(row['repriced_share_of_matched'])*100:.2f}%"
    interp = 'Higher continuity baseline' if row['pair']=='2025Q3_to_2025Q4' else 'High churn regime, decomposition required'
    r[4].text=interp

doc.add_heading('2) Generic tier worsening with same-NDC controls', level=1)
doc.add_paragraph(
    'Local generic QoQ outputs support a controlled narrative where tier worsening can occur even when same-NDC price changes are not uniformly inflationary. '
    'This is best framed as access friction under mixed repricing and mix-shift effects, with explicit caveats.'
)

doc.add_heading('3) Sponsor segmentation', level=1)
seg_para = ', '.join([f"{r['segment']}: {r['sponsor_count']}" for r in ss_sum])
doc.add_paragraph(f"Observed segment counts in 2025Q4 to 2026Q1: {seg_para}.")
doc.add_paragraph('Interpretation: sponsor behavior is heterogeneous and supports cluster-based story framing, not one market-average narrative.')

doc.add_heading('4) PBM mail/retail channel profile', level=1)
doc.add_paragraph(
    'PBM-level channel outputs provide a cross-sectional profile suitable for an Insight article. '
    'Claims should stay descriptive and avoid inferred intent without longitudinal confirmation.'
)

doc.add_heading('5) Beneficiary cost-sharing and tier linkage', level=1)
doc.add_paragraph(
    'A join-heavy linkage between plan-formulary tier structure and beneficiary design fields was executed for 2025Q4 and 2026Q1. '
    'Tier-level deductible-applies shares moved materially in upper tiers in this derived aggregation. '
    'Preferred cost amount fields were mostly zero in this extraction path, so cost trend claims require additional field-level validation before publication.'
)

# add a compact table of top tier deltas
bt = doc.add_table(rows=1, cols=4)
bh = bt.rows[0].cells
bh[0].text='Tier'; bh[1].text='NDC row delta'; bh[2].text='Deductible share delta'; bh[3].text='Plan count delta'
for row in bt_delta:
    r = bt.add_row().cells
    r[0].text = row['tier']
    r[1].text = row['linked_formulary_ndc_rows_delta']
    r[2].text = row['mean_plan_ded_applies_share_delta']
    r[3].text = row['plan_count_delta']

doc.add_heading('Best website stories to publish', level=1)
for s in [
    'Repricing vs churn in Part D (Research)',
    'Generic access friction despite stable same-NDC pricing (Research)',
    'Sponsor behavioral clustering (Insight)',
    'Generic coverage contraction signals (Insight)',
    'PBM network/channel fingerprints (Insight)',
]:
    doc.add_paragraph(s, style='List Bullet')

doc.add_heading('Sources used', level=1)
for src in [
    '/Volumes/SPUF/cms_partd_pricing_pipeline/data/analysis_full_parquet_metrics.json',
    '/Volumes/SPUF/cms_partd_pricing_pipeline/data/analysis/generic_qoq/2025Q4_to_2026Q1/*',
    '/Volumes/SPUF/cms_partd_pricing_pipeline/data/published/analysis/pricing_signals/2025Q4_to_2026Q1/*',
    '/Volumes/SPUF/cms_partd_pricing_pipeline/data/analysis/pbm_channel_steering/2026Q1/*',
    '/Volumes/SPUF/cms_partd_pricing_pipeline/data/extracted/2025Q4/*',
    '/Volumes/SPUF/cms_partd_pricing_pipeline/data/extracted/2026Q1/*',
]:
    doc.add_paragraph(src, style='List Bullet')

doc.save(docx_path)

# PDF
c = canvas.Canvas(str(pdf_path), pagesize=LETTER)
width, height = LETTER
x = 54
y = height - 54

def line(text, size=10, gap=14, bold=False):
    global y
    c.setFont('Helvetica-Bold' if bold else 'Helvetica', size)
    c.drawString(x, y, text)
    y -= gap

line('SPUF Part D Analysis Brief', size=16, gap=22, bold=True)
line('Date: April 22, 2026', size=10)
line('')
line('Executive summary', bold=True)
line('Decomposition is required for credible interpretation: 2025Q4 to 2026Q1 shows high churn.', size=10)
line('Top story set: repricing vs churn, generic access friction, sponsor clusters, PBM channel profile.', size=10)
line('')
line('Repricing vs churn metrics', bold=True)
for row in rc:
    line(f"{row['pair']}: matched={float(row['matched_share'])*100:.2f}%, add/drop={float(row['add_drop_share'])*100:.2f}%, repriced/matched={float(row['repriced_share_of_matched'])*100:.2f}%", size=9, gap=12)
line('')
line('Sponsor segment counts', bold=True)
for r in ss_sum:
    line(f"{r['segment']}: {r['sponsor_count']}", size=9, gap=12)
line('')
line('Beneficiary-tier linkage note', bold=True)
line('Upper-tier deductible-share shifts are visible in derived joins, with cost fields requiring validation.', size=9, gap=12)
line('')
line('Best website stories', bold=True)
for s in [
    '1) Repricing vs churn in Part D',
    '2) Generic access friction despite stable same-NDC pricing',
    '3) Sponsor behavioral clustering',
    '4) Generic coverage contraction signals',
    '5) PBM network/channel fingerprints',
]:
    line(s, size=9, gap=12)

c.showPage()
c.save()

print(docx_path)
print(pdf_path)
