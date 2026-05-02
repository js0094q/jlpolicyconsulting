from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

md_path = Path('/Users/josephstewart/Documents/JLPolicyConsulting/output/exports/spuf-site-drafts/research/part-d-generic-tier-defensiveness-research.mdx')
out_docx = Path('/Users/josephstewart/Documents/JLPolicyConsulting/output/exports/spuf-site-drafts/research/part-d-generic-tier-defensiveness-research.docx')
out_pdf = Path('/Users/josephstewart/Documents/JLPolicyConsulting/output/exports/spuf-site-drafts/research/part-d-generic-tier-defensiveness-research.pdf')


def strip_frontmatter(lines):
    if not lines or lines[0].strip() != '---':
        return lines
    end_idx = 1
    while end_idx < len(lines) and lines[end_idx].strip() != '---':
        end_idx += 1
    if end_idx >= len(lines):
        return lines
    return lines[end_idx + 1 :]


def is_table_line(line: str) -> bool:
    s = line.strip()
    return s.startswith('|') and s.endswith('|')


def parse_blocks(lines):
    content_lines = strip_frontmatter(lines)
    blocks = []
    i = 0
    while i < len(content_lines):
        line = content_lines[i].rstrip('\n')
        if not line:
            i += 1
            continue

        if line.startswith('# '):
            blocks.append(('h1', [line[2:].strip()]))
            i += 1
            continue
        if line.startswith('## '):
            blocks.append(('h2', [line[3:].strip()]))
            i += 1
            continue
        if line.startswith('- '):
            items = []
            while i < len(content_lines) and content_lines[i].startswith('- '):
                items.append(content_lines[i][2:].strip())
                i += 1
            blocks.append(('ul', items))
            continue

        if is_table_line(line):
            table_rows = []
            while i < len(content_lines) and is_table_line(content_lines[i]):
                row = [c.strip() for c in content_lines[i].strip().strip('|').split('|')]
                table_rows.append(row)
                i += 1
            if len(table_rows) >= 2:
                blocks.append(('table', table_rows))
            continue

        if line.startswith('### '):
            blocks.append(('h3', [line[4:].strip()]))
            i += 1
            continue

        para = []
        while i < len(content_lines):
            cur = content_lines[i].strip()
            if (
                not cur
                or cur.startswith('# ')
                or cur.startswith('## ')
                or cur.startswith('### ')
                or cur.startswith('- ')
                or is_table_line(cur)
            ):
                break
            para.append(cur)
            i += 1
        if para:
            blocks.append(('p', [' '.join(para)]))
        else:
            i += 1

    return blocks


def to_docx(blocks):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)

    for kind, payload in blocks:
        if kind == 'h1':
            doc.add_heading(payload[0], level=1)
        elif kind == 'h2':
            doc.add_heading(payload[0], level=2)
        elif kind == 'h3':
            doc.add_heading(payload[0], level=3)
        elif kind == 'ul':
            for item in payload:
                doc.add_paragraph(item, style='List Bullet')
        elif kind == 'table':
            header = payload[0]
            rows = payload[1:] if len(payload) > 1 else []
            if rows:
                table = doc.add_table(rows=1 + len(rows), cols=len(header))
                for j, h in enumerate(header):
                    table.cell(0, j).text = h
                for r_idx, row in enumerate(rows, start=1):
                    for c_idx, value in enumerate(row):
                        if c_idx < len(header):
                            table.cell(r_idx, c_idx).text = value
        elif kind == 'p':
            for line in payload:
                doc.add_paragraph(line)

    doc.save(out_docx)


def to_pdf(blocks):
    styles = getSampleStyleSheet()
    body = []
    body.append(Paragraph('Part D Generic Tier Defensiveness: Decomposition, Generics, and Sponsor Behavior', styles['Heading1']))
    body.append(Spacer(1, 0.12 * 72 / 72))

    for kind, payload in blocks:
        if kind == 'h1':
            body.append(Spacer(1, 0.08 * 72 / 72))
            body.append(Paragraph(payload[0], styles['Heading2']))
        elif kind == 'h2':
            body.append(Paragraph(payload[0], styles['Heading3']))
        elif kind == 'h3':
            body.append(Paragraph(payload[0], styles['Heading4']))
        elif kind == 'ul':
            for item in payload:
                body.append(Paragraph(f"• {item}", styles['BodyText']))
        elif kind == 'table':
            data = payload
            tbl = Table(data, repeatRows=1)
            tbl.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f2f4f7')),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('LEFTPADDING', (0, 0), (-1, -1), 4),
                ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ]))
            body.append(tbl)
            body.append(Spacer(1, 0.12 * 72 / 72))
        elif kind == 'p':
            for text in payload:
                body.append(Paragraph(text, styles['BodyText']))

    body.append(Spacer(1, 0.12 * 72 / 72))
    body.append(Paragraph('Prepared for: JL Policy Consulting research pipeline', styles['Normal']))

    doc_pdf = SimpleDocTemplate(str(out_pdf), pagesize=LETTER)
    doc_pdf.build(body)


if __name__ == '__main__':
    lines = md_path.read_text(encoding='utf-8').splitlines()
    blocks = parse_blocks(lines)
    to_docx(blocks)
    to_pdf(blocks)
    print(out_docx)
    print(out_pdf)
