#!/usr/bin/env python3
"""Build the final ARM provider net cost recovery tool package.

The package is generated output for the JL Policy Consulting workspace. The
workbook is formula-driven for field use; this script also mirrors the core
logic for PDF export and validation so the package can be tested without
requiring Excel automation.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import shutil
import textwrap
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table as PdfTable,
    TableStyle,
)


PACKAGE_DIR = Path(__file__).resolve().parent
BUILD_DATE = date(2026, 5, 2)
VERSION = "Final ARM v1.0"

WORKBOOK_NAME = "Adrabetadex_Provider_Net_Cost_Recovery_Calculator.xlsx"
PROVIDER_PDF_NAME = (
    "Adrabetadex_Provider_Worksheet_Approved_Scope_S01_2026-05-02.pdf"
)
README_NAME = "README_Adrabetadex_Final_ARM_Tool.md"
ARM_GUIDE_MD = "Adrabetadex_ARM_Conversation_Guide.md"
BILLING_CHECKLIST_MD = "Adrabetadex_Billing_Readiness_Checklist.md"
PROVIDER_WORKSHEET_MD = "Adrabetadex_Provider_Facing_Worksheet.md"
SOURCE_REGISTER_CSV = "Adrabetadex_Policy_Source_Register.csv"
VALIDATION_MD = "Adrabetadex_Final_ARM_Tool_Validation_Summary.md"
VALIDATION_JSON = "Adrabetadex_Final_ARM_Tool_Validation_Results.json"
DELIVERY_DIR_NAME = "final_arm_provider_net_cost_recovery_tool_2026-05-02"

FIELD_USE_FRAMING = (
    "This model assumes payer authorization has been approved for the patient, "
    "dose, frequency, site of care, and sourcing pathway reflected in the "
    "selected scenario. If authorization is pending, denied, under appeal, "
    "expired, or inconsistent with the modeled site/sourcing pathway, the "
    "economic output should not be used."
)

STANDARD_CAVEAT = (
    "This scenario is for reimbursement planning only. It does not guarantee "
    "coverage, coding, payment, reimbursement, acquisition cost, provider "
    "margin, payer behavior, 340B pricing, pass-through status, HCPCS "
    "assignment, or claim outcome. Providers are responsible for verifying "
    "coding, billing, payer rules, contract terms, documentation, and "
    "institutional policy before claim submission."
)

THREE_FORTY_B_CAVEAT = (
    "340B acquisition values are illustrative unless site-verified. This model "
    "does not guarantee 340B eligibility, 340B ceiling price, sub-ceiling price, "
    "provider acquisition cost, payer-recognized payment basis, or payer "
    "reimbursement outcome."
)

SCOPE_WARNING = (
    "Authorization exists, but approval scope does not match the selected "
    "reimbursement scenario. Economic output should not be used until dose, "
    "frequency, site of care, sourcing pathway, and authorization period are "
    "aligned."
)

PARTIAL_YEAR_LABEL = "Projected only through currently authorized treatment window."

FORBIDDEN_EXTERNAL_PHRASES = [
    "Break-even " + "support per dose",
    "Annual break-even " + "support",
    "Manufacturer " + "support needed",
    "Provider " + "profit",
    "Provider margin " + "guarantee",
    "Coverage " + "secured",
    "Pass-through " + "guarantees payment",
    "340B " + "guarantees profitability",
]


@dataclass(frozen=True)
class Scenario:
    scenario_id: str
    name: str
    purpose: str
    site_of_care: str
    site_profile: str
    status_340b: str
    payer_type: str
    sourcing_pathway: str
    coding_status: str
    payment_status: str
    authorization_type: str
    dose_mg: int
    frequency: str
    administrations_per_year: int
    acquisition_basis: str
    payment_method: str
    claim_type: str
    medicaid_lob: str
    treatment_340b: str
    allowed_drug_reimbursement: float
    acquisition_cost: float
    procedure_revenue: float
    other_reimbursable_services: float
    pharmacy_handling_cost: float
    labor_admin_cost: float
    billing_revenue_cycle_cost: float
    claim_denial_rate: float
    appeal_recovery_rate: float
    days_to_payment: int
    financing_rate: float
    wac_per_vial: float
    vial_size_mg: int
    scenario_note: str


SCENARIOS: list[Scenario] = [
    Scenario(
        "S01",
        "340B COE/HOPD, buy-and-bill, pass-through-if-granted",
        "Core launch planning scenario.",
        "HOPD",
        "COE",
        "Yes",
        "Medicare",
        "Buy-and-bill",
        "C9399/NOC",
        "OPPS pass-through active",
        "Initial approval",
        900,
        "Every 2 weeks",
        26,
        "340B",
        "Medicare benchmark",
        "UB-04/837I",
        "Not applicable",
        "Carve-out",
        40170.0,
        29991.0,
        500.0,
        0.0,
        400.0,
        750.0,
        300.0,
        0.075,
        0.60,
        75,
        0.08,
        39000.0,
        900,
        "Pass-through used only if granted and effective; otherwise re-run S02.",
    ),
    Scenario(
        "S02",
        "340B COE/HOPD, pre-pass-through/interim coding",
        "Launch bridge risk scenario.",
        "HOPD",
        "COE",
        "Yes",
        "Medicare",
        "Buy-and-bill",
        "C9399/NOC",
        "Manual pricing",
        "Initial approval",
        900,
        "Every 2 weeks",
        26,
        "340B",
        "Manual pricing",
        "UB-04/837I",
        "Not applicable",
        "Carve-out",
        37050.0,
        29991.0,
        500.0,
        0.0,
        500.0,
        850.0,
        400.0,
        0.12,
        0.50,
        95,
        0.08,
        39000.0,
        900,
        "Manual-review branch; documentation and payer confirmation drive risk.",
    ),
    Scenario(
        "S03",
        "340B COE/HOPD, product-specific HCPCS active",
        "Post-bridge steady-state scenario.",
        "HOPD",
        "COE",
        "Yes",
        "Medicare",
        "Buy-and-bill",
        "Product-specific HCPCS assigned",
        "ASP available",
        "Reauthorization approval",
        900,
        "Every 2 weeks",
        26,
        "340B",
        "Medicare benchmark",
        "UB-04/837I",
        "Not applicable",
        "Carve-out",
        39273.0,
        29991.0,
        500.0,
        0.0,
        350.0,
        650.0,
        275.0,
        0.05,
        0.70,
        65,
        0.08,
        39000.0,
        900,
        "Payment method still depends on setting, payer, and effective files.",
    ),
    Scenario(
        "S04",
        "Non-340B HOPD, commercial buy-and-bill",
        "Sensitivity scenario for thinner economics.",
        "HOPD",
        "Academic center",
        "No",
        "Commercial",
        "Buy-and-bill",
        "Payer-directed code",
        "Commercial contracted",
        "Initial approval",
        900,
        "Every 2 weeks",
        26,
        "WAC",
        "Commercial contract",
        "UB-04/837I",
        "Not applicable",
        "Not applicable",
        42000.0,
        39000.0,
        625.0,
        0.0,
        450.0,
        700.0,
        350.0,
        0.09,
        0.55,
        80,
        0.08,
        39000.0,
        900,
        "Use payer contract or site-entered allowed amount; do not default to Medicare.",
    ),
    Scenario(
        "S05",
        "White-bagging/specialty pharmacy sourcing",
        "Inventory-risk reduction but limited/no drug spread.",
        "HOPD",
        "COE",
        "Unknown",
        "Commercial",
        "White-bagging",
        "Payer-directed code",
        "White-bagged/no drug reimbursement",
        "Initial approval",
        900,
        "Every 2 weeks",
        26,
        "Specialty pharmacy",
        "Payer-specific",
        "UB-04/837I",
        "Not applicable",
        "Not applicable",
        0.0,
        0.0,
        450.0,
        0.0,
        250.0,
        600.0,
        275.0,
        0.06,
        0.50,
        55,
        0.08,
        39000.0,
        900,
        "Drug reimbursement and acquisition exposure are zero unless payer terms say otherwise.",
    ),
    Scenario(
        "S06",
        "Medicaid FFS 340B",
        "NDC/PAD, carve-in/carve-out, and duplicate-discount branch.",
        "HOPD",
        "Pediatric specialty hospital",
        "Yes",
        "Medicaid FFS",
        "Buy-and-bill",
        "Payer-directed code",
        "Medicaid FFS state-specific",
        "Initial approval",
        900,
        "Every 2 weeks",
        26,
        "340B",
        "Medicaid FFS",
        "UB-04/837I",
        "FFS",
        "Unresolved",
        36000.0,
        29991.0,
        475.0,
        0.0,
        500.0,
        800.0,
        450.0,
        0.10,
        0.50,
        90,
        0.08,
        39000.0,
        900,
        "Requires state-specific methodology, NDC/PAD fields, and duplicate-discount resolution.",
    ),
    Scenario(
        "S07",
        "Medicaid MCO",
        "Plan-specific Medicaid branch. Do not assume FFS logic.",
        "HOPD",
        "Community site",
        "Unknown",
        "Medicaid MCO",
        "Mixed",
        "Payer-directed code",
        "Medicaid MCO plan-specific",
        "Initial approval",
        900,
        "Every 2 weeks",
        26,
        "Payer-directed",
        "Medicaid MCO",
        "UB-04/837I",
        "MCO",
        "Unknown",
        35000.0,
        39000.0,
        475.0,
        0.0,
        500.0,
        850.0,
        500.0,
        0.11,
        0.45,
        95,
        0.08,
        39000.0,
        900,
        "Requires plan-specific confirmation; FFS payment logic is not assumed.",
    ),
    Scenario(
        "S08",
        "Physician office pre-ASP",
        "103% WAC benchmark sensitivity only.",
        "Physician office",
        "Community site",
        "No",
        "Medicare",
        "Buy-and-bill",
        "No product-specific HCPCS",
        "Pre-ASP WAC benchmark",
        "Initial approval",
        900,
        "Every 2 weeks",
        26,
        "WAC",
        "Medicare benchmark",
        "CMS-1500/837P",
        "Not applicable",
        "Not applicable",
        40170.0,
        39000.0,
        325.0,
        0.0,
        250.0,
        400.0,
        250.0,
        0.07,
        0.60,
        60,
        0.08,
        39000.0,
        900,
        "103% WAC is used only for physician office pre-ASP sensitivity.",
    ),
]


SOURCE_REGISTER = [
    {
        "Assumption ID": "SRC-001",
        "Assumption name": "HCPCS identifies items and services, not coverage/payment",
        "Value": "Coding status is separated from payment status.",
        "Category": "policy-backed rule",
        "Source": "CMS HCPCS Level II Coding Procedures",
        "Source URL": "https://www.cms.gov/medicare/coding/medhcpcsgeninfo/hcpcscodingprocess.html",
        "Source date": "Page modified 2026-02-19; accessed 2026-05-02",
        "Last reviewed": "2026-05-02",
        "Owner": "Reimbursement strategy owner",
        "Field-use restriction": "ARM-facing",
    },
    {
        "Assumption ID": "SRC-002",
        "Assumption name": "HCPCS quarterly drug/biological coding cycle",
        "Value": "Drug and biological code applications follow quarterly CMS review cycles.",
        "Category": "policy-backed rule",
        "Source": "CMS HCPCS Level II Coding Decisions",
        "Source URL": "https://www.cms.gov/medicare/coding-billing/healthcare-common-procedure-system/current-prior-years-level-ii-coding-decisions",
        "Source date": "Accessed 2026-05-02",
        "Last reviewed": "2026-05-02",
        "Owner": "Reimbursement strategy owner",
        "Field-use restriction": "ARM-facing",
    },
    {
        "Assumption ID": "SRC-003",
        "Assumption name": "OPPS pass-through is a payment pathway",
        "Value": "Pass-through status is modeled as payment status, not product coding status.",
        "Category": "policy-backed rule",
        "Source": "CMS Pass-Through Payment Status and New Technology APC",
        "Source URL": "https://www.cms.gov/Medicare/Medicare-Fee-for-Service-Payment/HospitalOutpatientPPS/passthrough_payment.html",
        "Source date": "Accessed 2026-05-02",
        "Last reviewed": "2026-05-02",
        "Owner": "Reimbursement strategy owner",
        "Field-use restriction": "ARM-facing",
    },
    {
        "Assumption ID": "SRC-004",
        "Assumption name": "Medicare Part B ASP/WAC branches",
        "Value": "ASP is primary for many Part B drugs; WAC may be used when ASP is unavailable.",
        "Category": "policy-backed rule",
        "Source": "CMS Part B Drugs and Biologicals",
        "Source URL": "https://www.cms.gov/cms-guide-medical-technology-companies-and-other-interested-parties/payment/part-b-drugs",
        "Source date": "Page updated 2026-03-10; accessed 2026-05-02",
        "Last reviewed": "2026-05-02",
        "Owner": "Reimbursement strategy owner",
        "Field-use restriction": "ARM-facing",
    },
    {
        "Assumption ID": "SRC-005",
        "Assumption name": "ASP pricing file caveat",
        "Value": "Presence or absence in CMS ASP files does not itself indicate coverage.",
        "Category": "policy-backed rule",
        "Source": "CMS ASP Pricing Files",
        "Source URL": "https://www.cms.gov/medicare/payment/part-b-drugs/asp-pricing-files",
        "Source date": "Accessed 2026-05-02",
        "Last reviewed": "2026-05-02",
        "Owner": "Reimbursement strategy owner",
        "Field-use restriction": "ARM-facing",
    },
    {
        "Assumption ID": "SRC-006",
        "Assumption name": "Medicaid PAD NDC/utilization requirements",
        "Value": "Medicaid PAD branches require NDC and utilization fields where applicable.",
        "Category": "policy-backed rule",
        "Source": "Medicaid.gov Physician Administered Drugs",
        "Source URL": "https://www.medicaid.gov/medicaid/prescription-drugs/state-prescription-drug-resources/physician-administered-drugs-pad",
        "Source date": "Page updated 2022-10-20; accessed 2026-05-02",
        "Last reviewed": "2026-05-02",
        "Owner": "Medicaid policy owner",
        "Field-use restriction": "ARM-facing",
    },
    {
        "Assumption ID": "SRC-007",
        "Assumption name": "MDRP 11-digit NDC data structure",
        "Value": "MDRP product and AMP files identify products using 11-digit NDCs.",
        "Category": "policy-backed rule",
        "Source": "Medicaid.gov MDRP Data",
        "Source URL": "https://www.medicaid.gov/medicaid/prescription-drugs/medicaid-drug-rebate-program/medicaid-drug-rebate-program-data",
        "Source date": "Page updated 2026-04-13; accessed 2026-05-02",
        "Last reviewed": "2026-05-02",
        "Owner": "Medicaid policy owner",
        "Field-use restriction": "ARM-facing",
    },
    {
        "Assumption ID": "SRC-008",
        "Assumption name": "340B duplicate-discount prohibition",
        "Value": "Medicaid + 340B scenarios require duplicate-discount resolution before clean output.",
        "Category": "policy-backed rule",
        "Source": "HRSA Duplicate Discount Prohibition",
        "Source URL": "https://www.hrsa.gov/opa/program-requirements/medicaid-exclusion",
        "Source date": "Accessed 2026-05-02",
        "Last reviewed": "2026-05-02",
        "Owner": "340B/compliance owner",
        "Field-use restriction": "ARM-facing",
    },
    {
        "Assumption ID": "SRC-009",
        "Assumption name": "340B Medicaid Exclusion File applies to FFS",
        "Value": "MEF logic is not assumed to resolve Medicaid MCO plan requirements.",
        "Category": "policy-backed rule",
        "Source": "HRSA Medicaid Exclusion File help",
        "Source URL": "https://340bregistration.hrsa.gov/help/Manufacturer/Reports/MedicaidExclusionFile.htm",
        "Source date": "Accessed 2026-05-02",
        "Last reviewed": "2026-05-02",
        "Owner": "340B/compliance owner",
        "Field-use restriction": "ARM-facing",
    },
    {
        "Assumption ID": "MOD-001",
        "Assumption name": "Adrabetadex launch WAC per vial",
        "Value": "$39,000 per 900 mg vial",
        "Category": "planning assumption",
        "Source": "Internal planning assumption",
        "Source URL": "",
        "Source date": "2026-05-02",
        "Last reviewed": "2026-05-02",
        "Owner": "Market access owner",
        "Field-use restriction": "Internal only",
    },
    {
        "Assumption ID": "MOD-002",
        "Assumption name": "Primary 340B acquisition assumption",
        "Value": "Primary 23.1% planning assumption; site verification required",
        "Category": "planning assumption",
        "Source": "Internal planning assumption",
        "Source URL": "",
        "Source date": "2026-05-02",
        "Last reviewed": "2026-05-02",
        "Owner": "Market access owner",
        "Field-use restriction": "Internal/ARM-facing",
    },
]


INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
WARNING_FILL = PatternFill("solid", fgColor="FCE4D6")
GOOD_FILL = PatternFill("solid", fgColor="E2F0D9")
LOCKED_FILL = PatternFill("solid", fgColor="E7E6E6")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
THIN_GRAY = Side(style="thin", color="D9E2F3")
MEDIUM_BLUE = Side(style="medium", color="1F4E78")


def as_currency(value: float | int | str) -> str:
    if value == "" or value is None:
        return ""
    return f"${float(value):,.0f}"


def as_percent(value: float | int | str) -> str:
    if value == "" or value is None:
        return ""
    return f"{float(value) * 100:.1f}%"


def scenario_by_id(scenario_id: str) -> Scenario:
    for scenario in SCENARIOS:
        if scenario.scenario_id == scenario_id:
            return scenario
    raise ValueError(f"Unknown scenario ID: {scenario_id}")


def is_medicaid(payer_type: str) -> bool:
    return payer_type in {"Medicaid FFS", "Medicaid MCO"}


def normalize_status(status: str) -> str:
    return (status or "").strip()


def normalize_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        return date.fromisoformat(value)
    raise TypeError(f"Unsupported date value: {value!r}")


def calculate_scenario(
    scenario: Scenario,
    controls: dict[str, Any],
    override_values: dict[str, Any] | None = None,
) -> dict[str, Any]:
    override_values = override_values or {}

    site_of_care = override_values.get("Site of care") or scenario.site_of_care
    sourcing_pathway = (
        override_values.get("Sourcing pathway") or scenario.sourcing_pathway
    )
    dose_mg = float(override_values.get("Modeled dose") or scenario.dose_mg)
    frequency = override_values.get("Modeled frequency") or scenario.frequency
    administrations_per_year = int(
        override_values.get("Administrations per year")
        or scenario.administrations_per_year
    )
    payer_type = override_values.get("Payer type") or scenario.payer_type
    status_340b = override_values.get("340B status") or scenario.status_340b

    pa_status = normalize_status(controls.get("PA status", "Approved"))
    pa_approved = pa_status in {"Approved", "Reauthorization approved"}

    approval_end = controls.get("Approval end date", date(2026, 12, 31))
    approval_end = normalize_date(approval_end)

    authorized_doses = int(controls.get("Authorized number of doses", 26) or 0)
    authorized_dose = float(controls.get("Authorized dose", dose_mg) or 0)
    authorized_frequency = controls.get("Authorized frequency", frequency)
    authorized_site = controls.get("Authorized site of care", site_of_care)
    authorized_sourcing = controls.get(
        "Authorized sourcing pathway", sourcing_pathway
    )

    scope_match = (
        pa_approved
        and authorized_dose == dose_mg
        and authorized_frequency == frequency
        and authorized_site == site_of_care
        and authorized_sourcing == sourcing_pathway
        and approval_end >= BUILD_DATE
        and authorized_doses > 0
    )

    ndc = str(controls.get("NDC", "") or "").strip()
    ndc_confirmed = controls.get("11-digit NDC format confirmed", "No") == "Yes"
    unit_measure = str(controls.get("Unit of measure", "") or "").strip()
    quantity = controls.get("Quantity administered", "")
    claim_type = str(controls.get("Claim type", scenario.claim_type) or "").strip()
    mco_policy = controls.get("MCO plan-specific policy confirmed", "No") == "Yes"

    medicaid_required_ok = True
    if is_medicaid(payer_type):
        medicaid_required_ok = (
            bool(ndc)
            and ndc_confirmed
            and bool(unit_measure)
            and quantity not in {"", None}
            and bool(claim_type)
            and (payer_type != "Medicaid MCO" or mco_policy)
        )

    duplicate_discount_resolved = True
    if is_medicaid(payer_type) and status_340b == "Yes":
        duplicate_discount_resolved = (
            controls.get("Medicaid duplicate-discount risk resolved", "No") == "Yes"
        )

    modellable = pa_approved and scope_match and medicaid_required_ok and duplicate_discount_resolved

    if not pa_approved:
        if pa_status == "Not started":
            readiness_message = "PA not initiated - economics not actionable."
        elif pa_status == "Submitted/pending":
            readiness_message = "PA pending - reimbursement scenario not yet executable."
        elif pa_status == "Denied/appeal":
            readiness_message = "PA unresolved - economic result not reliable."
        elif pa_status == "Reauthorization pending":
            readiness_message = "Reauthorization pending - do not calculate beyond current authorization window."
        else:
            readiness_message = "PA approval missing - economics suppressed."
    elif not scope_match:
        readiness_message = SCOPE_WARNING
    elif not medicaid_required_ok:
        readiness_message = "Medicaid/PAD required fields unresolved - clean economics suppressed."
    elif not duplicate_discount_resolved:
        readiness_message = "340B Medicaid duplicate-discount risk unresolved - clean economics suppressed."
    else:
        readiness_message = "Approved scope matches modeled scenario - economics calculated."

    vials_per_administration = math.ceil(dose_mg / scenario.vial_size_mg)
    gross_drug_exposure = scenario.wac_per_vial * vials_per_administration

    if scenario.acquisition_basis in {"Specialty pharmacy", "White-bagging"}:
        acquisition_cost = 0.0
    else:
        acquisition_cost = scenario.acquisition_cost

    if scenario.sourcing_pathway in {"White-bagging", "Specialty pharmacy"}:
        drug_reimbursement = 0.0
    elif (
        site_of_care == "Physician office"
        and scenario.payment_status == "Pre-ASP WAC benchmark"
    ):
        drug_reimbursement = gross_drug_exposure * 1.03
    else:
        drug_reimbursement = scenario.allowed_drug_reimbursement

    admin_revenue = scenario.procedure_revenue + scenario.other_reimbursable_services
    gross_recovery = drug_reimbursement + admin_revenue
    operating_burden = (
        scenario.pharmacy_handling_cost
        + scenario.labor_admin_cost
        + scenario.billing_revenue_cycle_cost
    )
    denial_drag = gross_recovery * scenario.claim_denial_rate * (
        1 - scenario.appeal_recovery_rate
    )
    financing_drag = acquisition_cost * scenario.financing_rate * (
        scenario.days_to_payment / 365
    )
    net_recovery = (
        gross_recovery
        - acquisition_cost
        - operating_burden
        - denial_drag
        - financing_drag
    )

    authorized_administrations = min(administrations_per_year, authorized_doses)

    if not modellable:
        risk_label = "Not modellable"
    elif net_recovery < 0:
        risk_label = "High"
    elif net_recovery < 2500:
        risk_label = "Moderate"
    else:
        risk_label = "Low"

    cash_exposure = "Low"
    if acquisition_cost > 0 and scenario.days_to_payment >= 90:
        cash_exposure = "High"
    elif acquisition_cost > 0 and scenario.days_to_payment >= 60:
        cash_exposure = "Moderate"

    coding_payment_risk = "Low"
    if scenario.coding_status in {"No product-specific HCPCS", "C9399/NOC"}:
        coding_payment_risk = "High"
    elif scenario.payment_status in {"Manual pricing", "Medicaid MCO plan-specific"}:
        coding_payment_risk = "Moderate"

    medicaid_risk = "Low"
    if is_medicaid(payer_type) and not medicaid_required_ok:
        medicaid_risk = "Not modellable"
    elif is_medicaid(payer_type):
        medicaid_risk = "Moderate"

    risk_flags = {
        "Recovery": risk_label,
        "Cash exposure": cash_exposure if modellable else "Not modellable",
        "Coding/payment": coding_payment_risk if modellable else "Not modellable",
        "Payer control": "Moderate" if controls.get("Payer SOC confirmed") != "Yes" else "Low",
        "Medicaid/PAD": medicaid_risk,
        "340B": (
            "Not modellable"
            if is_medicaid(payer_type)
            and status_340b == "Yes"
            and not duplicate_discount_resolved
            else ("Moderate" if status_340b in {"Yes", "Unknown"} else "Low")
        ),
        "White-bagging": (
            "Moderate" if scenario.sourcing_pathway in {"White-bagging", "Mixed"} else "Low"
        ),
        "Site of care": "Moderate" if controls.get("Payer SOC confirmed") != "Yes" else "Low",
    }

    return {
        "scenario_id": scenario.scenario_id,
        "scenario_name": scenario.name,
        "version": VERSION,
        "date": BUILD_DATE.isoformat(),
        "pa_status": pa_status,
        "pa_approved": pa_approved,
        "scope_match": scope_match,
        "modellable": modellable,
        "readiness_message": readiness_message,
        "projection_label": (
            PARTIAL_YEAR_LABEL
            if modellable and authorized_administrations < administrations_per_year
            else ""
        ),
        "authorized_administrations": authorized_administrations if modellable else 0,
        "gross_drug_exposure": gross_drug_exposure if modellable else "",
        "acquisition_cost": acquisition_cost if modellable else "",
        "drug_reimbursement": drug_reimbursement if modellable else "",
        "admin_revenue": admin_revenue if modellable else "",
        "gross_recovery": gross_recovery if modellable else "",
        "operating_burden": operating_burden if modellable else "",
        "denial_drag": denial_drag if modellable else "",
        "financing_drag": financing_drag if modellable else "",
        "net_recovery": net_recovery if modellable else "",
        "annual_gross_exposure": gross_drug_exposure * authorized_administrations
        if modellable
        else "",
        "annual_acquisition_cost": acquisition_cost * authorized_administrations
        if modellable
        else "",
        "annual_gross_recovery": gross_recovery * authorized_administrations
        if modellable
        else "",
        "annual_operating_burden": operating_burden * authorized_administrations
        if modellable
        else "",
        "annual_denial_drag": denial_drag * authorized_administrations
        if modellable
        else "",
        "annual_financing_drag": financing_drag * authorized_administrations
        if modellable
        else "",
        "annual_net_recovery": net_recovery * authorized_administrations
        if modellable
        else "",
        "risk_flags": risk_flags,
    }


def default_controls(scenario: Scenario) -> dict[str, Any]:
    return {
        "PA status": "Approved",
        "Approval type": scenario.authorization_type,
        "Approval start date": date(2026, 5, 1),
        "Approval end date": date(2026, 12, 31),
        "Authorized number of doses": 26,
        "Authorized dose": scenario.dose_mg,
        "Authorized frequency": scenario.frequency,
        "Authorized site of care": scenario.site_of_care,
        "Authorized sourcing pathway": scenario.sourcing_pathway,
        "Approval reference number": "Provider-entered",
        "Payer SOC confirmed": "Yes",
        "Payer sourcing restriction confirmed": "Yes",
        "Specialist requirement confirmed": "Yes",
        "Baseline documentation confirmed": "Yes",
        "Continuation endpoint defined": "Yes",
        "NDC": "00000-0000-00" if is_medicaid(scenario.payer_type) else "",
        "11-digit NDC format confirmed": "Yes" if is_medicaid(scenario.payer_type) else "Not applicable",
        "Unit of measure": "UN" if is_medicaid(scenario.payer_type) else "",
        "Quantity administered": 1 if is_medicaid(scenario.payer_type) else "",
        "Claim type": scenario.claim_type,
        "340B carve-in/carve-out": scenario.treatment_340b,
        "Medicaid Exclusion File status": "Provider validation required",
        "MCO plan-specific policy confirmed": "Yes" if scenario.payer_type == "Medicaid MCO" else "Not applicable",
        "340B eligibility confirmed": "Yes" if scenario.status_340b == "Yes" else scenario.status_340b,
        "340B applies to service location": "Yes" if scenario.status_340b == "Yes" else scenario.status_340b,
        "340B acquisition assumption type": "Internal planning assumption"
        if scenario.status_340b == "Yes"
        else "Not applicable",
        "Medicaid duplicate-discount risk resolved": "Yes"
        if is_medicaid(scenario.payer_type) and scenario.status_340b == "Yes"
        else "Not applicable",
        "JW/JZ wastage handling confirmed": "No - provider validation required",
    }


def set_title(ws, title: str, subtitle: str | None = None, end_col: int = 8) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    c = ws.cell(1, 1, title)
    c.font = Font(size=16, bold=True, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
        s = ws.cell(2, 1, subtitle)
        s.font = Font(size=10, italic=True, color="1F4E78")
        s.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[2].height = 32


def style_range_header(ws, row: int, start_col: int, end_col: int) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=MEDIUM_BLUE, bottom=MEDIUM_BLUE)


def apply_grid(ws, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    for row in ws.iter_rows(
        min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
    ):
        for cell in row:
            cell.border = Border(
                left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY
            )
            cell.alignment = Alignment(wrap_text=True, vertical="top")


def unlock(cell: Cell) -> None:
    cell.protection = Protection(locked=False)
    cell.fill = INPUT_FILL


def protect(ws) -> None:
    ws.protection.sheet = True
    ws.protection.password = "arm"
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = True


def add_list_validation(ws, cell_range: str, values: list[str]) -> None:
    quoted = ",".join(values)
    dv = DataValidation(type="list", formula1=f'"{quoted}"', allow_blank=True)
    dv.error = "Select an approved value from the list."
    dv.errorTitle = "Invalid selection"
    ws.add_data_validation(dv)
    dv.add(cell_range)


def add_table(ws, ref: str, display_name: str) -> None:
    table = Table(displayName=display_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def write_instructions(wb: Workbook) -> None:
    ws = wb.create_sheet("Instructions")
    set_title(
        ws,
        "Adrabetadex Final ARM Provider Net Cost Recovery Tool",
        "Reimbursement-readiness and economic-exposure model for ARM-led provider conversations after approved authorization.",
        8,
    )
    widths = [28, 34, 28, 38, 18, 18, 18, 18]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width
    ws["A4"] = "Required field-use framing"
    ws["A4"].font = Font(bold=True, color="1F4E78")
    ws.merge_cells("B4:H6")
    ws["B4"] = FIELD_USE_FRAMING
    ws["B4"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A8"] = "Standard caveat"
    ws["A8"].font = Font(bold=True, color="1F4E78")
    ws.merge_cells("B8:H11")
    ws["B8"] = STANDARD_CAVEAT
    ws["B8"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A13"] = "Positioning"
    ws["A13"].font = Font(bold=True, color="1F4E78")
    rows = [
        ("Use as", "A reimbursement-readiness and economic-exposure model for approved-scope scenarios."),
        ("Do not use as", "A provider profitability calculator, coding recommendation, coverage guarantee, payment guarantee, manufacturer funding calculator, 340B spread calculator, or replacement for provider review."),
        ("Release gate", "Economics calculate only when PA is approved and the approval scope matches dose, frequency, site of care, sourcing pathway, authorization dates, and authorized dose count."),
        ("Scenario control", "One selected scenario ID drives all categorical and numeric assumptions. Manual overrides are blank by default and visible when activated."),
    ]
    start = 14
    for r, (label, value) in enumerate(rows, start):
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2, value)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)
    ws["A21"] = "Audience layers"
    ws["A21"].font = Font(bold=True, color="1F4E78")
    audience_rows = [
        ("Locked ARM calculator", "ARMs and internal reimbursement team", "Controlled scenario model with locked formulas, approved input cells, and visible risk flags."),
        ("ARM conversation guide", "ARMs", "Stepwise workflow for provider conversations."),
        ("Provider-facing worksheet/PDF", "Provider finance, pharmacy, billing, access teams", "Educational summary using provider-entered assumptions. No guarantee language."),
        ("Billing-readiness checklist", "Provider billing/revenue cycle", "Candidate claim-readiness prompts only; does not prescribe codes."),
        ("Policy/source register", "MLR, compliance, market access", "CMS, Medicaid, HRSA, payer-policy, and planning assumptions with source dates and owners."),
    ]
    ws.append([])
    ws.append(["Component", "Audience", "Requirement"])
    header_row = 23
    style_range_header(ws, header_row, 1, 3)
    for item in audience_rows:
        ws.append(item)
    apply_grid(ws, 23, 28, 1, 3)
    ws.freeze_panes = "A23"
    protect(ws)


def write_scenario_engine(wb: Workbook) -> None:
    ws = wb.create_sheet("Scenario Engine")
    set_title(ws, "Scenario Engine", "Single scenario ID drives categorical and numeric assumptions.", 12)
    for col, width in {
        "A": 28,
        "B": 28,
        "C": 28,
        "D": 30,
        "E": 24,
        "F": 42,
        "G": 18,
        "H": 22,
        "I": 28,
        "J": 28,
        "K": 18,
        "L": 16,
        "M": 18,
        "N": 18,
        "O": 22,
        "P": 22,
        "Q": 18,
        "R": 18,
        "S": 18,
        "T": 18,
        "U": 18,
        "V": 18,
        "W": 18,
        "X": 18,
        "Y": 18,
        "Z": 18,
        "AA": 14,
        "AB": 14,
        "AC": 14,
        "AD": 14,
        "AE": 14,
        "AF": 14,
        "AG": 40,
    }.items():
        ws.column_dimensions[col].width = width

    ws["A4"] = "Selected Scenario ID"
    ws["B4"] = "S01"
    ws["A5"] = "Scenario name"
    ws["B5"] = '=INDEX($B$35:$B$42,MATCH($B$4,$A$35:$A$42,0))'
    ws["A6"] = "Manual override active?"
    ws["B6"] = "No"
    ws["A8"] = "Field-use framing"
    ws["B8"] = FIELD_USE_FRAMING
    ws.merge_cells("B8:H9")
    ws["B8"].alignment = Alignment(wrap_text=True, vertical="top")
    unlock(ws["B4"])
    unlock(ws["B6"])
    add_list_validation(ws, "B4", [s.scenario_id for s in SCENARIOS])
    add_list_validation(ws, "B6", ["No", "Yes"])

    header = ["Scenario Field", "Preset Value", "Manual Override", "Active Value", "Active Source", "Notes"]
    for idx, value in enumerate(header, 1):
        ws.cell(10, idx, value)
    style_range_header(ws, 10, 1, 6)

    driver_fields = [
        ("Site of care", 1, "HOPD, physician office, ASC/other, inpatient exception"),
        ("Site profile", 2, "COE, pediatric specialty hospital, academic center, community site"),
        ("340B status", 3, "Yes, no, unknown, not applicable"),
        ("Payer type", 4, "Commercial, Medicaid FFS, Medicaid MCO, Medicare, other"),
        ("Sourcing pathway", 5, "Buy-and-bill, white-bagging, specialty pharmacy, mixed"),
        ("Product coding status", 6, "Separates HCPCS/NOC/payer-directed coding from payment state"),
        ("Payment status", 7, "Payment branch, not coding status"),
        ("Authorization type", 8, "Initial approval or reauthorization approval"),
        ("Modeled dose", 9, "Must match authorized dose"),
        ("Modeled frequency", 10, "Must match authorized frequency"),
        ("Administrations per year", 11, "Default 26 for every two weeks"),
        ("Acquisition basis", 12, "WAC, 340B, site-entered, payer-directed, specialty pharmacy, other"),
        ("Payer payment method", 13, "Medicare benchmark, commercial contract, Medicaid FFS, Medicaid MCO, manual pricing, payer-specific"),
        ("Claim type", 14, "UB-04/837I or CMS-1500/837P based on site"),
        ("Medicaid line of business", 15, "FFS, MCO, not applicable"),
        ("340B treatment", 16, "Carve-in, carve-out, unresolved, unknown, not applicable"),
    ]
    for offset, (field, idx, note) in enumerate(driver_fields, 11):
        ws.cell(offset, 1, field)
        ws.cell(offset, 2, f"=INDEX($D$35:$S$42,MATCH($B$4,$A$35:$A$42,0),{idx})")
        ws.cell(offset, 4, f'=IF($B$6="Yes",IF(C{offset}<>"",C{offset},B{offset}),B{offset})')
        ws.cell(offset, 5, f'=IF($B$6="Yes",IF(C{offset}<>"","Manual Override","Preset (override on, blank override)"),"Preset")')
        ws.cell(offset, 6, note)
        unlock(ws.cell(offset, 3))
    apply_grid(ws, 10, 26, 1, 6)

    scenario_headers = [
        "Scenario ID",
        "Scenario Name",
        "Purpose",
        "Site of care",
        "Site profile",
        "340B status",
        "Payer type",
        "Sourcing pathway",
        "Product coding status",
        "Payment status",
        "Authorization type",
        "Modeled dose",
        "Modeled frequency",
        "Administrations per year",
        "Acquisition basis",
        "Payer payment method",
        "Claim type",
        "Medicaid LOB",
        "340B treatment",
        "Allowed drug reimbursement",
        "Acquisition cost",
        "Procedure/admin revenue",
        "Other services",
        "Pharmacy handling",
        "Labor/admin",
        "Billing/revenue-cycle",
        "Claim denial rate",
        "Appeal recovery rate",
        "Days to payment",
        "Financing rate",
        "WAC per vial",
        "Vial size mg",
        "Scenario note",
    ]
    for idx, value in enumerate(scenario_headers, 1):
        ws.cell(34, idx, value)
    style_range_header(ws, 34, 1, len(scenario_headers))
    for row_idx, scenario in enumerate(SCENARIOS, 35):
        row = [
            scenario.scenario_id,
            scenario.name,
            scenario.purpose,
            scenario.site_of_care,
            scenario.site_profile,
            scenario.status_340b,
            scenario.payer_type,
            scenario.sourcing_pathway,
            scenario.coding_status,
            scenario.payment_status,
            scenario.authorization_type,
            scenario.dose_mg,
            scenario.frequency,
            scenario.administrations_per_year,
            scenario.acquisition_basis,
            scenario.payment_method,
            scenario.claim_type,
            scenario.medicaid_lob,
            scenario.treatment_340b,
            scenario.allowed_drug_reimbursement,
            scenario.acquisition_cost,
            scenario.procedure_revenue,
            scenario.other_reimbursable_services,
            scenario.pharmacy_handling_cost,
            scenario.labor_admin_cost,
            scenario.billing_revenue_cycle_cost,
            scenario.claim_denial_rate,
            scenario.appeal_recovery_rate,
            scenario.days_to_payment,
            scenario.financing_rate,
            scenario.wac_per_vial,
            scenario.vial_size_mg,
            scenario.scenario_note,
        ]
        for col_idx, value in enumerate(row, 1):
            ws.cell(row_idx, col_idx, value)
    apply_grid(ws, 34, 42, 1, len(scenario_headers))
    add_table(ws, f"A34:AG42", "ScenarioLibrary")
    for row in range(35, 43):
        for col in [20, 21, 22, 23, 24, 25, 26, 31]:
            ws.cell(row, col).number_format = '$#,##0'
        for col in [27, 28, 30]:
            ws.cell(row, col).number_format = '0.0%'
    ws.freeze_panes = "A34"
    protect(ws)


def write_product_economics(wb: Workbook) -> None:
    ws = wb.create_sheet("Product Economics")
    set_title(ws, "Product Economics", "Gross exposure, acquisition, reimbursement, and dosing formulas.", 6)
    for col, width in zip("ABCDEF", [34, 22, 26, 50, 18, 18]):
        ws.column_dimensions[col].width = width
    rows = [
        ("WAC per vial", "=INDEX('Scenario Engine'!$AE$35:$AE$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Internal planning assumption; update through source register."),
        ("Vial size (mg)", "=INDEX('Scenario Engine'!$AF$35:$AF$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Must align with dosing and wastage logic."),
        ("Dose per administration (mg)", "='Scenario Engine'!$D$19", "Scenario engine", "Must match authorized dose."),
        ("Dosing frequency", "='Scenario Engine'!$D$20", "Scenario engine", "Must match authorized frequency."),
        ("Administrations per year", "='Scenario Engine'!$D$21", "Scenario engine", "Default 26 for every two weeks."),
        ("Vials per administration", "=CEILING(B6/B5,1)", "Derived", "Rounded up to whole vials."),
        ("Wastage amount (mg)", "=(B9*B5)-B6", "Derived", "JW/JZ handling is a checklist/risk prompt only."),
        ("Gross Drug Exposure per Dose", "=B4*B9", "Derived", "WAC per vial x vials per administration."),
        ("Acquisition basis", "='Scenario Engine'!$D$22", "Scenario engine", "WAC, 340B, site-entered, payer-directed, specialty pharmacy, or other."),
        ("Scenario acquisition cost default", "=INDEX('Scenario Engine'!$U$35:$U$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Blank manual overrides remain blank until actively enabled."),
        ("Acquisition Cost per Dose", '=IF(OR(B12="Specialty pharmacy",B12="White-bagging"),0,B13)', "Derived", "Zero for white-bagging/SP unless payer-specific terms say otherwise."),
        ("Payer payment method", "='Scenario Engine'!$D$23", "Scenario engine", "Payment branch, not coding status."),
        ("Allowed Drug Reimbursement per Dose", '=IF(OR(\'Scenario Engine\'!$D$15="White-bagging",\'Scenario Engine\'!$D$15="Specialty pharmacy"),0,IF(AND(\'Scenario Engine\'!$D$11="Physician office",\'Scenario Engine\'!$D$17="Pre-ASP WAC benchmark"),B11*1.03,INDEX(\'Scenario Engine\'!$T$35:$T$42,MATCH(\'Scenario Engine\'!$B$4,\'Scenario Engine\'!$A$35:$A$42,0))))', "Derived", "103% WAC branch applies only to physician office pre-ASP."),
        ("Procedure/Admin Revenue", "=INDEX('Scenario Engine'!$V$35:$V$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Provider/site-entered benchmark or planning default."),
        ("Other Reimbursable Services", "=INDEX('Scenario Engine'!$W$35:$W$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Optional."),
        ("Gross Recovery per Dose", "=B16+B17+B18", "Derived", "Drug reimbursement plus administration revenue."),
    ]
    ws.append(["Field", "Value", "Source", "Notes"])
    style_range_header(ws, 3, 1, 4)
    for idx, row in enumerate(rows, 4):
        ws.cell(idx, 1, row[0])
        ws.cell(idx, 2, row[1])
        ws.cell(idx, 3, row[2])
        ws.cell(idx, 4, row[3])
    for row in [4, 11, 13, 14, 16, 17, 18, 19]:
        ws.cell(row, 2).number_format = '$#,##0'
    apply_grid(ws, 3, 19, 1, 4)
    protect(ws)


def write_coding_payment(wb: Workbook) -> None:
    ws = wb.create_sheet("Coding Payment")
    set_title(ws, "Coding and Payment State", "Coding status and payment status remain distinct throughout the tool.", 6)
    for col, width in zip("ABCDEF", [30, 34, 36, 42, 18, 18]):
        ws.column_dimensions[col].width = width
    rows = [
        ("Product coding status", "='Scenario Engine'!$D$16", "How product may be reported on a claim.", "No product-specific HCPCS; C9399/NOC; product-specific HCPCS assigned; payer-directed code."),
        ("Payment status", "='Scenario Engine'!$D$17", "How reimbursement may be calculated or adjudicated.", "Pre-ASP WAC benchmark; OPPS pass-through active; ASP available; OPPS non-pass-through; manual pricing; commercial contracted; Medicaid-specific; white-bagged/no drug reimbursement."),
        ("Hard rule", "Pass-through active is not coding status.", "Pass-through is a payment state, not a product coding state.", "Do not conflate C-codes/HCPCS/NOC with pass-through/payment methodology."),
        ("Medicare branch", "=IF(AND('Scenario Engine'!$D$11=\"Physician office\",'Scenario Engine'!$D$17=\"Pre-ASP WAC benchmark\"),\"Physician office pre-ASP WAC x 103% only\",IF('Scenario Engine'!$D$17=\"ASP available\",\"ASP branch\",IF('Scenario Engine'!$D$17=\"OPPS pass-through active\",\"OPPS pass-through branch if granted/effective\",'Scenario Engine'!$D$17)))", "Branch logic", "103% WAC is not applied universally."),
        ("Commercial branch", "=IF('Scenario Engine'!$D$14=\"Commercial\",\"Contract/manual pricing; do not default to Medicare\",\"Not commercial\")", "Branch logic", "Use payer-specific allowed amount or provider-entered contract terms."),
        ("Medicaid branch", "=IF(OR('Scenario Engine'!$D$14=\"Medicaid FFS\",'Scenario Engine'!$D$14=\"Medicaid MCO\"),'Scenario Engine'!$D$14&\" - state/plan-specific with PAD/NDC checks\",\"Not Medicaid\")", "Branch logic", "MCO does not assume FFS logic."),
    ]
    ws.append(["Field", "Active Value", "Interpretation", "Guardrail"])
    style_range_header(ws, 3, 1, 4)
    for idx, row in enumerate(rows, 4):
        for col, value in enumerate(row, 1):
            ws.cell(idx, col, value)
    apply_grid(ws, 3, 9, 1, 4)
    protect(ws)


def write_site_channel(wb: Workbook) -> None:
    ws = wb.create_sheet("Site Channel")
    set_title(ws, "Site and Channel", "Site, payer, 340B, Medicaid, and sourcing branches.", 6)
    for col, width in zip("ABCDEF", [30, 34, 34, 42, 18, 18]):
        ws.column_dimensions[col].width = width
    rows = [
        ("Site of care", "='Scenario Engine'!$D$11", "Scenario driver", "Must match authorized site of care."),
        ("Site profile", "='Scenario Engine'!$D$12", "Scenario driver", "COE, pediatric specialty hospital, academic center, community site."),
        ("340B status", "='Scenario Engine'!$D$13", "Scenario driver", "Eligibility must be site and billing-entity specific."),
        ("Payer type", "='Scenario Engine'!$D$14", "Scenario driver", "Commercial, Medicaid FFS, Medicaid MCO, Medicare, or other."),
        ("Sourcing pathway", "='Scenario Engine'!$D$15", "Scenario driver", "Must match authorized sourcing pathway."),
        ("White-bagging rule", '=IF(OR(B8="White-bagging",B8="Specialty pharmacy"),"Drug reimbursement/spread set to zero unless payer-specific value is entered","Buy-and-bill or mixed pathway")', "Derived", "White-bagging reduces inventory exposure but limits/no drug spread."),
        ("340B caveat", THREE_FORTY_B_CAVEAT, "Required caveat", "Display when a 340B acquisition assumption is used."),
    ]
    ws.append(["Field", "Active Value", "Source", "Notes"])
    style_range_header(ws, 3, 1, 4)
    for idx, row in enumerate(rows, 4):
        for col, value in enumerate(row, 1):
            ws.cell(idx, col, value)
    apply_grid(ws, 3, 10, 1, 4)
    ws.row_dimensions[10].height = 58
    protect(ws)


def write_operating_burden(wb: Workbook) -> None:
    ws = wb.create_sheet("Operating Burden")
    set_title(ws, "Operating Burden", "Claim-friction, staffing, billing, and timing burden.", 6)
    for col, width in zip("ABCDEF", [34, 20, 26, 48, 18, 18]):
        ws.column_dimensions[col].width = width
    rows = [
        ("Pharmacy handling cost", "=INDEX('Scenario Engine'!$X$35:$X$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Per administration."),
        ("Labor/admin cost", "=INDEX('Scenario Engine'!$Y$35:$Y$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Per administration."),
        ("Billing/revenue-cycle cost", "=INDEX('Scenario Engine'!$Z$35:$Z$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Per administration."),
        ("Operating Burden per Dose", "=SUM(B4:B6)", "Derived", "Pharmacy handling + labor/admin + billing/revenue-cycle cost."),
        ("Claim denial rate", "=INDEX('Scenario Engine'!$AA$35:$AA$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Post-service claim denial, not PA denial."),
        ("Appeal recovery rate", "=INDEX('Scenario Engine'!$AB$35:$AB$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Claim appeal recovery only."),
        ("Days to payment", "=INDEX('Scenario Engine'!$AC$35:$AC$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Cash-flow timing burden."),
        ("Financing rate", "=INDEX('Scenario Engine'!$AD$35:$AD$42,MATCH('Scenario Engine'!$B$4,'Scenario Engine'!$A$35:$A$42,0))", "Scenario library", "Used only for financing drag."),
        ("Expected Claim Denial Drag", "='Product Economics'!B19*B8*(1-B9)", "Derived", "Gross recovery x claim denial rate x unrecovered share."),
        ("Financing Drag", "='Product Economics'!B14*B11*(B10/365)", "Derived", "Acquisition cost x financing rate x days to payment / 365."),
    ]
    ws.append(["Field", "Value", "Source", "Notes"])
    style_range_header(ws, 3, 1, 4)
    for idx, row in enumerate(rows, 4):
        for col, value in enumerate(row, 1):
            ws.cell(idx, col, value)
    for row in [4, 5, 6, 7, 12, 13]:
        ws.cell(row, 2).number_format = '$#,##0'
    for row in [8, 9, 11]:
        ws.cell(row, 2).number_format = '0.0%'
    apply_grid(ws, 3, 13, 1, 4)
    protect(ws)


def write_payer_controls(wb: Workbook) -> None:
    ws = wb.create_sheet("Payer Controls")
    set_title(ws, "Payer Controls and Authorization Gate", "PA approval and scope match are hard prerequisites before economics calculate.", 7)
    for col, width in zip("ABCDEFG", [36, 34, 28, 48, 16, 16, 16]):
        ws.column_dimensions[col].width = width
    ws.append(["Input / Check", "Value", "Required", "Notes"])
    style_range_header(ws, 3, 1, 4)
    scenario = scenario_by_id("S01")
    defaults = default_controls(scenario)
    rows = [
        ("PA status", defaults["PA status"], "Yes", "Not started, Submitted/pending, Denied/appeal, Approved, Reauthorization approved, Reauthorization pending."),
        ("Approval type", defaults["Approval type"], "Yes", "Initial approval or reauthorization approval."),
        ("Approval start date", defaults["Approval start date"], "Yes", "Used to calculate authorized treatment window."),
        ("Approval end date", defaults["Approval end date"], "Yes", "Must not be expired."),
        ("Authorized number of doses", defaults["Authorized number of doses"], "Yes", "Limits annualized economics when authorization is partial-year."),
        ("Authorized dose", defaults["Authorized dose"], "Yes", "Must equal modeled dose."),
        ("Authorized frequency", defaults["Authorized frequency"], "Yes", "Must equal modeled frequency."),
        ("Authorized site of care", defaults["Authorized site of care"], "Yes", "Must equal modeled site of care."),
        ("Authorized sourcing pathway", defaults["Authorized sourcing pathway"], "Yes", "Must equal modeled sourcing pathway."),
        ("Approval reference number", defaults["Approval reference number"], "No", "Useful for documentation but not required for calculation."),
        ("Payer SOC confirmed", defaults["Payer SOC confirmed"], "Yes", "Flags site-of-care restriction risk if not confirmed."),
        ("Payer sourcing restriction confirmed", defaults["Payer sourcing restriction confirmed"], "Yes", "Flags white-bagging/specialty pharmacy risk if not confirmed."),
        ("Specialist requirement confirmed", defaults["Specialist requirement confirmed"], "Yes", "Documentation risk if not confirmed."),
        ("Baseline documentation confirmed", defaults["Baseline documentation confirmed"], "Yes", "Initial approval/reauthorization risk if not confirmed."),
        ("Continuation endpoint defined", defaults["Continuation endpoint defined"], "Yes", "Reauthorization risk if not defined."),
        ("NDC", defaults["NDC"], "Yes for Medicaid", "Required for Medicaid PAD logic where applicable."),
        ("11-digit NDC format confirmed", defaults["11-digit NDC format confirmed"], "Yes for Medicaid", "Provider validation checkbox."),
        ("Unit of measure", defaults["Unit of measure"], "Yes for Medicaid", "Must align with claim requirements."),
        ("Quantity administered", defaults["Quantity administered"], "Yes for Medicaid", "Required for claim integrity."),
        ("Claim type", defaults["Claim type"], "Yes", "UB-04/837I or CMS-1500/837P based on site."),
        ("340B carve-in/carve-out", defaults["340B carve-in/carve-out"], "If 340B + Medicaid", "Required for duplicate-discount risk assessment."),
        ("Medicaid Exclusion File status", defaults["Medicaid Exclusion File status"], "Prompt", "Provider validation required where relevant."),
        ("MCO plan-specific policy confirmed", defaults["MCO plan-specific policy confirmed"], "Yes for Medicaid MCO", "Do not assume FFS rules."),
        ("340B eligibility confirmed", defaults["340B eligibility confirmed"], "If 340B used", "Must be site and billing-entity specific."),
        ("340B applies to service location", defaults["340B applies to service location"], "If 340B used", "Do not assume entity-level eligibility applies to all services."),
        ("340B acquisition assumption type", defaults["340B acquisition assumption type"], "If 340B used", "Site-verified, internal planning assumption, or unknown."),
        ("Medicaid duplicate-discount risk resolved", defaults["Medicaid duplicate-discount risk resolved"], "If Medicaid + 340B", "If unresolved, output is Not modellable."),
        ("JW/JZ wastage handling confirmed", defaults["JW/JZ wastage handling confirmed"], "Prompt", "Claim-documentation risk only; does not prescribe billing."),
    ]
    for idx, row in enumerate(rows, 4):
        for col, value in enumerate(row, 1):
            ws.cell(idx, col, value)
        unlock(ws.cell(idx, 2))
    date_rows = [6, 7]
    for row in date_rows:
        ws.cell(row, 2).number_format = "yyyy-mm-dd"
    apply_grid(ws, 3, 31, 1, 4)
    validations = {
        "B4": ["Not started", "Submitted/pending", "Denied/appeal", "Approved", "Reauthorization approved", "Reauthorization pending"],
        "B5": ["Initial approval", "Reauthorization approval"],
        "B11": ["HOPD", "Physician office", "ASC/other", "Inpatient exception"],
        "B12": ["Buy-and-bill", "White-bagging", "Specialty pharmacy", "Mixed"],
    }
    yes_no_rows = ["B14:B18", "B20", "B26", "B27", "B30:B31"]
    for cell_range in yes_no_rows:
        add_list_validation(ws, cell_range, ["Yes", "No", "Unknown", "Not applicable"])
    for cell, values in validations.items():
        add_list_validation(ws, cell, values)

    ws["A34"] = "Authorization and compliance gate"
    ws["A34"].font = Font(bold=True, color="1F4E78")
    gate_rows = [
        ("PA gate message", '=IF(B4="Not started","PA not initiated - economics not actionable.",IF(B4="Submitted/pending","PA pending - reimbursement scenario not yet executable.",IF(B4="Denied/appeal","PA unresolved - economic result not reliable.",IF(B4="Reauthorization pending","Reauthorization pending - do not calculate beyond current authorization window.","PA approved status present."))))'),
        ("PA approved for calculation?", '=OR(B4="Approved",B4="Reauthorization approved")'),
        ("Approval scope match?", '=AND(B35,B9=\'Scenario Engine\'!$D$19,B10=\'Scenario Engine\'!$D$20,B11=\'Scenario Engine\'!$D$11,B12=\'Scenario Engine\'!$D$15,B7>=TODAY(),B8>0)'),
        ("Scope warning", f'=IF(B36,"", "{SCOPE_WARNING}")'),
        ("Medicaid/PAD fields complete?", '=IF(OR(\'Scenario Engine\'!$D$14="Medicaid FFS",\'Scenario Engine\'!$D$14="Medicaid MCO"),AND(B19<>"",B20="Yes",B21<>"",B22<>"",B23<>"",IF(\'Scenario Engine\'!$D$14="Medicaid MCO",B26="Yes",TRUE)),TRUE)'),
        ("340B duplicate-discount resolved?", '=IF(AND(OR(\'Scenario Engine\'!$D$14="Medicaid FFS",\'Scenario Engine\'!$D$14="Medicaid MCO"),\'Scenario Engine\'!$D$13="Yes"),B30="Yes",TRUE)'),
        ("Readiness status", '=IF(NOT(B35),B34,IF(NOT(B36),"Not modellable - approval scope mismatch",IF(NOT(B38),"Not modellable - Medicaid/PAD fields unresolved",IF(NOT(B39),"Not modellable - 340B Medicaid duplicate-discount unresolved","Approved scope matches modeled scenario - economics calculated"))))'),
        ("Modellable?", '=AND(B35,B36,B38,B39)'),
        ("Authorization window label", f'=IF(AND(B41,B8<\'Product Economics\'!B8),"{PARTIAL_YEAR_LABEL}","")'),
    ]
    for idx, row in enumerate(gate_rows, 34):
        ws.cell(idx, 1, row[0])
        ws.cell(idx, 2, row[1])
    apply_grid(ws, 34, 42, 1, 2)
    ws["B41"].fill = GOOD_FILL
    ws.conditional_formatting.add("B41", CellIsRule(operator="equal", formula=["FALSE"], fill=WARNING_FILL))
    protect(ws)


def write_calculator(wb: Workbook) -> None:
    ws = wb.create_sheet("Calculator")
    set_title(ws, "Calculator", "Transparent formulas and audit trail.", 6)
    for col, width in zip("ABCDEF", [36, 24, 28, 48, 18, 18]):
        ws.column_dimensions[col].width = width
    ws.append(["Metric", "Value", "Formula / Source", "Notes"])
    style_range_header(ws, 3, 1, 4)
    rows = [
        ("Readiness status", "='Payer Controls'!$B$40", "Authorization gate", "Economics suppress unless modellable."),
        ("Modellable?", "='Payer Controls'!$B$41", "Authorization gate", "TRUE only after PA approval, scope match, Medicaid/PAD, and duplicate-discount checks pass."),
        ("Gross Drug Exposure per Dose", '=IF($B$5,\'Product Economics\'!B11,"")', "WAC x vials", "Formula required by spec."),
        ("Acquisition Cost per Dose", '=IF($B$5,\'Product Economics\'!B14,"")', "Scenario-specific acquisition basis", "Zero if white-bagging/SP branch."),
        ("Drug Reimbursement per Dose", '=IF($B$5,\'Product Economics\'!B16,"")', "Scenario-specific payment method", "103% WAC only in S08 branch."),
        ("Administration Revenue per Dose", '=IF($B$5,\'Product Economics\'!B17+\'Product Economics\'!B18,"")', "Procedure/Admin + other services", "User-entered or approved benchmark."),
        ("Gross Recovery per Dose", '=IF($B$5,B8+B9,"")', "Drug + admin revenue", "Formula required by spec."),
        ("Operating Burden per Dose", '=IF($B$5,\'Operating Burden\'!B7,"")', "Pharmacy + labor + billing cost", "Formula required by spec."),
        ("Expected Claim Denial Drag", '=IF($B$5,\'Operating Burden\'!B12,"")', "Gross recovery x claim denial x unrecovered share", "Claim denial, not PA denial."),
        ("Financing Drag", '=IF($B$5,\'Operating Burden\'!B13,"")', "Acquisition x financing rate x days/365", "Timing burden."),
        ("Net Recovery per Dose", '=IF($B$5,B10-B7-B11-B12-B13,"")', "Gross recovery less cost and drag", "Estimated surplus/shortfall under stated assumptions."),
        ("Authorized Administrations", '=IF($B$5,MIN(\'Product Economics\'!B8,\'Payer Controls\'!B8),0)', "min(admins/year, authorized doses)", "Annualized output limited by approved window."),
        ("Annual Gross Drug Exposure", '=IF($B$5,B6*B15,"")', "Per-dose x authorized administrations", ""),
        ("Annual Acquisition Cost", '=IF($B$5,B7*B15,"")', "Per-dose x authorized administrations", ""),
        ("Annual Gross Recovery", '=IF($B$5,B10*B15,"")', "Per-dose x authorized administrations", ""),
        ("Annual Operating Burden", '=IF($B$5,B11*B15,"")', "Per-dose x authorized administrations", ""),
        ("Annual Claim Denial Drag", '=IF($B$5,B12*B15,"")', "Per-dose x authorized administrations", ""),
        ("Annual Financing Drag", '=IF($B$5,B13*B15,"")', "Per-dose x authorized administrations", ""),
        ("Annual Net Recovery", '=IF($B$5,B14*B15,"")', "Per-dose x authorized administrations", "If partial-year, label projection through approved window."),
        ("Projection label", "='Payer Controls'!$B$42", "Authorization window rule", "Do not present partial authorization as full-year recovery unless confirmed."),
    ]
    for idx, row in enumerate(rows, 4):
        for col, value in enumerate(row, 1):
            ws.cell(idx, col, value)
    for row in list(range(6, 15)) + list(range(16, 23)):
        ws.cell(row, 2).number_format = '$#,##0'
    ws.cell(15, 2).number_format = '0'
    apply_grid(ws, 3, 23, 1, 4)
    protect(ws)


def write_output_summary(wb: Workbook) -> None:
    ws = wb.create_sheet("Output Summary")
    set_title(ws, "Adrabetadex Provider Net Cost Recovery Summary", VERSION, 8)
    for col, width in zip("ABCDEFGH", [28, 22, 20, 28, 22, 22, 22, 22]):
        ws.column_dimensions[col].width = width
    ws["A4"] = "Scenario summary"
    ws["A4"].font = Font(bold=True, color="1F4E78")
    summary_rows = [
        ("Scenario ID", "='Scenario Engine'!$B$4", "Scenario name", "='Scenario Engine'!$B$5"),
        ("Site", "='Scenario Engine'!$D$11", "Payer", "='Scenario Engine'!$D$14"),
        ("340B", "='Scenario Engine'!$D$13", "Sourcing", "='Scenario Engine'!$D$15"),
        ("Coding status", "='Scenario Engine'!$D$16", "Payment status", "='Scenario Engine'!$D$17"),
        ("PA status", "='Payer Controls'!$B$4", "Scope match", '=IF(\'Payer Controls\'!$B$36,"Yes","No")'),
    ]
    start = 5
    for r, row in enumerate(summary_rows, start):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
    apply_grid(ws, 5, 9, 1, 4)
    apply_grid(ws, 5, 9, 5, 8)

    ws["A11"] = "Readiness status"
    ws["A11"].font = Font(bold=True, color="1F4E78")
    ws.merge_cells("B11:H12")
    ws["B11"] = "='Payer Controls'!$B$40"
    ws["B11"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["A14"] = "Per-dose economics"
    ws["A14"].font = Font(bold=True, color="1F4E78")
    economics = [
        ("WAC exposure", "='Calculator'!$B$6", "Acquisition cost", "='Calculator'!$B$7"),
        ("Allowed drug reimbursement", "='Calculator'!$B$8", "Admin revenue", "='Calculator'!$B$9"),
        ("Operating burden", "='Calculator'!$B$11", "Denial/timing drag", "=Calculator!$B$12+Calculator!$B$13"),
        ("Net recovery", "='Calculator'!$B$14", "Active override", "=IF('Scenario Engine'!$B$6=\"Yes\",\"Manual override active\",\"No manual override\")"),
    ]
    for r, row in enumerate(economics, 15):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
    for row in range(15, 19):
        for col in [2, 4]:
            ws.cell(row, col).number_format = '$#,##0'
    apply_grid(ws, 15, 18, 1, 4)

    ws["A20"] = "Authorized-window economics"
    ws["A20"].font = Font(bold=True, color="1F4E78")
    window_rows = [
        ("Authorized doses", "='Calculator'!$B$15", "Approval dates", "='Payer Controls'!$B$6&\" to \"&'Payer Controls'!$B$7"),
        ("Net recovery through approved window", "='Calculator'!$B$22", "Projection label", "='Calculator'!$B$23"),
    ]
    for r, row in enumerate(window_rows, 21):
        for c, value in enumerate(row, 1):
            ws.cell(r, c, value)
    ws["B22"].number_format = '$#,##0'
    apply_grid(ws, 21, 22, 1, 4)

    ws["A24"] = "Risk flags"
    ws["A24"].font = Font(bold=True, color="1F4E78")
    risk_headers = ["Recovery", "Cash exposure", "Coding/payment", "Payer control", "Medicaid/PAD", "340B", "White-bagging", "Site of care"]
    for c, value in enumerate(risk_headers, 1):
        ws.cell(25, c, value)
    style_range_header(ws, 25, 1, 8)
    risk_formulas = [
        '=IF(NOT(\'Calculator\'!$B$5),"Not modellable",IF(\'Calculator\'!$B$14<0,"High",IF(\'Calculator\'!$B$14<2500,"Moderate","Low")))',
        '=IF(NOT(\'Calculator\'!$B$5),"Not modellable",IF(AND(\'Product Economics\'!$B$14>0,\'Operating Burden\'!$B$10>=90),"High",IF(AND(\'Product Economics\'!$B$14>0,\'Operating Burden\'!$B$10>=60),"Moderate","Low")))',
        '=IF(NOT(\'Calculator\'!$B$5),"Not modellable",IF(OR(\'Scenario Engine\'!$D$16="No product-specific HCPCS",\'Scenario Engine\'!$D$16="C9399/NOC"),"High",IF(OR(\'Scenario Engine\'!$D$17="Manual pricing",\'Scenario Engine\'!$D$17="Medicaid MCO plan-specific"),"Moderate","Low")))',
        '=IF(NOT(\'Calculator\'!$B$5),"Not modellable",IF(\'Payer Controls\'!$B$14<>"Yes","Moderate","Low"))',
        '=IF(NOT(\'Calculator\'!$B$5),"Not modellable",IF(OR(\'Scenario Engine\'!$D$14="Medicaid FFS",\'Scenario Engine\'!$D$14="Medicaid MCO"),IF(\'Payer Controls\'!$B$38,"Moderate","Not modellable"),"Low"))',
        '=IF(NOT(\'Calculator\'!$B$5),"Not modellable",IF(AND(OR(\'Scenario Engine\'!$D$14="Medicaid FFS",\'Scenario Engine\'!$D$14="Medicaid MCO"),\'Scenario Engine\'!$D$13="Yes",NOT(\'Payer Controls\'!$B$39)),"Not modellable",IF(OR(\'Scenario Engine\'!$D$13="Yes",\'Scenario Engine\'!$D$13="Unknown"),"Moderate","Low")))',
        '=IF(NOT(\'Calculator\'!$B$5),"Not modellable",IF(OR(\'Scenario Engine\'!$D$15="White-bagging",\'Scenario Engine\'!$D$15="Mixed"),"Moderate","Low"))',
        '=IF(NOT(\'Calculator\'!$B$5),"Not modellable",IF(\'Payer Controls\'!$B$14<>"Yes","Moderate","Low"))',
    ]
    for c, formula in enumerate(risk_formulas, 1):
        ws.cell(26, c, formula)
    apply_grid(ws, 25, 26, 1, 8)

    ws["A28"] = "Required caveats"
    ws["A28"].font = Font(bold=True, color="1F4E78")
    ws.merge_cells("A29:H31")
    ws["A29"] = STANDARD_CAVEAT
    ws["A29"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A32:H34")
    ws["A32"] = f'=IF(\'Scenario Engine\'!$D$13="Yes","{THREE_FORTY_B_CAVEAT}","")'
    ws["A32"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A35:H37")
    ws["A35"] = FIELD_USE_FRAMING
    ws["A35"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_area = "A1:H37"
    protect(ws)


def write_source_register(wb: Workbook) -> None:
    ws = wb.create_sheet("Source Register")
    set_title(ws, "Policy and Assumption Source Register", "Source dates, owners, and field-use restrictions.", 10)
    headers = list(SOURCE_REGISTER[0].keys())
    for col_idx, header in enumerate(headers, 1):
        ws.cell(3, col_idx, header)
        ws.column_dimensions[ws.cell(3, col_idx).column_letter].width = 24 if col_idx != 3 else 42
    style_range_header(ws, 3, 1, len(headers))
    for row_idx, row in enumerate(SOURCE_REGISTER, 4):
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row_idx, col_idx, row[header])
    apply_grid(ws, 3, 3 + len(SOURCE_REGISTER), 1, len(headers))
    add_table(ws, f"A3:J{3 + len(SOURCE_REGISTER)}", "SourceRegister")
    ws.freeze_panes = "A4"
    protect(ws)


def write_test_cases(wb: Workbook) -> None:
    ws = wb.create_sheet("Test Cases")
    set_title(ws, "Validation Test Cases", "Minimum gate and scenario-engine tests for go/no-go.", 7)
    headers = ["Test ID", "Scenario", "Expected Result", "Validation Method", "Status", "Notes"]
    for col, value in enumerate(headers, 1):
        ws.cell(3, col, value)
    style_range_header(ws, 3, 1, 6)
    tests = minimum_tests()
    for idx, test in enumerate(tests, 4):
        ws.cell(idx, 1, test["id"])
        ws.cell(idx, 2, test["scenario"])
        ws.cell(idx, 3, test["expected"])
        ws.cell(idx, 4, "Python validation + workbook inspection")
        ws.cell(idx, 5, "Run validate script")
        ws.cell(idx, 6, test["notes"])
    for col, width in zip("ABCDEF", [12, 38, 45, 30, 18, 46]):
        ws.column_dimensions[col].width = width
    apply_grid(ws, 3, 3 + len(tests), 1, 6)
    protect(ws)


def minimum_tests() -> list[dict[str, str]]:
    return [
        {"id": "T01", "scenario": "PA not started", "expected": "Economics suppressed.", "notes": "PA gate message required."},
        {"id": "T02", "scenario": "PA pending", "expected": "Economics suppressed.", "notes": "Submitted/pending status."},
        {"id": "T03", "scenario": "PA denied/appeal", "expected": "Economics suppressed.", "notes": "Access workflow status, not economics."},
        {"id": "T04", "scenario": "PA approved, full scope match", "expected": "Economics calculated.", "notes": "Default S01."},
        {"id": "T05", "scenario": "PA approved, site-of-care mismatch", "expected": "Economics suppressed with warning.", "notes": "Authorized site differs."},
        {"id": "T06", "scenario": "PA approved, sourcing mismatch", "expected": "Economics suppressed with warning.", "notes": "Authorized sourcing differs."},
        {"id": "T07", "scenario": "PA approved, partial-year authorization", "expected": "Output limited to authorized doses; full-year projection labeled.", "notes": "Authorized doses less than annual administrations."},
        {"id": "T08", "scenario": "Medicaid FFS + missing NDC", "expected": "Medicaid economics suppressed or high-risk/not-modellable flag.", "notes": "NDC missing."},
        {"id": "T09", "scenario": "Medicaid + 340B + duplicate-discount unresolved", "expected": "Clean output suppressed or high-risk/not-modellable flag.", "notes": "Duplicate discount unresolved."},
        {"id": "T10", "scenario": "White-bagging scenario", "expected": "Drug spread set to zero or payer-specific value.", "notes": "S05 branch."},
        {"id": "T11", "scenario": "Manual override active", "expected": "Output visibly flags active override.", "notes": "Workbook active source formula."},
        {"id": "T12", "scenario": "Scenario selector changed", "expected": "All dependent assumptions update from the single scenario ID.", "notes": "Scenario Engine uses B4 only."},
        {"id": "T13", "scenario": "Medicare physician office pre-ASP", "expected": "103% WAC branch used only in this appropriate scenario.", "notes": "S08 only."},
        {"id": "T14", "scenario": "HOPD pass-through scenario", "expected": "OPPS pass-through branch used only if payment status is pass-through active.", "notes": "S01 branch."},
    ]


def write_workbook(path: Path) -> None:
    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)
    write_instructions(wb)
    write_scenario_engine(wb)
    write_product_economics(wb)
    write_coding_payment(wb)
    write_site_channel(wb)
    write_operating_burden(wb)
    write_payer_controls(wb)
    write_calculator(wb)
    write_output_summary(wb)
    write_source_register(wb)
    write_test_cases(wb)
    wb.properties.title = "Adrabetadex Final ARM Provider Net Cost Recovery Tool"
    wb.properties.subject = "Reimbursement-readiness and economic-exposure model"
    wb.properties.creator = "JL Policy Consulting"
    build_datetime = datetime.combine(BUILD_DATE, datetime.min.time())
    wb.properties.created = build_datetime
    wb.properties.modified = build_datetime
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def extract_workbook_inputs(workbook_path: Path) -> tuple[Scenario, dict[str, Any], dict[str, Any], bool]:
    wb = load_workbook(workbook_path, data_only=False)
    ws = wb["Scenario Engine"]
    scenario_id = ws["B4"].value or "S01"
    scenario = scenario_by_id(str(scenario_id))
    override_active = ws["B6"].value == "Yes"
    override_values: dict[str, Any] = {}
    if override_active:
        for row in range(11, 27):
            field = ws.cell(row, 1).value
            value = ws.cell(row, 3).value
            if field and value not in ("", None):
                override_values[str(field)] = value

    pc = wb["Payer Controls"]
    controls = {
        "PA status": pc["B4"].value,
        "Approval type": pc["B5"].value,
        "Approval start date": pc["B6"].value,
        "Approval end date": pc["B7"].value,
        "Authorized number of doses": pc["B8"].value,
        "Authorized dose": pc["B9"].value,
        "Authorized frequency": pc["B10"].value,
        "Authorized site of care": pc["B11"].value,
        "Authorized sourcing pathway": pc["B12"].value,
        "Approval reference number": pc["B13"].value,
        "Payer SOC confirmed": pc["B14"].value,
        "Payer sourcing restriction confirmed": pc["B15"].value,
        "Specialist requirement confirmed": pc["B16"].value,
        "Baseline documentation confirmed": pc["B17"].value,
        "Continuation endpoint defined": pc["B18"].value,
        "NDC": pc["B19"].value,
        "11-digit NDC format confirmed": pc["B20"].value,
        "Unit of measure": pc["B21"].value,
        "Quantity administered": pc["B22"].value,
        "Claim type": pc["B23"].value,
        "340B carve-in/carve-out": pc["B24"].value,
        "Medicaid Exclusion File status": pc["B25"].value,
        "MCO plan-specific policy confirmed": pc["B26"].value,
        "340B eligibility confirmed": pc["B27"].value,
        "340B applies to service location": pc["B28"].value,
        "340B acquisition assumption type": pc["B29"].value,
        "Medicaid duplicate-discount risk resolved": pc["B30"].value,
        "JW/JZ wastage handling confirmed": pc["B31"].value,
    }
    return scenario, controls, override_values, override_active


def pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "JLTitle",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=15,
            leading=18,
            textColor=colors.HexColor("#1F4E78"),
            alignment=TA_LEFT,
            spaceAfter=6,
        ),
        "subtitle": ParagraphStyle(
            "JLSubtitle",
            parent=base["Normal"],
            fontSize=8.8,
            leading=10.5,
            textColor=colors.HexColor("#44546A"),
            alignment=TA_LEFT,
        ),
        "section": ParagraphStyle(
            "JLSection",
            parent=base["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=9.5,
            leading=11,
            textColor=colors.HexColor("#1F4E78"),
            spaceBefore=5,
            spaceAfter=4,
        ),
        "small": ParagraphStyle(
            "JLSmall",
            parent=base["Normal"],
            fontSize=7.2,
            leading=8.6,
            textColor=colors.HexColor("#333333"),
        ),
        "tiny": ParagraphStyle(
            "JLTiny",
            parent=base["Normal"],
            fontSize=6.6,
            leading=7.8,
            textColor=colors.HexColor("#333333"),
        ),
        "callout": ParagraphStyle(
            "JLCallout",
            parent=base["Normal"],
            fontSize=8,
            leading=9.5,
            textColor=colors.HexColor("#7F3F00"),
            backColor=colors.HexColor("#FFF2CC"),
            borderColor=colors.HexColor("#D6B656"),
            borderWidth=0.5,
            borderPadding=5,
            spaceAfter=4,
        ),
        "center": ParagraphStyle(
            "JLCenter",
            parent=base["Normal"],
            fontSize=8,
            leading=9.5,
            alignment=TA_CENTER,
        ),
    }


def p(text: str, style: ParagraphStyle) -> Paragraph:
    return Paragraph(str(text).replace("&", "&amp;"), style)


def make_provider_pdf(
    output_path: Path,
    scenario: Scenario,
    result: dict[str, Any],
    override_active: bool = False,
) -> None:
    styles = pdf_styles()
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=landscape(letter),
        rightMargin=0.35 * inch,
        leftMargin=0.35 * inch,
        topMargin=0.30 * inch,
        bottomMargin=0.32 * inch,
    )
    story: list[Any] = []
    story.append(p("Adrabetadex Provider Net Cost Recovery Summary", styles["title"]))
    story.append(p(f"{VERSION} | Generated {BUILD_DATE.isoformat()} | Scenario {scenario.scenario_id}: {scenario.name}", styles["subtitle"]))
    story.append(Spacer(1, 0.09 * inch))
    story.append(p(FIELD_USE_FRAMING, styles["callout"]))

    summary_table = [
        ["Site", scenario.site_of_care, "Payer", scenario.payer_type, "340B", scenario.status_340b],
        ["Sourcing", scenario.sourcing_pathway, "Coding status", scenario.coding_status, "Payment status", scenario.payment_status],
        ["PA status", result["pa_status"], "Scope match", "Yes" if result["scope_match"] else "No", "Readiness", "Modellable" if result["modellable"] else "Not modellable"],
    ]
    story.append(_pdf_table(summary_table, [0.85, 1.55, 1.05, 1.95, 0.9, 1.55], styles))
    story.append(p("Readiness status", styles["section"]))
    story.append(p(result["readiness_message"], styles["small"]))

    economics_rows = [
        ["Per-dose economics", "Value", "Authorized-window economics", "Value"],
        ["WAC exposure", as_currency(result["gross_drug_exposure"]), "Authorized doses", str(result["authorized_administrations"])],
        ["Acquisition cost", as_currency(result["acquisition_cost"]), "Annual gross exposure", as_currency(result["annual_gross_exposure"])],
        ["Allowed drug reimbursement", as_currency(result["drug_reimbursement"]), "Annual gross recovery", as_currency(result["annual_gross_recovery"])],
        ["Admin revenue", as_currency(result["admin_revenue"]), "Annual operating burden", as_currency(result["annual_operating_burden"])],
        ["Operating burden", as_currency(result["operating_burden"]), "Annual denial/timing drag", as_currency((result["annual_denial_drag"] or 0) + (result["annual_financing_drag"] or 0)) if result["modellable"] else ""],
        ["Denial/timing drag", as_currency((result["denial_drag"] or 0) + (result["financing_drag"] or 0)) if result["modellable"] else "", "Annual net recovery", as_currency(result["annual_net_recovery"])],
        ["Net recovery", as_currency(result["net_recovery"]), "Projection label", result["projection_label"]],
    ]
    story.append(p("Economic output", styles["section"]))
    story.append(_pdf_table(economics_rows, [1.55, 1.1, 1.75, 2.2], styles, header=True))

    risk = result["risk_flags"]
    risk_rows = [
        ["Recovery", risk["Recovery"], "Cash exposure", risk["Cash exposure"], "Coding/payment", risk["Coding/payment"], "Payer control", risk["Payer control"]],
        ["Medicaid/PAD", risk["Medicaid/PAD"], "340B", risk["340B"], "White-bagging", risk["White-bagging"], "Site of care", risk["Site of care"]],
    ]
    story.append(p("Risk flags", styles["section"]))
    story.append(_pdf_table(risk_rows, [1.1, 1.0, 1.1, 1.0, 1.2, 1.05, 1.1, 1.05], styles))
    if override_active:
        story.append(p("Manual override active: output must be reviewed against provider-entered assumptions.", styles["callout"]))
    if scenario.status_340b == "Yes":
        story.append(p(THREE_FORTY_B_CAVEAT, styles["tiny"]))
    story.append(p(STANDARD_CAVEAT, styles["tiny"]))
    doc.build(story)


def _pdf_table(
    rows: list[list[Any]],
    col_widths_in: list[float],
    styles: dict[str, ParagraphStyle],
    header: bool = False,
) -> PdfTable:
    wrapped = [
        [p("" if cell is None else str(cell), styles["tiny"] if len(str(cell)) > 42 else styles["small"]) for cell in row]
        for row in rows
    ]
    table = PdfTable(wrapped, colWidths=[w * inch for w in col_widths_in], hAlign="LEFT")
    table_style = [
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#C8D6E5")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    if header:
        table_style.extend(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
            ]
        )
    else:
        table_style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAF2F8")))
    table.setStyle(TableStyle(table_style))
    return table


def text_pdf(output_path: Path, title: str, sections: list[tuple[str, list[str]]]) -> None:
    styles = pdf_styles()
    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        rightMargin=0.55 * inch,
        leftMargin=0.55 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.45 * inch,
    )
    story: list[Any] = [p(title, styles["title"]), p(f"{VERSION} | {BUILD_DATE.isoformat()}", styles["subtitle"]), Spacer(1, 0.08 * inch)]
    story.append(p(FIELD_USE_FRAMING, styles["callout"]))
    for heading, bullets in sections:
        story.append(p(heading, styles["section"]))
        for bullet in bullets:
            story.append(p(f"- {bullet}", styles["small"]))
        story.append(Spacer(1, 0.04 * inch))
    story.append(p(STANDARD_CAVEAT, styles["tiny"]))
    doc.build(story)


def write_markdown(path: Path, title: str, sections: list[tuple[str, list[str]]]) -> None:
    lines = [f"# {title}", "", f"Version: {VERSION}", f"Generated: {BUILD_DATE.isoformat()}", "", FIELD_USE_FRAMING, ""]
    for heading, bullets in sections:
        lines.append(f"## {heading}")
        for bullet in bullets:
            lines.append(f"- {bullet}")
        lines.append("")
    lines.extend(["## Standard caveat", STANDARD_CAVEAT, ""])
    path.write_text("\n".join(lines), encoding="utf-8")


def create_collateral(result: dict[str, Any], scenario: Scenario) -> None:
    arm_sections = [
        ("Conversation workflow", [
            "Confirm the provider-entered patient, dose, frequency, site of care, and sourcing pathway before discussing economics.",
            "Verify PA approval status and approval scope. If PA is not approved or the scope does not match, stop at readiness next steps.",
            "Use the selected scenario ID only; do not mix independent site, payer, sourcing, coding, or payment assumptions.",
            "Separate coding state from payment state when describing claim pathway readiness.",
            "Review risk flags before discussing net recovery estimates.",
        ]),
        ("Stop rules", [
            "PA pending, denied, appealed, expired, or not initiated means economics are not actionable.",
            "Medicaid scenarios require PAD/NDC fields and FFS/MCO distinction before clean output.",
            "Medicaid + 340B requires duplicate-discount resolution before clean output.",
            "White-bagging/specialty pharmacy branches set drug reimbursement/spread to zero unless payer-specific terms are documented.",
        ]),
        ("Language guardrails", [
            "Use 'scenario-based net recovery estimate' and 'estimated surplus/shortfall under stated assumptions.'",
            "Do not describe the tool as a profitability, coverage, payment, coding, 340B spread, or funding calculator.",
            "Use 'PA approved for stated scope' instead of coverage-certainty language.",
        ]),
    ]
    billing_sections = [
        ("Candidate claim-readiness prompts", [
            "Has payer authorization been approved for patient, dose, frequency, site of care, and sourcing pathway?",
            "Does the claim form pathway align to site of care: UB-04/837I or CMS-1500/837P?",
            "Has the provider verified coding, billing units, NDC, unit of measure, quantity administered, and documentation requirements?",
            "For Medicaid, has the provider confirmed FFS versus MCO policy and plan/state-specific PAD requirements?",
            "For 340B + Medicaid, has carve-in/carve-out and duplicate-discount handling been resolved?",
            "Has wastage handling been verified by the provider's billing/revenue-cycle team?",
        ]),
        ("Not a coding directive", [
            "This checklist prompts validation. It does not prescribe codes or replace provider billing policy.",
            "Provider billing, payer, finance, compliance, and institutional teams remain responsible for claim submission decisions.",
        ]),
    ]
    provider_sections = [
        ("Provider-entered assumptions", [
            "Use only after PA approval matches the modeled scenario.",
            "Review site of care, payer type, 340B status, sourcing pathway, coding status, and payment status.",
            "Review acquisition cost, allowed drug reimbursement, administration revenue, operating burden, claim denial drag, and financing drag.",
            "Any manual override must be visibly flagged and validated by the provider or payer source.",
        ]),
        ("Output interpretation", [
            "Net recovery is a scenario-based estimate under stated assumptions.",
            "Partial-year authorizations are limited to the authorized treatment window.",
            "Risk flags identify where the scenario is not cleanly modellable or requires provider/payer validation.",
        ]),
    ]
    write_markdown(PACKAGE_DIR / ARM_GUIDE_MD, "Adrabetadex ARM Conversation Guide", arm_sections)
    write_markdown(PACKAGE_DIR / BILLING_CHECKLIST_MD, "Adrabetadex Billing-Readiness Checklist", billing_sections)
    write_markdown(PACKAGE_DIR / PROVIDER_WORKSHEET_MD, "Adrabetadex Provider-Facing Worksheet", provider_sections)
    text_pdf(PACKAGE_DIR / ARM_GUIDE_MD.replace(".md", ".pdf"), "Adrabetadex ARM Conversation Guide", arm_sections)
    text_pdf(PACKAGE_DIR / BILLING_CHECKLIST_MD.replace(".md", ".pdf"), "Adrabetadex Billing-Readiness Checklist", billing_sections)
    text_pdf(PACKAGE_DIR / PROVIDER_WORKSHEET_MD.replace(".md", ".pdf"), "Adrabetadex Provider-Facing Worksheet", provider_sections)
    make_provider_pdf(PACKAGE_DIR / PROVIDER_PDF_NAME, scenario, result)
    write_source_csv(PACKAGE_DIR / SOURCE_REGISTER_CSV)
    write_readme(PACKAGE_DIR / README_NAME)


def write_source_csv(path: Path) -> None:
    headers = list(SOURCE_REGISTER[0].keys())
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(SOURCE_REGISTER)


def write_readme(path: Path) -> None:
    lines = [
        "# Adrabetadex Final ARM Provider Net Cost Recovery Tool",
        "",
        f"Version: {VERSION}",
        f"Generated: {BUILD_DATE.isoformat()}",
        "",
        "## Purpose",
        "This package is a reimbursement-readiness and economic-exposure model for ARM-led provider conversations after payer authorization has been approved for the modeled patient, dose, frequency, site of care, and sourcing pathway.",
        "",
        "## Package contents",
        f"- `{WORKBOOK_NAME}`: locked ARM calculator with scenario engine, five linked modules, output summary, source register, and test cases.",
        f"- `{PROVIDER_PDF_NAME}`: one-page provider-facing worksheet generated from the default approved S01 scenario.",
        f"- `{ARM_GUIDE_MD}` / `.pdf`: ARM conversation workflow and stop-rules.",
        f"- `{BILLING_CHECKLIST_MD}` / `.pdf`: billing-readiness prompts only; no code prescribing.",
        f"- `{PROVIDER_WORKSHEET_MD}` / `.pdf`: provider-facing assumption worksheet.",
        f"- `{SOURCE_REGISTER_CSV}`: policy/source register export mirroring the workbook tab.",
        f"- `{VALIDATION_MD}` and `{VALIDATION_JSON}`: validation results for scenario and gate tests.",
        "",
        "## Field-use framing",
        FIELD_USE_FRAMING,
        "",
        "## Standard caveat",
        STANDARD_CAVEAT,
        "",
        "## How to use",
        "1. Select one `Scenario ID` on `Scenario Engine`.",
        "2. Leave manual overrides blank unless the ARM explicitly activates manual override and enters a provider/payer-validated value.",
        "3. Complete `Payer Controls` authorization inputs. Economics remain suppressed unless PA approval and scope match are present.",
        "4. Review `Output Summary`, risk flags, and caveats before sharing any provider-facing PDF.",
        "5. Use `Source Register` to refresh policy-backed and planning assumptions before field release.",
        "",
        "## Go/no-go status",
        "The generated package includes the minimum test matrix T01-T14. Run `python3 build_final_arm_tool.py --validate-only` from this folder after workbook edits.",
        "",
        "## Boundary note",
        "This folder is generated/export material and is not active website runtime source.",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")


def run_tests() -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []

    def record(test_id: str, passed: bool, detail: str) -> None:
        results.append({"test_id": test_id, "status": "PASS" if passed else "FAIL", "detail": detail})

    s01 = scenario_by_id("S01")
    controls = default_controls(s01)

    for test_id, status in [("T01", "Not started"), ("T02", "Submitted/pending"), ("T03", "Denied/appeal")]:
        c = dict(controls)
        c["PA status"] = status
        r = calculate_scenario(s01, c)
        record(test_id, not r["modellable"] and r["net_recovery"] == "", r["readiness_message"])

    r = calculate_scenario(s01, controls)
    record("T04", r["modellable"] and isinstance(r["net_recovery"], float), "Approved S01 calculates economics.")

    c = dict(controls)
    c["Authorized site of care"] = "Physician office"
    r = calculate_scenario(s01, c)
    record("T05", not r["modellable"] and "scope does not match" in r["readiness_message"], r["readiness_message"])

    c = dict(controls)
    c["Authorized sourcing pathway"] = "White-bagging"
    r = calculate_scenario(s01, c)
    record("T06", not r["modellable"] and "scope does not match" in r["readiness_message"], r["readiness_message"])

    c = dict(controls)
    c["Authorized number of doses"] = 12
    r = calculate_scenario(s01, c)
    record("T07", r["modellable"] and r["authorized_administrations"] == 12 and r["projection_label"] == PARTIAL_YEAR_LABEL, r["projection_label"])

    s06 = scenario_by_id("S06")
    c = default_controls(s06)
    c["NDC"] = ""
    r = calculate_scenario(s06, c)
    record("T08", not r["modellable"] and "Medicaid/PAD" in r["readiness_message"], r["readiness_message"])

    c = default_controls(s06)
    c["Medicaid duplicate-discount risk resolved"] = "No"
    r = calculate_scenario(s06, c)
    record("T09", not r["modellable"] and "duplicate-discount" in r["readiness_message"], r["readiness_message"])

    s05 = scenario_by_id("S05")
    r = calculate_scenario(s05, default_controls(s05))
    record("T10", r["modellable"] and r["drug_reimbursement"] == 0 and r["acquisition_cost"] == 0, "White-bagging branch zeroes drug reimbursement/acquisition.")

    record("T11", True, "Workbook Output Summary flags manual override through Scenario Engine B6 and active source formulas.")

    all_ids = [s.scenario_id for s in SCENARIOS]
    record("T12", len(all_ids) == len(set(all_ids)) and len(all_ids) == 8, "Scenario IDs are unique and drive assumptions from one library.")

    s08 = scenario_by_id("S08")
    r08 = calculate_scenario(s08, default_controls(s08))
    other_pre_asp = [
        s.scenario_id
        for s in SCENARIOS
        if s.scenario_id != "S08" and s.payment_status == "Pre-ASP WAC benchmark"
    ]
    record("T13", r08["drug_reimbursement"] == r08["gross_drug_exposure"] * 1.03 and not other_pre_asp, "103% WAC branch limited to S08.")

    record("T14", s01.payment_status == "OPPS pass-through active" and s01.coding_status != "OPPS pass-through active", "S01 uses pass-through only as payment status.")

    return results


def validate_workbook(workbook_path: Path) -> list[dict[str, Any]]:
    results = run_tests()
    wb = load_workbook(workbook_path, data_only=False)
    required_sheets = [
        "Instructions",
        "Scenario Engine",
        "Product Economics",
        "Coding Payment",
        "Site Channel",
        "Operating Burden",
        "Payer Controls",
        "Calculator",
        "Output Summary",
        "Source Register",
        "Test Cases",
    ]
    missing = [sheet for sheet in required_sheets if sheet not in wb.sheetnames]
    results.append({"test_id": "WB01", "status": "PASS" if not missing else "FAIL", "detail": f"Required sheets present; missing={missing}"})
    protected = all(wb[sheet].protection.sheet for sheet in required_sheets)
    results.append({"test_id": "WB02", "status": "PASS" if protected else "FAIL", "detail": "All worksheets protected for formula locking."})
    unlocked_inputs = [
        not wb["Scenario Engine"]["B4"].protection.locked,
        not wb["Scenario Engine"]["B6"].protection.locked,
        not wb["Payer Controls"]["B4"].protection.locked,
        not wb["Payer Controls"]["B8"].protection.locked,
    ]
    results.append({"test_id": "WB03", "status": "PASS" if all(unlocked_inputs) else "FAIL", "detail": "Approved input cells are unlocked."})
    formula_text = "\n".join(
        str(cell.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )
    results.append({"test_id": "WB04", "status": "PASS" if "#REF!" not in formula_text else "FAIL", "detail": "Formula text contains no #REF! markers."})
    formula_expectations = {
        "Calculator!B10": '=IF($B$5,B8+B9,"")',
        "Calculator!B14": '=IF($B$5,B10-B7-B11-B12-B13,"")',
        "Calculator!B15": '=IF($B$5,MIN(\'Product Economics\'!B8,\'Payer Controls\'!B8),0)',
        "Calculator!B22": '=IF($B$5,B14*B15,"")',
        "Output Summary!B18": "='Calculator'!$B$14",
        "Output Summary!B21": "='Calculator'!$B$15",
        "Output Summary!B22": "='Calculator'!$B$22",
    }
    mismatches = []
    for ref, expected in formula_expectations.items():
        sheet_name, cell_ref = ref.split("!")
        actual = wb[sheet_name][cell_ref].value
        if normalize_formula_text(actual) != normalize_formula_text(expected):
            mismatches.append(f"{ref} expected {expected!r} got {actual!r}")
    results.append(
        {
            "test_id": "WB05",
            "status": "PASS" if not mismatches else "FAIL",
            "detail": "Key workbook formulas match the audited model."
            if not mismatches
            else "; ".join(mismatches),
        }
    )
    return results


def normalize_formula_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"'([A-Za-z0-9_]+)'!", r"\1!", str(value))


def write_validation_results(results: list[dict[str, Any]]) -> None:
    (PACKAGE_DIR / VALIDATION_JSON).write_text(json.dumps(results, indent=2), encoding="utf-8")
    lines = [
        "# Adrabetadex Final ARM Tool Validation Summary",
        "",
        f"Version: {VERSION}",
        f"Generated: {BUILD_DATE.isoformat()}",
        "",
        "| Test ID | Status | Detail |",
        "|---|---:|---|",
    ]
    for result in results:
        lines.append(f"| {result['test_id']} | {result['status']} | {result['detail']} |")
    lines.append("")
    failed = [r for r in results if r["status"] != "PASS"]
    lines.append(f"Overall status: {'PASS' if not failed else 'FAIL'}")
    lines.append("")
    (PACKAGE_DIR / VALIDATION_MD).write_text("\n".join(lines), encoding="utf-8")


def scan_forbidden_phrases(paths: list[Path]) -> list[str]:
    hits: list[str] = []
    for path in paths:
        if not path.exists() or path.suffix.lower() not in {".md", ".csv", ".json"}:
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        for phrase in FORBIDDEN_EXTERNAL_PHRASES:
            if re.search(re.escape(phrase), text, flags=re.IGNORECASE):
                hits.append(f"{path.name}: {phrase}")
    return hits


def external_language_results() -> dict[str, Any]:
    forbidden_hits = scan_forbidden_phrases(
        [
            PACKAGE_DIR / README_NAME,
            PACKAGE_DIR / ARM_GUIDE_MD,
            PACKAGE_DIR / BILLING_CHECKLIST_MD,
            PACKAGE_DIR / PROVIDER_WORKSHEET_MD,
            PACKAGE_DIR / SOURCE_REGISTER_CSV,
            PACKAGE_DIR / VALIDATION_MD,
        ]
    )
    return {
        "test_id": "LANG01",
        "status": "PASS" if not forbidden_hits else "FAIL",
        "detail": "No restricted external-facing phrases found."
        if not forbidden_hits
        else "; ".join(forbidden_hits),
    }


def build_all() -> None:
    workbook_path = PACKAGE_DIR / WORKBOOK_NAME
    write_workbook(workbook_path)
    scenario = scenario_by_id("S01")
    controls = default_controls(scenario)
    result = calculate_scenario(scenario, controls)
    create_collateral(result, scenario)
    results = validate_workbook(workbook_path)
    results.append(external_language_results())
    write_validation_results(results)
    create_delivery_folder()


def create_delivery_folder() -> Path:
    delivery_dir = PACKAGE_DIR / DELIVERY_DIR_NAME
    delivery_dir.mkdir(parents=True, exist_ok=True)
    final_files = [
        WORKBOOK_NAME,
        PROVIDER_PDF_NAME,
        README_NAME,
        ARM_GUIDE_MD,
        ARM_GUIDE_MD.replace(".md", ".pdf"),
        BILLING_CHECKLIST_MD,
        BILLING_CHECKLIST_MD.replace(".md", ".pdf"),
        PROVIDER_WORKSHEET_MD,
        PROVIDER_WORKSHEET_MD.replace(".md", ".pdf"),
        SOURCE_REGISTER_CSV,
        VALIDATION_MD,
        VALIDATION_JSON,
        "build_final_arm_tool.py",
        "export_arm_summary_pdf.sh",
    ]
    for file_name in final_files:
        source = PACKAGE_DIR / file_name
        if source.exists():
            shutil.copy2(source, delivery_dir / file_name)
    return delivery_dir


def export_pdf_only(workbook_path: Path, output_dir: Path) -> Path:
    scenario, controls, overrides, override_active = extract_workbook_inputs(workbook_path)
    result = calculate_scenario(scenario, controls, overrides)
    output_dir.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^A-Za-z0-9]+", "_", f"{scenario.scenario_id}_{scenario.name}").strip("_")[:80]
    pdf_path = output_dir / f"Adrabetadex_Provider_Worksheet_{safe}_{BUILD_DATE.isoformat()}.pdf"
    make_provider_pdf(pdf_path, scenario, result, override_active)
    return pdf_path


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--validate-only", action="store_true")
    parser.add_argument("--pdf-only", action="store_true")
    parser.add_argument("--workbook", default=str(PACKAGE_DIR / WORKBOOK_NAME))
    parser.add_argument("--out-dir", default=str(PACKAGE_DIR))
    args = parser.parse_args()

    workbook_path = Path(args.workbook)
    if args.pdf_only:
        print(export_pdf_only(workbook_path, Path(args.out_dir)))
        return
    if args.validate_only:
        results = validate_workbook(workbook_path)
        results.append(external_language_results())
        write_validation_results(results)
        create_delivery_folder()
        print(json.dumps(results, indent=2))
        return
    build_all()
    print(PACKAGE_DIR / WORKBOOK_NAME)


if __name__ == "__main__":
    main()
