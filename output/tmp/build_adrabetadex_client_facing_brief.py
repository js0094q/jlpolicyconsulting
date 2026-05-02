from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path("/Users/josephstewart/Documents/JLPolicyConsulting")
OUT_DIR = ROOT / "output" / "exports" / "adrabetadex_mlr_concept_brief"
TMP_DIR = ROOT / "output" / "tmp" / "adrabetadex_client_facing_brief"
TEMPLATE_DOCX = ROOT / "output" / "doc" / "Adrabetadex_Provider_Net_Cost_Calculator_Narrative_Brief_Navisync_Letterhead.docx"
DOCX_OUT = OUT_DIR / "Adrabetadex_Provider_Net_Cost_Calculator_Client_Facing_Brief.docx"
PDF_OUT = OUT_DIR / "Adrabetadex_Provider_Net_Cost_Calculator_Client_Facing_Brief.pdf"
BACKGROUND_OUT = TMP_DIR / "navisync_letterhead_background.jpg"
LOGO_OUT = TMP_DIR / "navisync_logo.png"


NAVY = "173B5F"
GREEN = "2F7D32"
TEAL = "0B7D6D"
LIGHT_GREEN = "EAF3E5"
LIGHT_TEAL = "EAF5F2"
LIGHT_BLUE = "EEF4F8"
GRID = "B8C9D4"
TEXT = "1C2D3A"
MUTED = "536678"


TITLE = "Adrabetadex Provider Net Cost Calculator Concept Brief"
SUBTITLE = "Prepared for the Beren Therapeutics Launch Team"
AUDIENCE = "Audience: Internal medical, commercial, market access, reimbursement, and finance stakeholders"
DATE_LINE = "Date: April 2026"
FOOTER_LEFT = "An NPG Health Company"


SECTION_1_TABLE = [
    ["Provider Question", "Why It Matters"],
    ["Who buys the drug?", "Determines whether the provider carries inventory and cash-flow risk."],
    ["Who bills for the drug?", "Determines whether the provider can recover drug cost or only bill for administration."],
    ["What does the provider pay?", "Acquisition cost may differ from WAC depending on 340B status, discounting, or channel."],
    ["What does the payer reimburse?", "Expected payment under Medicare, Medicaid, commercial, or contracted scenarios."],
    ["How long does payment take?", "Estimated days to payment and working-capital exposure."],
    ["What if the claim is denied?", "Denial exposure, appeal recovery, and rework burden."],
    ["What closes the gap?", "The support needed to make a thin or negative scenario financially workable."],
]


CORE_THESIS_TABLE = [
    ["Model Output", "Meaning for a Non-Technical Audience"],
    ["Expected reimbursement", "What the provider expects to be paid."],
    ["Inventory at risk", "How much drug cost the provider is carrying before payment."],
    ["Time to cash", "How long the provider waits to be paid."],
    ["Denial exposure", "How much money is at risk if the claim is denied or delayed."],
    ["Administration economics", "Whether the provider is paid enough for the work of giving the drug."],
    ["Break-even support", "The support needed so the provider does not lose money."],
]


LAUNCH_FRICTION_TABLE = [
    ["Launch Friction", "Simple Explanation", "Why It Matters"],
    ["Buy-and-bill", "The provider buys the drug, gives it, bills the payer, and waits for payment.", "Creates inventory and cash-flow risk."],
    ["White-bagging", "A specialty pharmacy sends the drug to the provider for administration.", "Reduces provider inventory risk but removes drug margin."],
    ["340B acquisition", "Eligible hospitals may acquire the drug at a lower price.", "Can make hospital economics stronger than non-340B settings."],
    ["Medicaid claim integrity", "Medicaid may require correct NDC, coding, and state-specific billing details.", "Coverage can exist while payment still fails."],
    ["HCPCS timing", "Product-specific billing codes may not be available at launch.", "Interim coding can create claim uncertainty."],
    ["Hospital contracting", "Hospitals may need contracting clarity before they stock or administer the drug.", "Internal hospital economics can shape adoption."],
    ["PBM / specialty pharmacy rules", "Payers may require use of a specific specialty pharmacy.", "Determines whether the provider can buy-and-bill."],
    ["Manufacturer support", "Support may be needed to keep the provider whole.", "Shows the minimum support needed to close the provider gap."],
]


FLOW_TABLE = [
    ["Buy-and-Bill Pathway", "White-Bagging Pathway"],
    [
        "Manufacturer\n-> Hospital / COE / Provider\n-> Provider acquires drug\n-> Provider administers drug\n-> Provider bills payer\n-> Provider waits for payment\n-> Provider net recovery or shortfall",
        "Manufacturer\n-> Specialty pharmacy\n-> Specialty pharmacy supplies drug\n-> Provider administers only\n-> Provider bills administration\n-> Provider administration economics only",
    ],
]


STEP_TABLE = [
    ["Step", "What the Calculator Does", "Plain-English Output"],
    ["1. Starts with gross drug exposure", "Uses WAC and dosing frequency.", "How large the drug cost is."],
    ["2. Applies acquisition basis", "Uses WAC, 340B, discount, rebate, or white-bagging.", "What the provider actually has at risk."],
    ["3. Applies reimbursement basis", "Uses payer, Medicare, Medicaid, contract, or manual payment input.", "What the provider expects to be paid."],
    ["4. Adds administration costs", "Adds labor, pharmacy handling, billing, overhead, and related costs.", "What it costs to provide the service."],
    ["5. Adds claim friction", "Adds denial risk, appeal recovery, rework, and days to payment.", "How payment delays or denials affect the provider."],
    ["6. Tests support levels", "Adds provider-facing support where needed.", "The support needed to avoid a provider loss."],
    ["7. Produces a result", "Calculates net recovery or shortfall.", "Workable, thin, negative, or site-specific."],
]


OUTPUT_TABLE = [
    ["Question", "Calculator Output", "Why the Audience Should Care"],
    ["Can the provider afford to start treatment?", "Inventory exposure and expected reimbursement.", "Shows whether a site may hesitate to stock or administer."],
    ["Does reimbursement cover drug cost?", "Drug spread before operating cost.", "Shows whether payment covers acquisition."],
    ["Does the provider remain whole after all costs?", "Net recovery after labor, handling, denial, and timing costs.", "Shows the true provider result."],
    ["Which sites are most viable?", "COE, HOPD, physician office, Medicaid, and white-bagging comparison.", "Helps prioritize launch sites."],
    ["Which payer rules create risk?", "Prior authorization, specialist, coding, sourcing, and payment flags.", "Shows where payer friction affects uptake."],
    ["What support closes the gap?", "Break-even manufacturer support.", "Shows what action is needed to make a scenario workable."],
    ["How does a comparator perform?", "Same economics applied to a relevant therapy.", "Shows whether Adrabetadex is better, worse, or similar under real reimbursement conditions."],
]


SOURCE_TABLE = [
    ["Source Category", "Purpose"],
    ["Beren internal Adrabetadex assumptions", "Supports WAC, dosing, and working model assumptions."],
    ["Aetna CPB 0442", "Supports payer-rule example for rare-disease CNS-administered therapy."],
    ["CMS Part B payment limit / initial-period WAC rule", "Supports initial reimbursement benchmark concept."],
    ["CMS HCPCS Level II process", "Supports product-specific coding timing."],
    ["HRSA 340B ceiling price / new drug estimation", "Supports 340B acquisition scenario logic."],
    ["Medicaid PAD / EPSDT resources", "Supports Medicaid claim-integrity and pediatric coverage framing."],
]


def hex_to_rgb(hex_value: str) -> RGBColor:
    return RGBColor(int(hex_value[0:2], 16), int(hex_value[2:4], 16), int(hex_value[4:6], 16))


def set_run_font(run, size=10.5, bold=False, color=TEXT, name="Arial"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = hex_to_rgb(color)


def clear_body(document: Document) -> None:
    body = document._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def clear_container(container) -> None:
    element = container._element
    for child in list(element):
        element.remove(child)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = GRID, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_grid(table, widths: list[float]) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_grid = tbl.find(qn("w:tblGrid"))
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        tbl_grid.append(grid_col)


def add_paragraph(doc: Document, text: str = "", size=10.5, color=TEXT, bold=False, space_after=5, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.08
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, color=color, bold=bold)
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=15 if level == 1 else 12, bold=False, color=TEAL if level == 1 else GREEN)
    return p


def add_callout(doc: Document, lines: list[str], fill=LIGHT_TEAL, border=TEAL, font_size=10.2):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid(table, [6.55])
    cell = table.cell(0, 0)
    cell.width = Inches(6.55)
    set_cell_shading(cell, fill)
    set_cell_border(cell, border, "8")
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
        run = p.add_run(line)
        set_run_font(run, size=font_size, color=TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_docx_table(doc: Document, data: list[list[str]], widths: list[float], font_size=8.5, header_fill=TEAL):
    table = doc.add_table(rows=len(data), cols=len(data[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid(table, widths)
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[col_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_idx == 0:
                set_cell_shading(cell, header_fill)
            elif row_idx % 2 == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            else:
                set_cell_shading(cell, "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            for existing in list(p.runs):
                existing._element.getparent().remove(existing._element)
            run = p.add_run(data[row_idx][col_idx])
            set_run_font(run, size=font_size, bold=(row_idx == 0), color="FFFFFF" if row_idx == 0 else TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    run._r.append(fld_begin)

    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run._r.append(instr)

    run = paragraph.add_run()
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    run._r.append(fld_sep)

    run = paragraph.add_run()
    text = OxmlElement("w:t")
    text.text = "1"
    run._r.append(text)
    set_run_font(run, size=8.5, color=MUTED)

    run = paragraph.add_run()
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_end)


def configure_docx_header(doc: Document) -> None:
    section = doc.sections[0]
    header = section.header
    clear_container(header)
    p = header.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("NAVISYNC")
    set_run_font(run, size=18, color=GREEN, bold=True)
    tagline = p.add_run("  A Managed Markets Agency | Payers. Providers. Patients.")
    set_run_font(tagline, size=7.5, color=MUTED)
    p2 = header.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    p_pr = p2._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), TEAL)
    border.append(bottom)
    p_pr.append(border)


def configure_docx_footer(doc: Document) -> None:
    section = doc.sections[0]
    footer = section.footer
    clear_container(footer)
    left_p = footer.add_paragraph()
    left_p.paragraph_format.space_after = Pt(0)
    left_run = left_p.add_run(FOOTER_LEFT)
    set_run_font(left_run, size=8, color=MUTED)
    right_p = footer.add_paragraph()
    right_p.paragraph_format.space_after = Pt(0)
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    page_run = right_p.add_run("Page ")
    set_run_font(page_run, size=8.5, color=MUTED)
    add_page_field(right_p)


def build_docx() -> None:
    extract_background()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.88)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.22)
    section.footer_distance = Inches(0.22)
    configure_docx_header(doc)
    configure_docx_footer(doc)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].font.color.rgb = hex_to_rgb(TEXT)

    add_paragraph(doc, "April 2026", size=10, color=MUTED, space_after=6)
    title_p = add_paragraph(doc, TITLE, size=23, color=GREEN, space_after=4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p = add_paragraph(doc, SUBTITLE, size=11, color=TEXT, bold=True, space_after=2)
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    audience_p = add_paragraph(doc, AUDIENCE, size=9.5, color=MUTED, space_after=1)
    audience_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p = add_paragraph(doc, DATE_LINE, size=9.5, color=MUTED, space_after=10)
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading(doc, "1. Opening Position")
    add_paragraph(doc, "The Adrabetadex Provider Net Cost Calculator helps Beren evaluate whether treatment centers can administer Adrabetadex under real-world reimbursement conditions.")
    add_paragraph(doc, "Adrabetadex may have a compelling clinical profile, but provider adoption will still depend on whether centers can acquire the drug, administer it, bill correctly, receive payment, and remain financially whole.")
    add_paragraph(doc, "The calculator translates payer rules, site-of-care requirements, channel decisions, coding status, and manufacturer support into a clear provider-level financial result.")
    add_docx_table(doc, SECTION_1_TABLE, [2.25, 4.55], font_size=8.2)

    add_heading(doc, "Aetna Example for Non-Technical Readers", level=2)
    add_callout(
        doc,
        [
            "Aetna's policy for Brineura, a rare-disease therapy administered into the central nervous system, shows why a calculator like this is necessary. The policy does not simply ask whether the product is clinically appropriate. It requires precertification, diagnosis by or in consultation with a specialist in metabolic disease or lysosomal storage disorders, confirmed diagnosis by enzyme assay or genetic testing, and administration by or under a physician knowledgeable in intraventricular administration. It also ties reimbursement to covered coding, including applicable HCPCS and procedure codes where criteria are met.",
            "The practical point is simple: for provider-administered rare-disease therapies, payer rules, specialist requirements, coding, and site capability affect whether the claim can be paid. The calculator must capture those mechanics, not just product price.",
            "Source note: Aetna Clinical Policy Bulletin 0442, Lysosomal Storage Disorder Treatments.",
        ],
        fill=LIGHT_TEAL,
        border=TEAL,
        font_size=9.4,
    )

    add_heading(doc, "2. Core Thesis")
    add_paragraph(doc, "The calculator converts payer rules, site-of-care dynamics, and channel mechanics into treatment-center economics.")
    add_callout(
        doc,
        [
            "It is not a list-price model. It is not a manufacturer gross-to-net model. It is not a clinical value model.",
            "In plain English, WAC helps size the starting exposure. It does not show what the provider pays, gets reimbursed, or carries while waiting for payment.",
        ],
        fill=LIGHT_GREEN,
        border=GREEN,
        font_size=10,
    )
    add_paragraph(doc, "It is a provider-facing reimbursement execution model that estimates:")
    add_docx_table(doc, CORE_THESIS_TABLE, [2.15, 4.65], font_size=8.4)

    add_heading(doc, "3. Why the Calculator Is Needed")
    add_paragraph(doc, "A provider can believe in the clinical value of a therapy and still hesitate to administer it if the financial pathway is unclear.")
    add_paragraph(doc, "For Adrabetadex, the calculator is needed because several reimbursement issues can change whether the provider is financially whole.")
    add_docx_table(doc, LAUNCH_FRICTION_TABLE, [1.65, 2.7, 2.45], font_size=7.8)

    add_heading(doc, "4. How the Calculator Works")
    add_paragraph(doc, "The calculator uses a simple product and payment flow.")
    add_docx_table(doc, FLOW_TABLE, [3.35, 3.35], font_size=8.6, header_fill=GREEN)
    add_paragraph(doc, "The calculator applies one formula across all scenarios:")
    add_callout(
        doc,
        [
            "Provider net position = drug reimbursement + administration revenue + provider-facing support - drug acquisition cost - labor and pharmacy handling - billing and denial rework - financing cost from delayed payment"
        ],
        fill=LIGHT_GREEN,
        border=GREEN,
        font_size=9.7,
    )
    add_docx_table(doc, STEP_TABLE, [1.75, 2.6, 2.45], font_size=7.8)

    add_heading(doc, "5. What the Calculator Will Show")
    add_paragraph(doc, "The calculator should make the economics visible to people who do not normally think in reimbursement terms.")
    add_docx_table(doc, OUTPUT_TABLE, [2.05, 2.35, 2.4], font_size=7.7)

    add_heading(doc, "Client-Facing Close")
    add_callout(
        doc,
        [
            "The purpose of the calculator is to make provider economics understandable.",
            "Adrabetadex may be clinically compelling, but launch success still depends on whether treatment centers can buy the drug, administer it, bill correctly, get paid, and avoid carrying unacceptable financial risk.",
            "The calculator shows where the economics work, where they are thin, where they fail, and what action is needed to improve the provider pathway.",
        ],
        fill=LIGHT_TEAL,
        border=TEAL,
        font_size=9.8,
    )

    add_heading(doc, "Source Notes")
    add_paragraph(doc, "Source notes are included to support the model framing without overloading the brief with technical citations.")
    add_docx_table(doc, SOURCE_TABLE, [2.65, 4.15], font_size=8.2)

    doc.save(str(DOCX_OUT))


def extract_background() -> None:
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    with ZipFile(TEMPLATE_DOCX) as zf:
        BACKGROUND_OUT.write_bytes(zf.read("word/media/image1.jpg"))
    image = Image.open(BACKGROUND_OUT)
    logo_crop = image.crop((115, 115, 1030, 315))
    logo_crop.save(LOGO_OUT)


def pdf_styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("title", parent=base["Title"], fontName="Helvetica", fontSize=22, leading=25, textColor=colors.HexColor("#2F7D32"), alignment=TA_CENTER, spaceAfter=6),
        "subtitle": ParagraphStyle("subtitle", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, textColor=colors.HexColor("#1C2D3A"), alignment=TA_CENTER, spaceAfter=2),
        "meta": ParagraphStyle("meta", parent=base["Normal"], fontName="Helvetica", fontSize=9.2, leading=12, textColor=colors.HexColor("#536678"), alignment=TA_CENTER, spaceAfter=2),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontName="Helvetica", fontSize=15, leading=18, textColor=colors.HexColor("#0B7D6D"), spaceBefore=8, spaceAfter=5, keepWithNext=True),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontName="Helvetica", fontSize=12, leading=15, textColor=colors.HexColor("#2F7D32"), spaceBefore=7, spaceAfter=4, keepWithNext=True),
        "body": ParagraphStyle("body", parent=base["Normal"], fontName="Helvetica", fontSize=9.4, leading=12.4, textColor=colors.HexColor("#1C2D3A"), spaceAfter=5),
        "small": ParagraphStyle("small", parent=base["Normal"], fontName="Helvetica", fontSize=8.0, leading=10, textColor=colors.HexColor("#1C2D3A"), spaceAfter=3),
        "callout": ParagraphStyle("callout", parent=base["Normal"], fontName="Helvetica", fontSize=9.2, leading=12, textColor=colors.HexColor("#1C2D3A"), spaceAfter=5, leftIndent=0),
        "table": ParagraphStyle("table", parent=base["Normal"], fontName="Helvetica", fontSize=7.7, leading=9.3, textColor=colors.HexColor("#1C2D3A")),
        "table_header": ParagraphStyle("table_header", parent=base["Normal"], fontName="Helvetica-Bold", fontSize=7.8, leading=9.3, textColor=colors.white),
    }


def p(text: str, style: ParagraphStyle) -> Paragraph:
    escaped = text.replace("&", "&amp;")
    escaped = escaped.replace("<", "&lt;").replace(">", "&gt;")
    return Paragraph(escaped.replace("\n", "<br/>"), style)


def pdf_table(data, widths, styles, header_fill=TEAL, font_size=None):
    rows = []
    for idx, row in enumerate(data):
        row_style = styles["table_header"] if idx == 0 else styles["table"]
        if font_size and idx > 0:
            row_style = ParagraphStyle(f"table_{font_size}_{idx}", parent=row_style, fontSize=font_size, leading=font_size + 1.7)
        rows.append([p(cell, row_style) for cell in row])
    table = Table(rows, colWidths=[w * inch for w in widths], repeatRows=1, hAlign="CENTER")
    table_style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + header_fill)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B8C9D4")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]
    for row_idx in range(1, len(data)):
        if row_idx % 2 == 0:
            table_style.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.HexColor("#EEF4F8")))
        else:
            table_style.append(("BACKGROUND", (0, row_idx), (-1, row_idx), colors.white))
    table.setStyle(TableStyle(table_style))
    return table


def pdf_callout(lines, styles, fill=LIGHT_TEAL, border=TEAL, font_size=9.2):
    style = ParagraphStyle(f"callout_{font_size}_{fill}", parent=styles["callout"], fontSize=font_size, leading=font_size + 2.8)
    content = [[p("\n\n".join(lines), style)]]
    table = Table(content, colWidths=[6.45 * inch], hAlign="CENTER")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#" + fill)),
        ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#" + border)),
        ("LEFTPADDING", (0, 0), (-1, -1), 16),
        ("RIGHTPADDING", (0, 0), (-1, -1), 16),
        ("TOPPADDING", (0, 0), (-1, -1), 10),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 10),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    return table


def on_page(canvas, doc):
    width, height = letter
    if BACKGROUND_OUT.exists():
        canvas.drawImage(str(BACKGROUND_OUT), 0, 0, width=width, height=height, preserveAspectRatio=False, mask="auto")
    canvas.saveState()
    canvas.setFont("Helvetica", 7.5)
    canvas.setFillColor(colors.HexColor("#536678"))
    canvas.drawString(0.52 * inch, 0.40 * inch, FOOTER_LEFT)
    canvas.drawRightString(width - 0.52 * inch, 0.40 * inch, f"Page {doc.page}")
    canvas.restoreState()


def build_pdf() -> None:
    extract_background()
    styles = pdf_styles()
    doc = BaseDocTemplate(
        str(PDF_OUT),
        pagesize=letter,
        leftMargin=0.68 * inch,
        rightMargin=0.68 * inch,
        topMargin=1.25 * inch,
        bottomMargin=0.72 * inch,
        title=TITLE,
        author="Navisync",
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates([PageTemplate(id="letterhead", frames=[frame], onPage=on_page)])

    story = [
        p("April 2026", styles["meta"]),
        p(TITLE, styles["title"]),
        p(SUBTITLE, styles["subtitle"]),
        p(AUDIENCE, styles["meta"]),
        p(DATE_LINE, styles["meta"]),
        Spacer(1, 0.08 * inch),
        p("1. Opening Position", styles["h1"]),
        p("The Adrabetadex Provider Net Cost Calculator helps Beren evaluate whether treatment centers can administer Adrabetadex under real-world reimbursement conditions.", styles["body"]),
        p("Adrabetadex may have a compelling clinical profile, but provider adoption will still depend on whether centers can acquire the drug, administer it, bill correctly, receive payment, and remain financially whole.", styles["body"]),
        p("The calculator translates payer rules, site-of-care requirements, channel decisions, coding status, and manufacturer support into a clear provider-level financial result.", styles["body"]),
        Spacer(1, 0.08 * inch),
        pdf_table(SECTION_1_TABLE, [2.18, 4.28], styles, font_size=7.5),
        p("Aetna Example for Non-Technical Readers", styles["h2"]),
        pdf_callout(
            [
                "Aetna's policy for Brineura, a rare-disease therapy administered into the central nervous system, shows why a calculator like this is necessary. The policy does not simply ask whether the product is clinically appropriate. It requires precertification, diagnosis by or in consultation with a specialist in metabolic disease or lysosomal storage disorders, confirmed diagnosis by enzyme assay or genetic testing, and administration by or under a physician knowledgeable in intraventricular administration. It also ties reimbursement to covered coding, including applicable HCPCS and procedure codes where criteria are met.",
                "The practical point is simple: for provider-administered rare-disease therapies, payer rules, specialist requirements, coding, and site capability affect whether the claim can be paid. The calculator must capture those mechanics, not just product price.",
                "Source note: Aetna Clinical Policy Bulletin 0442, Lysosomal Storage Disorder Treatments.",
            ],
            styles,
            font_size=8.25,
        ),
        PageBreak(),
        p("2. Core Thesis", styles["h1"]),
        p("The calculator converts payer rules, site-of-care dynamics, and channel mechanics into treatment-center economics.", styles["body"]),
        pdf_callout(
            [
                "It is not a list-price model. It is not a manufacturer gross-to-net model. It is not a clinical value model.",
                "In plain English, WAC helps size the starting exposure. It does not show what the provider pays, gets reimbursed, or carries while waiting for payment.",
            ],
            styles,
            fill=LIGHT_GREEN,
            border=GREEN,
            font_size=9,
        ),
        p("It is a provider-facing reimbursement execution model that estimates:", styles["body"]),
        pdf_table(CORE_THESIS_TABLE, [2.1, 4.36], styles, font_size=7.7),
        p("3. Why the Calculator Is Needed", styles["h1"]),
        p("A provider can believe in the clinical value of a therapy and still hesitate to administer it if the financial pathway is unclear.", styles["body"]),
        p("For Adrabetadex, the calculator is needed because several reimbursement issues can change whether the provider is financially whole.", styles["body"]),
        pdf_table(LAUNCH_FRICTION_TABLE, [1.55, 2.52, 2.39], styles, font_size=7.05),
        PageBreak(),
        p("4. How the Calculator Works", styles["h1"]),
        p("The calculator uses a simple product and payment flow.", styles["body"]),
        pdf_table(FLOW_TABLE, [3.18, 3.18], styles, header_fill=GREEN, font_size=8.15),
        p("The calculator applies one formula across all scenarios:", styles["body"]),
        pdf_callout(
            [
                "Provider net position = drug reimbursement + administration revenue + provider-facing support - drug acquisition cost - labor and pharmacy handling - billing and denial rework - financing cost from delayed payment"
            ],
            styles,
            fill=LIGHT_GREEN,
            border=GREEN,
            font_size=8.65,
        ),
        pdf_table(STEP_TABLE, [1.65, 2.45, 2.36], styles, font_size=7.0),
        PageBreak(),
        p("5. What the Calculator Will Show", styles["h1"]),
        p("The calculator should make the economics visible to people who do not normally think in reimbursement terms.", styles["body"]),
        pdf_table(OUTPUT_TABLE, [1.94, 2.23, 2.29], styles, font_size=7.0),
        p("Client-Facing Close", styles["h1"]),
        pdf_callout(
            [
                "The purpose of the calculator is to make provider economics understandable.",
                "Adrabetadex may be clinically compelling, but launch success still depends on whether treatment centers can buy the drug, administer it, bill correctly, get paid, and avoid carrying unacceptable financial risk.",
                "The calculator shows where the economics work, where they are thin, where they fail, and what action is needed to improve the provider pathway.",
            ],
            styles,
            font_size=9.0,
        ),
        p("Source Notes", styles["h1"]),
        p("Source notes are included to support the model framing without overloading the brief with technical citations.", styles["body"]),
        pdf_table(SOURCE_TABLE, [2.5, 3.96], styles, font_size=7.1),
    ]
    doc.build(story)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    build_docx()
    build_pdf()
    print(DOCX_OUT)
    print(PDF_OUT)


if __name__ == "__main__":
    main()
