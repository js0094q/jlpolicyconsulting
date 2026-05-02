from __future__ import annotations

import json
import math
import os
import shutil
import sys
from pathlib import Path

PPTX_VENDOR = Path("/tmp/adrabetadex_pptx_lib")
if PPTX_VENDOR.exists():
    sys.path.insert(0, str(PPTX_VENDOR))

from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
DELIVERY = ROOT / "output" / "exports" / "adrabetadex_provider_economics_executive_brief_2026-04-28"
PDF_DIR = ROOT / "output" / "pdf"
TMP_DIR = ROOT / "tmp" / "pdfs" / "adrabetadex_provider_economics_brief"
ASSET_DIR = TMP_DIR / "assets"
PREVIEW_DIR = DELIVERY / "previews"
LETTERHEAD = ROOT / "tmp" / "pdfs" / "adrabetadex_assets" / "img-000.png"

PDF_NAME = "Adrabetadex_Provider_Economics_and_Adoption_Risk_Executive_Brief.pdf"
PPTX_NAME = "Adrabetadex_Provider_Economics_and_Adoption_Risk_Executive_Brief.pptx"
REPORT_NAME = "build_manifest.json"

W, H = letter

NAVY = colors.HexColor("#1F2E2B")
GREEN = colors.HexColor("#4E8A2D")
LIGHT_GREEN = colors.HexColor("#EEF5EA")
MID_GREEN = colors.HexColor("#C7DDBB")
PALE = colors.HexColor("#F7FAF6")
GRAY = colors.HexColor("#4D5656")
LIGHT_GRAY = colors.HexColor("#D8E2D4")
AMBER = colors.HexColor("#B27A1B")
RED = colors.HexColor("#8C3D2B")
WHITE = colors.white
BLACK = colors.black

PPT_NAVY = RGBColor(31, 46, 43)
PPT_GREEN = RGBColor(78, 138, 45)
PPT_LIGHT_GREEN = RGBColor(238, 245, 234)
PPT_MID_GREEN = RGBColor(199, 221, 187)
PPT_PALE = RGBColor(247, 250, 246)
PPT_GRAY = RGBColor(77, 86, 86)
PPT_LIGHT_GRAY = RGBColor(216, 226, 212)
PPT_AMBER = RGBColor(178, 122, 27)
PPT_RED = RGBColor(140, 61, 43)
PPT_WHITE = RGBColor(255, 255, 255)


CONTENT = {
    "title": "Adrabetadex Provider Economics and Adoption Risk",
    "subtitle": "Why Net Cost Recovery Determines Use",
    "prepared": "Prepared for Beren Therapeutics | Navisync Letterhead Format | April 2026",
    "source_note": (
        "Prepared from supplied planning assumptions and the existing Adrabetadex concept brief. "
        "Illustrative economics are for management planning, not a reimbursement guarantee."
    ),
    "thesis": "Providers adopt when treatment is financially viable, not only when clinically convinced.",
    "purpose": (
        "The calculator converts coding, payment, payer controls, acquisition cost, and claim timing into a "
        "provider-level go / no-go adoption signal."
    ),
    "core_rows": [
        ("Acquisition cost", "Provider buys high-cost drug before payment", "Working-capital exposure"),
        ("Coding pathway", "No product-specific code at launch", "Manual claims and denial risk"),
        ("Payment timing", "30-90+ day lag", "Cash tied up across treatment cycles"),
        ("Denials and rework", "Prior authorization and documentation burden", "Admin cost and revenue leakage"),
        ("Site controls", "White-bagging or site-of-care limits", "Drug margin or ability to treat may disappear"),
    ],
    "coding_steps": [
        ("EAP", "No commercial claim"),
        ("NOC code", "Manual billing"),
        ("J-code", "Standardized billing"),
        ("Stable billing", "Routine pathway"),
    ],
    "pass_rows": [
        ("No pass-through", "Drug may be bundled or manually priced", "High margin uncertainty"),
        ("Pass-through active", "Temporary separate OPPS payment", "Improves early payment visibility"),
        ("After expiration", "Economics rely on coding and contracting", "Long-term viability must be proven"),
    ],
    "eap_vs_commercial": [
        ("Drug cost", "Often free or subsidized", "Provider purchases at WAC or 340B"),
        ("Billing work", "Minimal", "Full coding and claims burden"),
        ("Financial risk", "Low", "High until paid"),
        ("Adoption read", "Can overstate readiness", "Shows real launch viability"),
    ],
    "payer_friction": [
        ("Prior auth", "Approval before treatment. Delays care and weakens claim certainty."),
        ("White-bagging", "Payer pharmacy ships drug. Provider loses drug spread and control."),
        ("Site-of-care", "Payer steers setting. Hospital economics may not hold."),
    ],
    "impact_340b": [
        ("Non-340B", "~$39,000", "ASP/WAC-based", "Thin or negative after delays and denials"),
        ("340B COE", "~$30,000 planning assumption", "Similar benchmark", "Stronger per-dose margin runway"),
        ("White-bagged", "No provider purchase", "Specialty pharmacy bills drug", "No drug margin; admin economics matter"),
    ],
    "definitions": [
        ("HCPCS", "Billing code set used for drugs and services; a J-code makes product billing more routine."),
        ("Pass-through", "Temporary OPPS separate payment for certain new outpatient drugs and biologicals."),
        ("Buy-and-bill", "Provider buys the drug, administers it, then bills the payer."),
        ("340B", "Federal program allowing eligible hospitals to acquire outpatient drugs at discounted prices."),
    ],
    "inputs": [
        "Coding status",
        "Pass-through status",
        "Site of care",
        "Payer mix",
        "340B eligibility",
        "Buy-and-bill vs white-bagging",
        "Denials and payment timing",
    ],
    "outputs": ["Per-dose margin", "Annual exposure", "Cash flow burden", "Break-even support", "Adoption viability"],
}


def ensure_dirs() -> None:
    DELIVERY.mkdir(parents=True, exist_ok=True)
    PDF_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    PREVIEW_DIR.mkdir(parents=True, exist_ok=True)


def prepare_assets() -> dict[str, Path]:
    logo_path = ASSET_DIR / "navisync_logo_crop.png"
    if LETTERHEAD.exists():
        img = Image.open(LETTERHEAD).convert("RGBA")
        img.crop((95, 120, 1075, 375)).save(logo_path)
    else:
        logo = Image.new("RGBA", (980, 255), (255, 255, 255, 0))
        draw = ImageDraw.Draw(logo)
        draw.text((20, 70), "NAVISYNC", fill=(47, 98, 44), font=load_font(86, bold=True))
        draw.text((25, 160), "A Managed Markets Agency | Payers. Providers. Patients.", fill=(90, 90, 90), font=load_font(28))
        logo.save(logo_path)
    return {"logo": logo_path}


def draw_letterhead(c: canvas.Canvas, page_num: int) -> None:
    if LETTERHEAD.exists():
        c.drawImage(str(LETTERHEAD), 0, 0, width=W, height=H, mask="auto")
    else:
        c.setFillColor(colors.white)
        c.rect(0, 0, W, H, stroke=0, fill=1)
        c.setFillColor(GREEN)
        c.setFont("Helvetica-Bold", 24)
        c.drawString(42, H - 58, "NAVISYNC")
        c.setFont("Helvetica", 8)
        c.drawString(42, 29, "An NPG Health Company")
        c.setFillColor(GREEN)
        c.rect(0, 0, W, 8, stroke=0, fill=1)
    c.setFillColor(GRAY)
    c.setFont("Helvetica", 8)
    c.drawRightString(W - 39, 30, f"Page {page_num}")


def style(font_size: float = 8.5, leading: float | None = None, color=GRAY, bold: bool = False) -> ParagraphStyle:
    return ParagraphStyle(
        name=f"s-{font_size}-{bold}",
        fontName="Helvetica-Bold" if bold else "Helvetica",
        fontSize=font_size,
        leading=leading or font_size + 2.1,
        textColor=color,
        spaceAfter=0,
        spaceBefore=0,
    )


def para(c: canvas.Canvas, text: str, x: float, y_top: float, width: float, height: float, size: float = 8.5, color=GRAY, bold: bool = False, align: int = 0) -> float:
    st = style(size, color=color, bold=bold)
    st.alignment = align
    p = Paragraph(escape(text), st)
    w, h = p.wrap(width, height)
    p.drawOn(c, x, y_top - h)
    return h


def escape(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\n", "<br/>")
    )


def draw_section_title(c: canvas.Canvas, label: str, x: float, y: float, width: float) -> None:
    c.setStrokeColor(MID_GREEN)
    c.setLineWidth(0.8)
    c.line(x, y - 5, x + width, y - 5)
    c.setFillColor(GREEN)
    c.setFont("Helvetica-Bold", 11)
    c.drawString(x, y, label)


def draw_table(c: canvas.Canvas, x: float, y_top: float, col_widths: list[float], headers: list[str], rows: list[tuple[str, ...]], row_h: float, font_size: float = 7.4) -> float:
    total_w = sum(col_widths)
    header_h = row_h
    c.setFillColor(LIGHT_GREEN)
    c.rect(x, y_top - header_h, total_w, header_h, stroke=0, fill=1)
    c.setStrokeColor(MID_GREEN)
    c.setLineWidth(0.6)
    c.rect(x, y_top - header_h, total_w, header_h, stroke=1, fill=0)
    xpos = x
    for idx, header in enumerate(headers):
        para(c, header, xpos + 5, y_top - 5, col_widths[idx] - 10, header_h - 4, size=font_size, color=GREEN, bold=True)
        xpos += col_widths[idx]
        c.line(xpos, y_top, xpos, y_top - header_h - row_h * len(rows))
    current_y = y_top - header_h
    for r_idx, row_values in enumerate(rows):
        fill = PALE if r_idx % 2 == 0 else WHITE
        c.setFillColor(fill)
        c.rect(x, current_y - row_h, total_w, row_h, stroke=0, fill=1)
        c.setStrokeColor(LIGHT_GRAY)
        c.rect(x, current_y - row_h, total_w, row_h, stroke=1, fill=0)
        xpos = x
        for col_idx, value in enumerate(row_values):
            is_first = col_idx == 0
            para(
                c,
                value,
                xpos + 5,
                current_y - 5,
                col_widths[col_idx] - 10,
                row_h - 5,
                size=font_size,
                color=NAVY if is_first else GRAY,
                bold=is_first,
            )
            xpos += col_widths[col_idx]
        current_y -= row_h
    return header_h + row_h * len(rows)


def rounded_label(c: canvas.Canvas, x: float, y: float, w: float, h: float, title: str, body: str, fill=LIGHT_GREEN, stroke=MID_GREEN) -> None:
    c.setFillColor(fill)
    c.setStrokeColor(stroke)
    c.setLineWidth(0.8)
    c.roundRect(x, y, w, h, 7, stroke=1, fill=1)
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 8.2)
    c.drawString(x + 8, y + h - 15, title)
    para(c, body, x + 8, y + h - 21, w - 16, h - 20, size=7.0, color=GRAY)


def arrow(c: canvas.Canvas, x1: float, y1: float, x2: float, y2: float, color=GREEN) -> None:
    c.setStrokeColor(color)
    c.setFillColor(color)
    c.setLineWidth(1.5)
    c.line(x1, y1, x2, y2)
    angle = math.atan2(y2 - y1, x2 - x1)
    size = 5
    points = [
        (x2, y2),
        (x2 - size * math.cos(angle - math.pi / 6), y2 - size * math.sin(angle - math.pi / 6)),
        (x2 - size * math.cos(angle + math.pi / 6), y2 - size * math.sin(angle + math.pi / 6)),
    ]
    c.line(points[0][0], points[0][1], points[1][0], points[1][1])
    c.line(points[0][0], points[0][1], points[2][0], points[2][1])


def draw_pdf_page_one(c: canvas.Canvas) -> None:
    draw_letterhead(c, 1)
    left, right = 58, W - 58
    width = right - left
    y = H - 105
    c.setFillColor(GRAY)
    c.setFont("Helvetica", 9.5)
    c.drawString(left, y, "April 2026")
    y -= 20
    c.setFillColor(GREEN)
    c.setFont("Helvetica-Bold", 19.5)
    c.drawCentredString(W / 2, y, CONTENT["title"])
    y -= 20
    c.setFont("Helvetica", 13)
    c.drawCentredString(W / 2, y, CONTENT["subtitle"])
    y -= 18
    c.setFillColor(GRAY)
    c.setFont("Helvetica", 8.8)
    c.drawCentredString(W / 2, y, CONTENT["prepared"])
    y -= 22

    c.setFillColor(PALE)
    c.setStrokeColor(MID_GREEN)
    c.roundRect(left, y - 70, width, 70, 8, stroke=1, fill=1)
    c.setFillColor(NAVY)
    c.setFont("Helvetica-Bold", 10.5)
    c.drawString(left + 14, y - 17, "Purpose")
    para(c, CONTENT["purpose"], left + 14, y - 23, width - 28, 25, size=7.6, color=GRAY)
    c.setFillColor(GREEN)
    c.setFont("Helvetica-Bold", 9.8)
    c.drawString(left + 14, y - 59, CONTENT["thesis"])
    y -= 88

    draw_section_title(c, "1. The Core Problem: Provider Hesitation", left, y, width)
    y -= 16
    used = draw_table(
        c,
        left,
        y,
        [124, 185, 199],
        ["Issue", "Provider concern", "Impact on adoption"],
        CONTENT["core_rows"],
        row_h=28,
        font_size=7.2,
    )
    y -= used + 22

    draw_section_title(c, "2. Coding Pathway: Payment Becomes Routine Only After the Billing Path Stabilizes", left, y, width)
    y -= 67
    box_w, box_h, gap = 100, 42, 24
    start_x = left + 10
    for idx, (title, body) in enumerate(CONTENT["coding_steps"]):
        bx = start_x + idx * (box_w + gap)
        rounded_label(c, bx, y, box_w, box_h, title, body, fill=LIGHT_GREEN if idx in (0, 1) else PALE)
        if idx < len(CONTENT["coding_steps"]) - 1:
            arrow(c, bx + box_w + 3, y + box_h / 2, bx + box_w + gap - 4, y + box_h / 2)
    c.setFillColor(RED)
    c.setFont("Helvetica-Bold", 7.4)
    c.drawString(start_x + box_w + gap + 5, y - 11, "Launch risk: manual billing + high denial exposure")
    c.setFillColor(GREEN)
    c.drawString(start_x + 3 * (box_w + gap) + 8, y - 11, "Target state: standardized billing")
    y -= 32

    draw_section_title(c, "3. Pass-Through Status: Temporary Stabilizer, Not a Permanent Answer", left, y, width)
    y -= 16
    draw_table(
        c,
        left,
        y,
        [128, 205, 175],
        ["Scenario", "Payment structure", "Provider risk"],
        CONTENT["pass_rows"],
        row_h=29,
        font_size=7.4,
    )
    c.setFillColor(GRAY)
    c.setFont("Helvetica", 6.6)
    c.drawString(left, 44, CONTENT["source_note"])


def draw_pdf_page_two(c: canvas.Canvas) -> None:
    draw_letterhead(c, 2)
    left, right = 58, W - 58
    width = right - left
    y = H - 108

    draw_section_title(c, "4. EAP vs Commercial: Early Experience Can Overstate Launch Readiness", left, y, width)
    y -= 16
    used = draw_table(
        c,
        left,
        y,
        [110, 180, 218],
        ["Dimension", "EAP", "Commercial buy-and-bill"],
        CONTENT["eap_vs_commercial"],
        row_h=22,
        font_size=6.6,
    )
    y -= used + 15

    draw_section_title(c, "5. Payer Friction: Policy Rules Directly Shape Provider Behavior", left, y, width)
    y -= 44
    panel_w = (width - 22) / 3
    for idx, (title, body) in enumerate(CONTENT["payer_friction"]):
        x = left + idx * (panel_w + 11)
        c.setFillColor(PALE)
        c.setStrokeColor(MID_GREEN)
        c.roundRect(x, y, panel_w, 36, 7, stroke=1, fill=1)
        c.setFillColor(GREEN)
        c.circle(x + 15, y + 23, 5, stroke=0, fill=1)
        c.setFillColor(NAVY)
        c.setFont("Helvetica-Bold", 8.2)
        c.drawString(x + 27, y + 27, title)
        para(c, body, x + 27, y + 22, panel_w - 35, 20, size=5.9, color=GRAY)
    y -= 24

    draw_section_title(c, "6. 340B Impact: Acquisition Advantage Can Determine Early Site Activation", left, y, width)
    y -= 16
    used = draw_table(
        c,
        left,
        y,
        [92, 120, 126, 170],
        ["Scenario", "Acquisition", "Reimbursement", "Margin signal"],
        CONTENT["impact_340b"],
        row_h=23,
        font_size=5.9,
    )
    y -= used + 20

    draw_section_title(c, "7. Calculator Value: Inputs Become a Provider Go / No-Go Signal", left, y, width)
    y -= 61
    col_w = (width - 36) / 3
    rounded_label(c, left, y, col_w, 45, "Inputs", "Code, pass-through, site, payer mix, 340B, channel, denials, timing.", fill=LIGHT_GREEN)
    rounded_label(c, left + col_w + 18, y, col_w, 45, "Model", "Expected net provider position after payment, cost, denial, rework, and carrying cost.", fill=PALE)
    rounded_label(c, left + 2 * (col_w + 18), y, col_w, 45, "Outputs", "Margin, exposure, cash flow, break-even support, viability.", fill=LIGHT_GREEN)
    arrow(c, left + col_w + 3, y + 22, left + col_w + 15, y + 22)
    arrow(c, left + 2 * col_w + 21, y + 22, left + 2 * col_w + 33, y + 22)
    y -= 17

    draw_section_title(c, "8. Plain-English Terms", left, y, width)
    y -= 13
    term_w = (width - 12) / 2
    term_h = 19
    for idx, (term, definition) in enumerate(CONTENT["definitions"]):
        row = idx // 2
        col = idx % 2
        x = left + col * (term_w + 12)
        yy = y - row * (term_h + 4)
        c.setFillColor(WHITE)
        c.setStrokeColor(LIGHT_GRAY)
        c.roundRect(x, yy - term_h, term_w, term_h, 5, stroke=1, fill=1)
        c.setFillColor(GREEN)
        c.setFont("Helvetica-Bold", 5.9)
        c.drawString(x + 7, yy - 10, term)
        para(c, definition, x + 59, yy - 5, term_w - 66, term_h - 4, size=4.9, color=GRAY)
    y -= 53

    c.setFillColor(NAVY)
    c.roundRect(left, y - 40, width, 40, 7, stroke=0, fill=1)
    c.setFillColor(WHITE)
    c.setFont("Helvetica-Bold", 11.5)
    c.drawCentredString(W / 2, y - 17, "Providers adopt when financially viable, not only when clinically convinced.")
    c.setFont("Helvetica", 6.9)
    c.drawCentredString(W / 2, y - 29, "The calculator turns that viability test into a launch-planning and access-strategy tool.")
    c.setFillColor(GRAY)
    c.setFont("Helvetica", 6.6)
    c.drawString(left, 44, CONTENT["source_note"])


def build_pdf() -> Path:
    pdf_path = DELIVERY / PDF_NAME
    c = canvas.Canvas(str(pdf_path), pagesize=letter)
    c.setTitle("Adrabetadex Provider Economics and Adoption Risk")
    c.setAuthor("OpenAI Codex for JL Policy Consulting")
    draw_pdf_page_one(c)
    c.showPage()
    draw_pdf_page_two(c)
    c.save()
    shutil.copy2(pdf_path, PDF_DIR / PDF_NAME)
    return pdf_path


def ppt_add_text(slide, text, x, y, w, h, size=18, color=PPT_GRAY, bold=False, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    tx = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tx.text_frame.clear()
    tx.text_frame.margin_left = Inches(0.02)
    tx.text_frame.margin_right = Inches(0.02)
    tx.text_frame.margin_top = Inches(0.01)
    tx.text_frame.margin_bottom = Inches(0.01)
    tx.text_frame.vertical_anchor = valign
    p = tx.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tx


def ppt_box(slide, x, y, w, h, fill=PPT_PALE, line=PPT_MID_GREEN, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    box = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = line
    box.line.width = Pt(0.75)
    return box


def ppt_arrow(slide, x1, y1, x2, y2, color=PPT_GREEN):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(2)
    line.line.end_arrowhead = True
    return line


def ppt_header(slide, assets, label="Executive brief"):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PPT_WHITE
    slide.shapes.add_picture(str(assets["logo"]), Inches(0.35), Inches(0.18), width=Inches(2.45))
    ppt_add_text(slide, "Beren Therapeutics", 10.15, 0.29, 1.8, 0.24, size=9.5, color=PPT_GRAY, bold=True, align=PP_ALIGN.RIGHT)
    ppt_add_text(slide, label, 10.15, 0.53, 1.8, 0.2, size=7.5, color=PPT_GRAY, align=PP_ALIGN.RIGHT)
    rule = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), Inches(7.1), Inches(5.6), Inches(0.07))
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor(166, 210, 58)
    rule.line.fill.background()
    rule2 = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(6.25), Inches(7.1), Inches(6.7), Inches(0.07))
    rule2.fill.solid()
    rule2.fill.fore_color.rgb = RGBColor(38, 91, 44)
    rule2.line.fill.background()


def ppt_title(slide, title, subtitle=None, x=0.8, y=0.95, w=11.6):
    ppt_add_text(slide, title, x, y, w, 0.58, size=30, color=PPT_GREEN, bold=True)
    if subtitle:
        ppt_add_text(slide, subtitle, x, y + 0.58, w, 0.35, size=15.5, color=PPT_GRAY)


def ppt_table(slide, x, y, col_ws, headers, rows, row_h=0.48, font_size=9.5):
    total_w = sum(col_ws)
    ppt_box(slide, x, y, total_w, row_h, fill=PPT_LIGHT_GREEN, line=PPT_MID_GREEN, radius=False)
    cur_x = x
    for idx, header in enumerate(headers):
        ppt_add_text(slide, header, cur_x + 0.06, y + 0.08, col_ws[idx] - 0.12, row_h - 0.08, size=font_size, color=PPT_GREEN, bold=True)
        cur_x += col_ws[idx]
    for r_idx, row in enumerate(rows):
        ry = y + row_h * (r_idx + 1)
        ppt_box(slide, x, ry, total_w, row_h, fill=PPT_PALE if r_idx % 2 == 0 else PPT_WHITE, line=PPT_LIGHT_GRAY, radius=False)
        cur_x = x
        for c_idx, value in enumerate(row):
            ppt_add_text(
                slide,
                value,
                cur_x + 0.06,
                ry + 0.07,
                col_ws[c_idx] - 0.12,
                row_h - 0.06,
                size=font_size - 0.8,
                color=PPT_NAVY if c_idx == 0 else PPT_GRAY,
                bold=c_idx == 0,
            )
            cur_x += col_ws[c_idx]


def build_pptx(assets: dict[str, Path]) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Cover
    s = prs.slides.add_slide(blank)
    ppt_header(s, assets, "April 2026")
    ppt_add_text(s, "Adrabetadex", 0.85, 1.35, 4.4, 0.58, size=18, color=PPT_GRAY, bold=True)
    ppt_add_text(s, "Provider Economics\nand Adoption Risk", 0.85, 1.88, 8.9, 1.55, size=40, color=PPT_GREEN, bold=True)
    ppt_add_text(s, "Why Net Cost Recovery Determines Use", 0.9, 3.55, 6.2, 0.42, size=19, color=PPT_NAVY)
    ppt_box(s, 7.9, 1.18, 3.9, 3.3, fill=PPT_LIGHT_GREEN, line=PPT_MID_GREEN)
    ppt_add_text(s, "Management takeaway", 8.25, 1.55, 2.9, 0.35, size=14, color=PPT_GREEN, bold=True)
    ppt_add_text(s, CONTENT["thesis"], 8.25, 2.03, 3.05, 1.15, size=24, color=PPT_NAVY, bold=True)
    ppt_add_text(s, CONTENT["prepared"], 0.9, 5.85, 6.8, 0.3, size=10.5, color=PPT_GRAY)

    # 2 Executive thesis and definitions
    s = prs.slides.add_slide(blank)
    ppt_header(s, assets)
    ppt_title(s, "Clinical value does not make a provider financially able to treat", "The launch risk is operational economics at the treatment center.")
    ppt_box(s, 0.9, 2.05, 5.5, 1.2, fill=PPT_LIGHT_GREEN, line=PPT_MID_GREEN)
    ppt_add_text(s, "Decision question", 1.16, 2.28, 1.7, 0.25, size=12, color=PPT_GREEN, bold=True)
    ppt_add_text(s, "Does the provider make or lose money treating this patient?", 1.16, 2.62, 4.75, 0.35, size=20, color=PPT_NAVY, bold=True)
    ppt_box(s, 7.0, 1.9, 4.6, 2.0, fill=PPT_PALE, line=PPT_MID_GREEN)
    ppt_add_text(s, "The calculator translates reimbursement mechanics into a go / no-go adoption signal.", 7.35, 2.28, 3.85, 0.9, size=22, color=PPT_NAVY, bold=True)
    ppt_add_text(s, "Plain-English terms", 0.95, 4.0, 3.2, 0.3, size=16, color=PPT_GREEN, bold=True)
    for idx, (term, definition) in enumerate(CONTENT["definitions"]):
        x = 0.95 + (idx % 2) * 5.75
        y = 4.45 + (idx // 2) * 0.72
        ppt_add_text(s, term, x, y, 0.9, 0.25, size=10.5, color=PPT_GREEN, bold=True)
        ppt_add_text(s, definition, x + 0.95, y, 4.35, 0.4, size=9.2, color=PPT_GRAY)

    # 3 Core problem table
    s = prs.slides.add_slide(blank)
    ppt_header(s, assets)
    ppt_title(s, "The core problem is provider hesitation", "Each issue creates a direct barrier to adoption.")
    ppt_table(s, 0.85, 2.05, [2.35, 4.45, 4.6], ["Issue", "Provider concern", "Impact on adoption"], CONTENT["core_rows"], row_h=0.62, font_size=11)

    # 4 Coding pathway
    s = prs.slides.add_slide(blank)
    ppt_header(s, assets)
    ppt_title(s, "Coding is the gateway to payment", "Providers are evaluating the billing pathway, not only the drug.")
    start_x, y = 0.95, 2.55
    for idx, (title, body) in enumerate(CONTENT["coding_steps"]):
        x = start_x + idx * 2.85
        ppt_box(s, x, y, 2.2, 1.05, fill=PPT_LIGHT_GREEN if idx < 2 else PPT_PALE, line=PPT_MID_GREEN)
        ppt_add_text(s, title, x + 0.18, y + 0.2, 1.85, 0.28, size=14, color=PPT_NAVY, bold=True)
        ppt_add_text(s, body, x + 0.18, y + 0.55, 1.85, 0.28, size=10.5, color=PPT_GRAY)
        if idx < 3:
            ppt_arrow(s, x + 2.25, y + 0.52, x + 2.76, y + 0.52)
    ppt_add_text(s, "Launch risk: manual billing + high denial exposure", 2.95, 4.18, 3.6, 0.3, size=12, color=PPT_RED, bold=True)
    ppt_add_text(s, "Target state: standardized billing", 9.2, 4.18, 2.8, 0.3, size=12, color=PPT_GREEN, bold=True)
    ppt_box(s, 0.95, 5.05, 10.9, 0.82, fill=PPT_NAVY, line=PPT_NAVY)
    ppt_add_text(s, "Why it matters: coding determines how quickly the provider can submit a clean claim and get paid.", 1.25, 5.28, 10.1, 0.32, size=16, color=PPT_WHITE, bold=True)

    # 5 Pass-through and EAP
    s = prs.slides.add_slide(blank)
    ppt_header(s, assets)
    ppt_title(s, "Pass-through can stabilize early adoption, but only temporarily", "EAP experience does not prove commercial viability.")
    ppt_table(s, 0.75, 1.85, [2.25, 4.15, 3.5], ["Pass-through scenario", "Payment structure", "Provider risk"], CONTENT["pass_rows"], row_h=0.55, font_size=9.8)
    ppt_add_text(s, "EAP vs commercial buy-and-bill", 0.8, 4.25, 3.9, 0.3, size=16, color=PPT_GREEN, bold=True)
    ppt_table(s, 0.75, 4.65, [1.75, 3.1, 3.8], ["Dimension", "EAP", "Commercial"], CONTENT["eap_vs_commercial"][:3], row_h=0.42, font_size=8.6)
    ppt_box(s, 9.7, 4.55, 2.45, 1.35, fill=PPT_LIGHT_GREEN, line=PPT_MID_GREEN)
    ppt_add_text(s, "Key read", 10.0, 4.82, 1.5, 0.25, size=12, color=PPT_GREEN, bold=True)
    ppt_add_text(s, "A center that treats in EAP may still decline commercial launch.", 10.0, 5.18, 1.8, 0.42, size=13, color=PPT_NAVY, bold=True)

    # 6 Payer friction
    s = prs.slides.add_slide(blank)
    ppt_header(s, assets)
    ppt_title(s, "Payer friction is structural", "Payer rules can decide whether the provider can treat, will treat, or loses money treating.")
    for idx, (title, body) in enumerate(CONTENT["payer_friction"]):
        x = 1.0 + idx * 3.9
        ppt_box(s, x, 2.35, 3.15, 2.25, fill=PPT_PALE, line=PPT_MID_GREEN)
        dot = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x + 0.25), Inches(2.73), Inches(0.22), Inches(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = PPT_GREEN
        dot.line.fill.background()
        ppt_add_text(s, title, x + 0.58, 2.58, 2.1, 0.35, size=18, color=PPT_NAVY, bold=True)
        ppt_add_text(s, body, x + 0.58, 3.08, 2.25, 0.9, size=14, color=PPT_GRAY)
    ppt_add_text(s, "Why it matters: these controls can erase the provider's economic reason to activate a site.", 1.05, 5.35, 10.7, 0.38, size=17, color=PPT_GREEN, bold=True)

    # 7 340B impact
    s = prs.slides.add_slide(blank)
    ppt_header(s, assets)
    ppt_title(s, "340B can be the difference between early adoption and avoidance", "Acquisition economics shape which centers can take launch risk.")
    ppt_table(s, 0.8, 2.05, [1.8, 2.55, 2.55, 4.25], ["Scenario", "Acquisition", "Reimbursement", "Margin signal"], CONTENT["impact_340b"], row_h=0.66, font_size=10.2)
    ppt_box(s, 0.9, 5.25, 10.95, 0.74, fill=PPT_NAVY, line=PPT_NAVY)
    ppt_add_text(s, "Key read: early adoption likely concentrates where acquisition economics and billing readiness both work.", 1.25, 5.47, 10.0, 0.25, size=16, color=PPT_WHITE, bold=True)

    # 8 Calculator value
    s = prs.slides.add_slide(blank)
    ppt_header(s, assets)
    ppt_title(s, "The calculator turns complexity into a provider go / no-go answer", "Leadership sees margin, cash exposure, and break-even support by scenario.")
    panels = [
        ("Inputs", ", ".join(CONTENT["inputs"])),
        ("Model", "Expected net provider position after payment, acquisition, denial, rework, labor, and carrying cost."),
        ("Outputs", ", ".join(CONTENT["outputs"])),
    ]
    for idx, (title, body) in enumerate(panels):
        x = 0.9 + idx * 4.1
        ppt_box(s, x, 2.3, 3.15, 2.1, fill=PPT_LIGHT_GREEN if idx != 1 else PPT_PALE, line=PPT_MID_GREEN)
        ppt_add_text(s, title, x + 0.28, 2.65, 2.4, 0.34, size=19, color=PPT_NAVY, bold=True)
        ppt_add_text(s, body, x + 0.28, 3.16, 2.45, 0.82, size=12.6, color=PPT_GRAY)
        if idx < 2:
            ppt_arrow(s, x + 3.18, 3.35, x + 4.0, 3.35)
    ppt_add_text(s, "Primary output: Does this provider make or lose money treating this patient?", 1.05, 5.35, 10.6, 0.4, size=18, color=PPT_GREEN, bold=True)

    # 9 Bottom line
    s = prs.slides.add_slide(blank)
    ppt_header(s, assets)
    ppt_add_text(s, "Bottom line", 0.9, 1.15, 3.0, 0.42, size=18, color=PPT_GRAY, bold=True)
    ppt_add_text(s, CONTENT["thesis"], 0.9, 1.85, 9.9, 1.65, size=38, color=PPT_GREEN, bold=True)
    ppt_box(s, 0.95, 4.4, 3.25, 1.08, fill=PPT_PALE, line=PPT_MID_GREEN)
    ppt_add_text(s, "Without the model", 1.23, 4.68, 1.9, 0.24, size=12, color=PPT_RED, bold=True)
    ppt_add_text(s, "Adoption may be overestimated because EAP use and clinical interest are misread as commercial readiness.", 1.23, 5.0, 2.45, 0.3, size=10.4, color=PPT_GRAY)
    ppt_box(s, 4.95, 4.4, 3.25, 1.08, fill=PPT_LIGHT_GREEN, line=PPT_MID_GREEN)
    ppt_add_text(s, "With the model", 5.23, 4.68, 1.9, 0.24, size=12, color=PPT_GREEN, bold=True)
    ppt_add_text(s, "Leadership sees which sites, channels, and support levers make treatment financially viable.", 5.23, 5.0, 2.45, 0.3, size=10.4, color=PPT_GRAY)
    ppt_box(s, 8.95, 4.4, 2.85, 1.08, fill=PPT_NAVY, line=PPT_NAVY)
    ppt_add_text(s, "Use the result", 9.22, 4.68, 1.7, 0.24, size=12, color=PPT_WHITE, bold=True)
    ppt_add_text(s, "Prioritize launch sites and field reimbursement strategy by provider economics.", 9.22, 5.0, 2.0, 0.3, size=10.4, color=PPT_WHITE)

    pptx_path = DELIVERY / PPTX_NAME
    prs.save(pptx_path)
    return pptx_path


def load_font(size: int, bold: bool = False):
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
        "/Library/Fonts/Arial.ttf",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).exists():
            try:
                return ImageFont.truetype(candidate, size)
            except Exception:
                continue
    return ImageFont.load_default()


def wrap_draw(draw: ImageDraw.ImageDraw, text: str, xy: tuple[int, int], width: int, font, fill, line_gap: int = 4) -> int:
    words = text.split()
    lines = []
    current = ""
    for word in words:
        test = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), test, font=font)[2] <= width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    x, y = xy
    for line in lines:
        draw.text((x, y), line, fill=fill, font=font)
        y += font.size + line_gap
    return y


def draw_preview_slide(index: int, title: str, subtitle: str, body_lines: list[str], assets: dict[str, Path]) -> Path:
    img = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(img)
    logo = Image.open(assets["logo"]).convert("RGBA")
    logo.thumbnail((320, 90))
    img.paste(logo, (70, 36), logo)
    draw.text((1210, 44), "Beren Therapeutics", fill=(77, 86, 86), font=load_font(24, bold=True))
    draw.rectangle((75, 850, 720, 861), fill=(166, 210, 58))
    draw.rectangle((720, 850, 1520, 861), fill=(38, 91, 44))
    draw.text((105, 150), title, fill=(78, 138, 45), font=load_font(50, bold=True))
    if subtitle:
        wrap_draw(draw, subtitle, (107, 222), 1180, load_font(28), (77, 86, 86), line_gap=8)
    y = 330
    for line in body_lines:
        draw.rounded_rectangle((105, y, 1360, y + 82), radius=16, outline=(199, 221, 187), fill=(247, 250, 246), width=2)
        wrap_draw(draw, line, (135, y + 20), 1180, load_font(28, bold=True), (31, 46, 43), line_gap=7)
        y += 105
    path = PREVIEW_DIR / f"slide_{index:02d}.png"
    img.save(path)
    return path


def build_previews(assets: dict[str, Path]) -> list[Path]:
    slides = [
        (1, "Adrabetadex Provider Economics and Adoption Risk", "Why Net Cost Recovery Determines Use", [CONTENT["thesis"]]),
        (2, "Clinical value does not make a provider financially able to treat", "The launch risk is operational economics at the treatment center.", ["Decision question: Does the provider make or lose money treating this patient?"]),
        (3, "The core problem is provider hesitation", "Each issue creates a direct barrier to adoption.", ["Acquisition, coding, payment timing, denials, and site controls each change provider willingness to treat."]),
        (4, "Coding is the gateway to payment", "Providers are evaluating the billing pathway, not only the drug.", ["EAP to NOC code to J-code to stable billing."]),
        (5, "Pass-through can stabilize early adoption", "EAP experience does not prove commercial viability.", ["Temporary payment visibility helps, but commercial buy-and-bill still exposes the provider."]),
        (6, "Payer friction is structural", "Payer rules can decide whether the provider can treat.", ["Prior authorization, white-bagging, and site-of-care policies directly shape provider economics."]),
        (7, "340B can determine early site activation", "Acquisition economics shape which centers can take launch risk.", ["340B COEs may have a stronger per-dose margin runway than non-340B sites."]),
        (8, "The calculator turns complexity into a go / no-go answer", "Leadership sees margin, cash exposure, and break-even support.", ["Inputs to model to outputs: margin, cash flow, and adoption viability."]),
        (9, "Bottom line", "", [CONTENT["thesis"], "Use the result to prioritize launch sites and field reimbursement strategy."]),
    ]
    return [draw_preview_slide(*slide, assets=assets) for slide in slides]


def main() -> None:
    ensure_dirs()
    assets = prepare_assets()
    pdf_path = build_pdf()
    pptx_path = build_pptx(assets)
    preview_paths = build_previews(assets)
    manifest = {
        "pdf": str(pdf_path),
        "pdf_copy": str(PDF_DIR / PDF_NAME),
        "pptx": str(pptx_path),
        "previews": [str(path) for path in preview_paths],
        "source": "Generated from supplied April 2026 Adrabetadex management brief prompt and existing Navisync letterhead PDF.",
    }
    (DELIVERY / REPORT_NAME).write_text(json.dumps(manifest, indent=2) + "\n")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
