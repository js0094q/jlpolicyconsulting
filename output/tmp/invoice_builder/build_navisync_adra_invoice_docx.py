from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REPO_ROOT = Path("/Users/josephstewart/Documents/JLPolicyConsulting")
DATA_PATH = REPO_ROOT / "output/tmp/invoice_builder/navisync_adra_invoice_data.json"
OUTPUT_PATH = REPO_ROOT / "output/exports/invoices/JLPC-2026-0501_Navisync_Adrabetadex_Beren_Hourly_Invoice.docx"


def money(value: float) -> str:
    return f"${value:,.2f}"


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 8.0, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def set_cell_width(cell, width: float) -> None:
    cell.width = Inches(width)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table_borders(table, color: str = "C7D1DD", size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_label_value_table(doc: Document, rows: list[tuple[str, str]], widths: tuple[float, float]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    style_table_borders(table)
    for row_idx, (label, value) in enumerate(rows):
        row = table.rows[row_idx]
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_width(row.cells[0], widths[0])
        set_cell_width(row.cells[1], widths[1])
        set_cell_fill(row.cells[0], "EEF4F8")
        set_cell_text(row.cells[0], label, bold=True, size=8.5)
        set_cell_text(row.cells[1], value, size=8.5)


def add_line_item_table(doc: Document, items: list[dict], rate: float) -> None:
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    style_table_borders(table)
    widths = [1.55, 6.3, 0.55, 0.75, 0.9]
    headers = ["Date / time evidence", "Project / deliverable", "Hours", "Rate", "Amount"]
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for idx, cell in enumerate(header_row.cells):
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_width(cell, widths[idx])
        set_cell_fill(cell, "DCE8F2")
        set_cell_text(cell, headers[idx], bold=True, size=7.8)

    for item in items:
        row = table.add_row()
        values = [
            item["dateEvidence"],
            item["deliverable"],
            f"{float(item['hours']):.1f}",
            money(rate),
            money(float(item["hours"]) * rate),
        ]
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
            set_cell_width(cell, widths[idx])
            set_cell_text(cell, values[idx], size=7.4)
            if idx in (2, 3, 4):
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT


def build_docx() -> None:
    data = json.loads(DATA_PATH.read_text())
    items = data["items"]
    rate = float(data["rate"])
    total_hours = sum(float(item["hours"]) for item in items)
    total_amount = total_hours * rate

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(8.5)
    styles["Heading 1"].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(13)
    styles["Heading 1"].font.bold = True

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("JL Policy Consulting, LLC | Hourly Invoice")
    run.bold = True
    run.font.size = Pt(15)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    subtitle.add_run("Navisync / Beren Therapeutics Launch Team / Adrabetadex").font.size = Pt(9)

    header = doc.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    header.autofit = False
    style_table_borders(header, color="9AAABE")
    left, right = header.rows[0].cells
    set_cell_width(left, 5.05)
    set_cell_width(right, 5.05)
    set_cell_fill(left, "F8FAFC")
    set_cell_fill(right, "F8FAFC")
    set_cell_text(
        left,
        f"{data['seller']['name']}\n"
        f"{data['seller']['contact']}\n"
        f"{data['seller']['title']}\n"
        f"{data['seller']['email']}",
        size=8.5,
    )
    set_cell_text(
        right,
        f"Invoice Number: {data['invoiceNumber']}\n"
        f"Invoice Date: {data['invoiceDate']}\n"
        f"Service Period: {data['servicePeriod']}\n"
        f"Payment Terms: {data['paymentTerms']}\n"
        f"Currency: {data['currency']}",
        size=8.5,
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_label_value_table(
        doc,
        [
            ("Bill To", " / ".join(data["billTo"][:3])),
            ("Scope", data["billTo"][3]),
            ("Hours", f"{total_hours:.1f}"),
            ("Rate", f"{money(rate)} / hour"),
            ("Total Due", money(total_amount)),
        ],
        (1.1, 8.95),
    )

    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(4)
    heading.add_run("Line Items").bold = True

    add_line_item_table(doc, items[:7], rate)
    doc.add_page_break()
    continued = doc.add_paragraph()
    continued.paragraph_format.space_after = Pt(4)
    continued.add_run("Line Items Continued").bold = True
    add_line_item_table(doc, items[7:], rate)

    totals = doc.add_table(rows=3, cols=2)
    totals.alignment = WD_TABLE_ALIGNMENT.RIGHT
    totals.autofit = False
    style_table_borders(totals, color="C4A353")
    total_rows = [("Subtotal", money(total_amount)), ("Expenses", money(0)), ("TOTAL DUE", money(total_amount))]
    for idx, (label, value) in enumerate(total_rows):
        row = totals.rows[idx]
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_width(row.cells[0], 1.3)
        set_cell_width(row.cells[1], 1.25)
        fill = "F6E9C8" if idx == 2 else "FFFFFF"
        set_cell_fill(row.cells[0], fill)
        set_cell_fill(row.cells[1], fill)
        set_cell_text(row.cells[0], label, bold=idx == 2, size=8)
        set_cell_text(row.cells[1], value, bold=idx == 2, size=8)
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    notes_heading = doc.add_paragraph()
    notes_heading.paragraph_format.space_before = Pt(6)
    notes_heading.paragraph_format.space_after = Pt(2)
    notes_heading.add_run("Notes").bold = True
    for note in data["notes"]:
        paragraph = doc.add_paragraph(style=None)
        paragraph.paragraph_format.left_indent = Inches(0.12)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run("- " + note).font.size = Pt(7.8)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(
        json.dumps(
            {
                "outputPath": str(OUTPUT_PATH),
                "lineItems": len(items),
                "totalHours": total_hours,
                "totalAmount": total_amount,
            }
        )
    )


if __name__ == "__main__":
    build_docx()
