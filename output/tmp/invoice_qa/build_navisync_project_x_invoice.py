from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path("/Users/josephstewart/Documents/JLPolicyConsulting")
OUT_DIR = ROOT / "output" / "exports" / "invoices"
TMP_DIR = ROOT / "output" / "tmp" / "invoice_qa"
DOCX_PATH = OUT_DIR / "JLPC-2026-0502_Navisync_Adrabetadex_Beren_Hourly_Invoice_PROJECT_X_TEMPLATE.docx"

BLUE = "2F5FBE"
DARK_BLUE = "0F1A2F"
MUTED = "4D5A72"
LIGHT_BORDER = "CBD8EA"
LIGHT_FILL = "F4F6FA"
WHITE = "FFFFFF"
FONT_NAME = "Helvetica"
CONTENT_WIDTH = 7.0

LINE_ITEMS = [
    (
        "Mar. 18 – Mar. 22, 2026",
        "Source review, research and development of initial reimbursement and access materials, including synthesis of clinical inputs, market access considerations, and creation of foundational client-facing reimbursement documents and structured concept brief framework.",
        Decimal("11.0"),
        Decimal("200.00"),
        Decimal("2200.00"),
    ),
    (
        "Apr. 9 – Apr. 10, 2026",
        "Development of Adrabetadex Provider Net Cost Calculator executive summary materials, including presentation-ready outputs, structured narrative, and formatted client deliverables for internal and client review for April 13th in-person launch meeting.",
        Decimal("5.0"),
        Decimal("200.00"),
        Decimal("1000.00"),
    ),
    (
        "Apr. 17, 2026",
        "Coding and reimbursement pathway analysis, including hospital outpatient billing structure, HCPCS strategy, OPPS/pass-through considerations, and provider economics implications for launch planning.",
        Decimal("2.0"),
        Decimal("200.00"),
        Decimal("400.00"),
    ),
    (
        "Apr. 19 – Apr. 23, 2026",
        "Initial build of Adrabetadex Provider Net Cost Recovery Calculator, including multi-scenario modeling, workbook architecture, reimbursement logic integration, and supporting documentation for ARM and client use.",
        Decimal("9.0"),
        Decimal("200.00"),
        Decimal("1800.00"),
    ),
    (
        "Apr. 21 – Apr. 24, 2026",
        "Initial research and development of comprehensive ARM training package, including training guide, provider billing references, reimbursement materials, presentation decks, and supporting outputs for field deployment.",
        Decimal("12.5"),
        Decimal("200.00"),
        Decimal("2500.00"),
    ),
    (
        "Apr. 23, 2026",
        "Concept brief document revision and reformatting, including standardization of materials, terminology alignment, and final client-ready PDF outputs.",
        Decimal("2.5"),
        Decimal("200.00"),
        Decimal("500.00"),
    ),
    (
        "Apr. 24 – Apr. 26, 2026",
        "Development of client-facing concept brief outlining reimbursement model and provider economics, including structured narrative, payer dynamics framing, and supporting materials for MLR/client review.",
        Decimal("5.5"),
        Decimal("200.00"),
        Decimal("1100.00"),
    ),
    (
        "Apr. 27, 2026",
        "Revision of provider economics and adoption risk executive presentation materials  including PowerPoint/PDF outputs aligned to Navisync presentation standards.",
        Decimal("4.0"),
        Decimal("200.00"),
        Decimal("800.00"),
    ),
    (
        "Apr. 28, 2026",
        "Refinement of Provider Net Cost Recovery Model concept brief, including enhanced reimbursement model framing, structured presentation updates, and final deliverable outputs.",
        Decimal("5.0"),
        Decimal("200.00"),
        Decimal("1000.00"),
    ),
    (
        "Apr. 28, 2026",
        "Concept brief comparison and quality review, including validation of content accuracy, structural consistency, and presentation quality across versions.",
        Decimal("2.0"),
        Decimal("200.00"),
        Decimal("400.00"),
    ),
    (
        "Apr. 29, 2026",
        "ARM training material revisions, including updated billing references, structured training modules, and revised presentation decks based on Beren feedback and reimbursement positioning.",
        Decimal("10.0"),
        Decimal("200.00"),
        Decimal("2000.00"),
    ),
    (
        "Apr. 29 – Apr. 30, 2026",
        "Update of ARM training modules A/B/C, including integration of source materials, rebuild of training documents, and final validation of client-ready deliverables.",
        Decimal("2.0"),
        Decimal("200.00"),
        Decimal("400.00"),
    ),
]


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def money(value: Decimal) -> str:
    return f"${value:,.2f}"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=72, start=72, bottom=72, end=72) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=LIGHT_BORDER, size="6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def clear_table_borders(table) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_table_width(table, inches: float) -> None:
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(inches * 1440)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.first_child_found_in("w:tblLayout")
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER


def set_column_widths(table, widths: list[float]) -> None:
    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        tbl_grid.append(grid_col)
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_paragraph_text(
    paragraph,
    text: str,
    *,
    size: float = 9.0,
    color: str = DARK_BLUE,
    bold: bool = False,
    align=None,
    space_after: float = 0,
    line_spacing: float = 1.0,
) -> None:
    paragraph.clear()
    paragraph.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing
    run = paragraph.add_run(text)
    run.font.name = FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold


def add_text(cell, text: str, **kwargs) -> None:
    paragraph = cell.paragraphs[0]
    set_paragraph_text(paragraph, text, **kwargs)


def add_lines(cell, lines: list[tuple[str, dict]]) -> None:
    cell.text = ""
    for idx, (text, opts) in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        set_paragraph_text(p, text, **opts)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = FONT_NAME
    run.font.size = Pt(8.5)
    run.font.color.rgb = rgb(MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def make_logo(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    img = Image.new("RGBA", (320, 320), (255, 255, 255, 0))
    draw = ImageDraw.Draw(img)
    draw.rounded_rectangle((0, 0, 320, 320), radius=40, fill=(34, 69, 153, 255))
    try:
        font = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 126)
    except OSError:
        font = ImageFont.load_default()
    bbox = draw.textbbox((0, 0), "JL", font=font)
    x = (320 - (bbox[2] - bbox[0])) / 2
    y = (320 - (bbox[3] - bbox[1])) / 2 - 6
    draw.text((x, y), "JL", font=font, fill=(255, 255, 255, 255))
    img.save(path)


def add_header(section, logo_path: Path) -> None:
    header = section.header
    for p in header.paragraphs:
        p.text = ""
        p.paragraph_format.space_after = Pt(0)

    table = header.add_table(rows=1, cols=3, width=Inches(CONTENT_WIDTH))
    clear_table_borders(table)
    set_table_width(table, CONTENT_WIDTH)
    set_column_widths(table, [0.50, 4.45, 2.05])
    row = table.rows[0]
    row.height = Inches(0.40)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY

    logo_cell, brand_cell, date_cell = row.cells
    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)

    logo_p = logo_cell.paragraphs[0]
    logo_p.paragraph_format.space_after = Pt(0)
    logo_p.add_run().add_picture(str(logo_path), width=Inches(0.40))

    brand_cell.text = ""
    p1 = brand_cell.paragraphs[0]
    set_paragraph_text(p1, "JL Policy Consulting", size=11.6, color=BLUE, bold=True, space_after=0)
    p2 = brand_cell.add_paragraph()
    set_paragraph_text(p2, "LLC", size=8.5, color=MUTED, space_after=0)

    set_paragraph_text(date_cell.paragraphs[0], "05/02/26", size=8.0, color=MUTED, align=WD_ALIGN_PARAGRAPH.RIGHT)

    spacer = header.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.line_spacing = 0.25

    line = header.add_table(rows=1, cols=1, width=Inches(CONTENT_WIDTH))
    clear_table_borders(line)
    set_table_width(line, CONTENT_WIDTH)
    set_column_widths(line, [CONTENT_WIDTH])
    line.rows[0].height = Pt(2.0)
    line.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    set_cell_margins(line.cell(0, 0), top=0, start=0, bottom=0, end=0)
    set_cell_shading(line.cell(0, 0), BLUE)


def add_footer(section) -> None:
    footer = section.footer
    for p in footer.paragraphs:
        p.text = ""
        p.paragraph_format.space_after = Pt(0)

    line = footer.add_table(rows=1, cols=1, width=Inches(CONTENT_WIDTH))
    clear_table_borders(line)
    set_table_width(line, CONTENT_WIDTH)
    set_column_widths(line, [CONTENT_WIDTH])
    line.rows[0].height = Pt(2.0)
    line.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    set_cell_margins(line.cell(0, 0), top=0, start=0, bottom=0, end=0)
    set_cell_shading(line.cell(0, 0), BLUE)

    table = footer.add_table(rows=1, cols=2, width=Inches(CONTENT_WIDTH))
    clear_table_borders(table)
    set_table_width(table, CONTENT_WIDTH)
    set_column_widths(table, [CONTENT_WIDTH / 2, CONTENT_WIDTH / 2])
    for cell in table.rows[0].cells:
        set_cell_margins(cell, top=72, start=0, bottom=0, end=0)
    set_paragraph_text(table.cell(0, 0).paragraphs[0], "JL Policy Consulting", size=8.5, color=MUTED)
    add_page_number(table.cell(0, 1).paragraphs[0])


def style_grid_table(table, widths: list[float], header_rows: int = 0) -> None:
    set_table_width(table, sum(widths))
    set_column_widths(table, widths)
    for r_idx, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        set_row_cant_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell, top=70, start=70, bottom=70, end=70)
            if r_idx < header_rows:
                set_cell_shading(cell, LIGHT_FILL)
        if r_idx < header_rows:
            repeat_table_header(row)


def add_meta_table(doc: Document) -> None:
    table = doc.add_table(rows=3, cols=4)
    style_grid_table(table, [1.25, 2.10, 1.25, 2.40])
    data = [
        ("Invoice Number", "JLPC-2026-0502-ADRA-BEREN", "Invoice Date", "May 2, 2026"),
        ("Service Period", "March 18, 2026 – April 30, 2026", "Payment Terms", "Net 30"),
        ("Currency", "USD", "", ""),
    ]
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx]
        row.height = Inches(0.25)
        for c_idx, value in enumerate(row_data):
            cell = table.cell(r_idx, c_idx)
            if c_idx in (0, 2):
                set_cell_shading(cell, LIGHT_FILL)
            add_text(cell, value, size=8.9, color="000000")


def add_bill_engagement(doc: Document) -> None:
    table = doc.add_table(rows=1, cols=2)
    style_grid_table(table, [3.50, 3.50])
    left, right = table.rows[0].cells
    add_lines(
        left,
        [
            ("Bill To", {"size": 10.0, "color": DARK_BLUE, "space_after": 1}),
            ("Navisync, LLC", {"size": 8.9, "color": DARK_BLUE, "space_after": 0}),
            ("445 South Street, Suite 305", {"size": 8.9, "color": DARK_BLUE, "space_after": 0}),
            ("Morristown, NJ 07960", {"size": 8.9, "color": DARK_BLUE, "space_after": 0}),
        ],
    )
    add_lines(
        right,
        [
            ("Engagement Reference", {"size": 10.0, "color": DARK_BLUE, "space_after": 1}),
            ("Project Name: Adrabetadex / Beren Therapeutics Launch Planning", {"size": 8.9, "color": DARK_BLUE, "space_after": 0}),
            ("Service Type: Health policy, reimbursement, provider economics, and ARM training support", {"size": 8.9, "color": DARK_BLUE, "space_after": 0}),
        ],
    )


def add_line_items(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    set_paragraph_text(p, "Line Items", size=10.5, color=BLUE)

    def add_items_table(items) -> None:
        table = doc.add_table(rows=1, cols=5)
        style_grid_table(table, [0.90, 3.75, 0.50, 0.80, 1.05], header_rows=1)
        headers = ["Date", "Description", "Hours", "Rate", "Amount"]
        for idx, header in enumerate(headers):
            align = WD_ALIGN_PARAGRAPH.RIGHT if idx >= 2 else WD_ALIGN_PARAGRAPH.LEFT
            add_text(table.cell(0, idx), header, size=7.6, color="000000", align=align)

        for date, desc, hours, rate, amount in items:
            row = table.add_row()
            values = [date, desc, f"{hours:.1f}", money(rate), money(amount)]
            for idx, value in enumerate(values):
                cell = row.cells[idx]
                set_cell_border(cell)
                set_cell_margins(cell, top=62, start=62, bottom=62, end=62)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                align = WD_ALIGN_PARAGRAPH.RIGHT if idx >= 2 else WD_ALIGN_PARAGRAPH.LEFT
                size = 6.8 if idx == 1 else 7.0
                add_text(cell, value, size=size, color=DARK_BLUE, align=align, line_spacing=0.92)

    add_items_table(LINE_ITEMS)


def add_totals_payment_notes(doc: Document) -> None:
    total_hours = sum(item[2] for item in LINE_ITEMS)
    subtotal = sum(item[4] for item in LINE_ITEMS)
    assert total_hours == Decimal("70.5"), total_hours
    assert subtotal == Decimal("14100.00"), subtotal

    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)

    table = doc.add_table(rows=3, cols=2)
    style_grid_table(table, [1.60, 5.40])
    totals = [
        ("Subtotal", money(subtotal)),
        ("Expenses", money(Decimal("0.00"))),
        ("TOTAL DUE", money(subtotal)),
    ]
    for idx, (label, value) in enumerate(totals):
        row = table.rows[idx]
        row.height = Inches(0.24)
        left, right = row.cells
        set_cell_shading(left, LIGHT_FILL)
        bold = label == "TOTAL DUE"
        add_text(left, label, size=8.9, color="000000", bold=bold)
        add_text(right, value, size=8.9, color=DARK_BLUE, bold=bold)

    pay = doc.add_paragraph()
    pay.paragraph_format.keep_with_next = True
    pay.paragraph_format.space_before = Pt(8)
    pay.paragraph_format.space_after = Pt(2)
    set_paragraph_text(pay, "Payment Instructions", size=10.5, color=BLUE)

    payment_lines = [
        "Bank ACH",
        "Account Name: Joseph Stewart",
        "Bank: Capital One, N.A.",
        "Account Number: 36091605578",
        "Routing: 031176110",
        "SWIFT: HIBKUS44",
    ]
    for idx, line in enumerate(payment_lines):
        p = doc.add_paragraph()
        if idx < len(payment_lines) - 1:
            p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(0)
        set_paragraph_text(p, line, size=8.9, color=DARK_BLUE)

    notes_header = doc.add_paragraph()
    notes_header.paragraph_format.space_before = Pt(7)
    notes_header.paragraph_format.space_after = Pt(1)
    set_paragraph_text(notes_header, "Notes", size=10.5, color=BLUE)

    notes = [
        "Services reflect Navisync/Beren Adrabetadex Net Cost Recovery Modell including reimbursement mechanics, provider economics, launch-access, and ARM training support delivered during the stated service period.",
        "Line items preserve the underlying hourly allocation while presenting the work as deliverable-based consulting support.",
        "Supporting documentation and underlying work products are available upon request.",
    ]
    for note in notes:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.1)
        p.paragraph_format.first_line_indent = Inches(-0.1)
        p.paragraph_format.space_after = Pt(0)
        set_paragraph_text(p, "- " + note, size=8.1, color=MUTED, line_spacing=0.96)


def build() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    logo_path = TMP_DIR / "jl-logo.png"
    make_logo(logo_path)

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.95)
    section.bottom_margin = Inches(0.90)
    section.header_distance = Inches(0.18)
    section.footer_distance = Inches(0.25)

    styles = doc.styles
    styles["Normal"].font.name = FONT_NAME
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    styles["Normal"].font.size = Pt(9.3)
    styles["Normal"].font.color.rgb = rgb(DARK_BLUE)

    add_header(section, logo_path)
    add_footer(section)

    top = doc.add_table(rows=1, cols=2)
    clear_table_borders(top)
    set_table_width(top, CONTENT_WIDTH)
    set_column_widths(top, [4.80, 2.20])
    for cell in top.rows[0].cells:
        set_cell_margins(cell, top=0, start=0, bottom=0, end=0)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    add_lines(
        top.cell(0, 0),
        [
            ("JL Policy Consulting, LLC", {"size": 9.3, "color": DARK_BLUE, "bold": True, "space_after": 0}),
            ("Joseph L. Stewart, MPH, CPhT", {"size": 9.3, "color": DARK_BLUE, "space_after": 0}),
            ("Managing Director, Health Policy & Reimbursement", {"size": 9.3, "color": DARK_BLUE, "space_after": 0}),
        ],
    )
    set_paragraph_text(
        top.cell(0, 1).paragraphs[0],
        "INVOICE",
        size=24.0,
        color=DARK_BLUE,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)

    for line in [
        "45 Beyers Road",
        "Montgomery, NY 12549",
        "United States",
        "Email: Joseph.Stewart@JLPolicyConsulting.com",
        "Phone: +1 (845) 779-2447",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        set_paragraph_text(p, line, size=9.3, color=DARK_BLUE, line_spacing=1.0)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)

    add_meta_table(doc)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    spacer.paragraph_format.line_spacing = 0.35
    add_bill_engagement(doc)
    add_line_items(doc)
    add_totals_payment_notes(doc)

    doc.core_properties.title = "Invoice JLPC-2026-0502-ADRA-BEREN"
    doc.core_properties.author = "JL Policy Consulting, LLC"
    doc.core_properties.subject = "Navisync Adrabetadex / Beren Therapeutics hourly invoice"
    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    build()
