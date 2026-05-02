from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/josephstewart/Documents/JLPolicyConsulting")
TEMPLATE = Path("/Users/josephstewart/Downloads/Navisync Corporate Template.docx")
SOURCE_DOCX = Path("/Users/josephstewart/Downloads/Adrabetadex_Navisync_Final_Combined.docx")
SOURCE_PDF = Path("/Users/josephstewart/Downloads/Adrabetadex Provider Net Cost Calculator Concept Brief 4_1_26.pdf")
OUT_DIR = ROOT / "output" / "exports" / "adrabetadex_provider_net_cost_recovery_model_concept_brief_2026-04-28"
DOCX_PATH = OUT_DIR / "Adrabetadex_Provider_Net_Cost_Recovery_Model_Concept_Brief.docx"
PDF_PATH = OUT_DIR / "Adrabetadex_Provider_Net_Cost_Recovery_Model_Concept_Brief.pdf"
MANIFEST_PATH = OUT_DIR / "build_manifest.json"

TITLE = "Adrabetadex Provider Net Cost Recovery Model"
SUBTITLE = "Executive Brief | April 2026"

GREEN = "4E8A2D"
DARK_GREEN = "244B25"
TEXT = "394344"
PALE_GREEN = "EEF5EA"
PALE_GRAY = "F6F8F5"
LINE = "BFD6B4"
WHITE = "FFFFFF"

CONTENT_WIDTH = 7.05
CONTENT_TWIPS = int(CONTENT_WIDTH * 1440)


SOURCES = [
    (
        "[1] Beren Therapeutics / Navisync planning assumptions",
        "Client-supplied assumptions for PDUFA timing, clinical narrative, WAC, vial size, dosing frequency, and annual administrations.",
        "Internal planning assumptions supplied for this concept brief.",
    ),
    (
        "[2] CMS, HCPCS Level II Coding Procedures",
        "HCPCS drug and biological application deadlines, effective-date cycles, and CMS statement that coding does not itself determine coverage or payment.",
        "https://www.cms.gov/medicare/coding-billing/healthcare-common-procedure-system/level-ii-coding-process",
    ),
    (
        "[3] CMS, Part B Drug Payment Limits Overview",
        "Initial sales period WAC logic, ASP timing, AWP, contractor pricing, and OPPS payment for drugs without an assigned HCPCS code.",
        "https://www.cms.gov/files/document/part-b-drug-payment-limits-overview.pdf-0",
    ),
    (
        "[4] CMS, C9399 unclassified outpatient drug payment",
        "C9399 logic for certain new FDA-approved hospital outpatient drugs without product-specific HCPCS codes.",
        "https://www.cms.gov/newsroom/press-releases/medicare-pay-unclassified-fda-approved-drugs-administered-outpatient-departments",
    ),
    (
        "[5] CMS, OPPS Pass-Through Payment Status page",
        "MEARIS submission pathway for OPPS drug and biological pass-through applications.",
        "https://www.cms.gov/medicare/payment/prospective-payment-systems/hospital-outpatient/pass-through-payment-status-new-technology-ambulatory-payment-classification-apc",
    ),
    (
        "[6] CMS, Transitional Pass-Through Drug and Biological Eligibility Process",
        "Eligibility process and planning timing for transitional pass-through status.",
        "https://www.cms.gov/files/document/determine-eligibility-drugs-biologicals-transitional-pass-through-payment-under-hospital-outpatient.pdf",
    ),
    (
        "[7] HRSA, 340B ceiling price calculation",
        "340B ceiling price framework based on AMP minus URA.",
        "https://www.hrsa.gov/about/faqs/how-340b-ceiling-price-calculated",
    ),
    (
        "[8] HRSA, 340B pricing formulas",
        "Operational 340B package-size and case-pack pricing formulas.",
        "https://340bpricingsubmissions.hrsa.gov/Help/Manufacturer/Pricing%20Formulas/Pricing%20Formulas.htm",
    ),
    (
        "[9] Medicaid.gov, EPSDT",
        "Medicaid EPSDT coverage framework for children under age 21.",
        "https://www.medicaid.gov/medicaid/benefits/early-and-periodic-screening-diagnostic-and-treatment",
    ),
    (
        "[10] Medicaid.gov, State Drug Utilization Data FAQ",
        "NDC-based utilization reporting logic under the Medicaid Drug Rebate Program.",
        "https://www.medicaid.gov/faq/regarding-state-drug-utilization-data-sdud-there-official-crosswalk-of-national-drug-code-ndc-and-healthcare-common-procedure-coding-system-hcpcs-codes-which-providers-can-use/index.html",
    ),
    (
        "[11] UnitedHealthcare, medication sourcing protocol",
        "Public commercial-policy analogue for specialty-pharmacy sourcing and provider administration-only billing.",
        "https://www.uhcprovider.com/content/dam/provider/docs/public/resources/pharmacy/medication-expansion-sourcing-faq.pdf",
    ),
    (
        "[12] UnitedHealthcare, Provider Administered Drugs - Site of Care policy",
        "Public commercial-policy analogue for provider-administered drug site-of-care management.",
        "https://www.uhcprovider.com/content/dam/provider/docs/public/policies/index/commercial/provider-administered-drugs-soc-04012026.pdf",
    ),
]


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = rgb(color)


def add_para(
    doc,
    text: str = "",
    *,
    size: float = 9.1,
    color: str = TEXT,
    bold: bool = False,
    before: float = 0,
    after: float = 4.2,
    line_spacing: float = 1.04,
    align: int | None = None,
):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text: str, *, level: int = 1, before: float = 10, after: float = 4):
    size = 14 if level == 1 else 10.8
    p = add_para(doc, text, size=size, color=GREEN, bold=True, before=before, after=after, line_spacing=1.02)
    p.paragraph_format.keep_with_next = True
    return p


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=55, start=90, bottom=55, end=90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_twips: int = CONTENT_TWIPS) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "0")
    tbl_ind.set(qn("w:type"), "dxa")


def set_column_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(
    doc,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    *,
    font_size: float = 7.65,
    header_size: float | None = None,
    first_col_bold: bool = True,
    after: float = 5,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_width(table, int(sum(widths) * 1440))
    set_column_widths(table, widths)
    header = table.rows[0]
    repeat_header(header)
    for c_idx, text in enumerate(headers):
        cell = header.cells[c_idx]
        set_cell_shading(cell, PALE_GREEN)
        set_cell_border(cell, size="6")
        set_cell_margins(cell, top=70, start=90, bottom=70, end=90)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(text)
        set_run_font(r, size=header_size or font_size, bold=True, color=GREEN)
    for r_idx, row_values in enumerate(rows):
        row = table.add_row()
        for c_idx, value in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_shading(cell, WHITE if r_idx % 2 else PALE_GRAY)
            set_cell_border(cell, size="4")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(value)
            set_run_font(r, size=font_size, bold=(first_col_bold and c_idx == 0), color=DARK_GREEN if c_idx == 0 else TEXT)
    add_para(doc, "", size=1, after=after)
    return table


def add_callout(doc, label: str, body: str, *, fill: str = PALE_GREEN, after: float = 6):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_width(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, size="7")
    set_cell_margins(cell, top=85, start=120, bottom=85, end=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    set_run_font(r, size=8.8, bold=True, color=GREEN)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.02
    r2 = p2.add_run(body)
    set_run_font(r2, size=8.6, color=TEXT)
    add_para(doc, "", size=1, after=after)
    return table


def add_bullets(doc, items: list[str], *, size: float = 8.8):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.first_line_indent = Inches(-0.1)
        p.paragraph_format.space_after = Pt(1.8)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run("- " + item)
        set_run_font(r, size=size, color=TEXT)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1.08)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.73)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0.16)
    for style_name in ("Normal", "Body Text", "List Paragraph"):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(9.1)
        style.font.color.rgb = rgb(TEXT)
        style.paragraph_format.space_after = Pt(4.2)
        style.paragraph_format.line_spacing = 1.04


def normalize_letterhead_anchor(docx_path: Path) -> None:
    """Keep the full-page Navisync letterhead on-page in Word/PDF exports."""
    tmp_path = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as src, zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("word/header") and item.filename.endswith(".xml"):
                xml = data.decode("utf-8")
                xml = xml.replace(
                    '<wp:positionH relativeFrom="column"><wp:posOffset>-685800</wp:posOffset></wp:positionH>',
                    '<wp:positionH relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionH>',
                )
                xml = xml.replace(
                    '<wp:positionV relativeFrom="paragraph"><wp:posOffset>-731520</wp:posOffset></wp:positionV>',
                    '<wp:positionV relativeFrom="page"><wp:posOffset>0</wp:posOffset></wp:positionV>',
                )
                data = xml.encode("utf-8")
            dst.writestr(item, data)
    shutil.move(tmp_path, docx_path)


def add_title_page(doc: Document) -> None:
    add_para(doc, SUBTITLE, size=9.2, color=TEXT, before=0, after=10)
    p = add_para(doc, TITLE, size=21.5, color=GREEN, bold=True, before=0, after=12, line_spacing=1.0)
    p.paragraph_format.keep_with_next = True
    add_heading(doc, "Executive Summary", level=1, before=2, after=4)
    add_para(
        doc,
        "Beren is preparing Adrabetadex for an anticipated August 2026 PDUFA date in infantile-onset Niemann-Pick disease type C, based on client-supplied planning assumptions. The supplied clinical narrative includes a 71% reduction in mortality risk, and the supplied price assumption is $39,000 per 900 mg vial administered every two weeks. At 26 administrations per year, gross annual drug exposure is approximately $1,014,000 per patient, before administration, pharmacy handling, denial management, and financing costs. [1]",
    )
    add_para(
        doc,
        "The reimbursement issue is not whether the clinical profile matters. It does. The issue is whether hospitals and centers of excellence can administer Adrabetadex without becoming financially exposed.",
    )
    add_callout(
        doc,
        "Core point",
        "The Provider Net Cost Recovery Model should answer one operational question: after payment, acquisition cost, coding status, payer friction, denial risk, manufacturer support, and payment timing are applied, does a provider make or lose money treating this patient?",
    )


def add_provider_model(doc: Document) -> None:
    add_heading(doc, "1. Why the Provider Model Matters")
    add_para(
        doc,
        "Adrabetadex has the risk profile of a high-cost, provider-administered specialty therapy. Adoption will depend not only on payer coverage, but on whether the treating center can execute the full reimbursement pathway.",
    )
    add_table(
        doc,
        ["Launch issue", "What the provider sees", "Why it matters"],
        [
            ["Acquisition exposure", "Approximately $39,000 per vial before payment.", "Creates working-capital and inventory risk."],
            ["HCPCS uncertainty", "Early use may require NOC or temporary coding before a product-specific code.", "Manual billing increases documentation burden and denial risk. [2][4]"],
            ["Payment basis", "WAC, AWP, ASP, contractor pricing, or contracted rates may apply depending on site and timing.", "Small differences can determine whether treatment is profitable or loss-making. [3]"],
            ["Payer controls", "Prior authorization, site-of-care limits, and specialty-pharmacy sourcing.", "Clinical need does not eliminate administrative friction. [11][12]"],
            ["Medicaid / 340B complexity", "EPSDT may support medical necessity, but NDC, state payment, MCO, and duplicate-discount rules still matter.", "Coverage does not guarantee collectible payment. [7][8][9][10]"],
        ],
        [1.48, 3.05, 2.52],
    )
    add_callout(
        doc,
        "Positioning",
        "The model is not a price calculator. It is a reimbursement execution model that connects payer strategy, provider activation, field reimbursement, and finance.",
        fill=PALE_GRAY,
    )


def add_buy_bill(doc: Document) -> None:
    add_heading(doc, "2. Why Buy-and-Bill Is Different From EAP")
    add_para(
        doc,
        "Early Access Program experience can mislead launch planning. In an EAP, providers may not carry the same acquisition, billing, or payment risk they face after commercial launch.",
    )
    add_table(
        doc,
        ["Pathway", "Provider economics", "Launch implication"],
        [
            ["EAP", "Drug cost and reimbursement exposure may be limited or removed.", "Treatment experience does not prove commercial adoption readiness."],
            ["Commercial buy-and-bill", "Provider acquires the drug, bills payer, waits for reimbursement, and bears denial/payment risk.", "Viable only if payment, denial rate, and days to payment support a positive net position."],
            ["White-bagging / specialty-pharmacy sourcing", "Specialty pharmacy bills for the drug; provider may bill only administration.", "Drug margin is removed, but operational burden remains. [11]"],
            ["Site-of-care steering", "Payer may restrict where the drug may be administered.", "COE strategy must match clinical capability with payable site-of-care rules. [12]"],
        ],
        [1.55, 3.0, 2.5],
    )
    add_para(
        doc,
        "UnitedHealthcare's public sourcing materials describe a pathway where the specialty pharmacy bills the payer for the drug and outpatient providers may only seek reimbursement for administration, illustrating how commercial payer rules can remove provider drug economics from the treatment pathway. [11]",
    )


def add_model_calculation(doc: Document) -> None:
    add_heading(doc, "3. What the Model Should Calculate")
    add_para(doc, "The model should calculate provider economics by scenario, payer type, site of care, and channel.")
    add_table(
        doc,
        ["Model domain", "Inputs", "Output"],
        [
            ["Product and administration", "Dose, WAC, vial size, dosing frequency, procedure assumptions, pharmacy handling, labor, overhead.", "Per-dose and annual economics."],
            ["Coding and payment", "HCPCS status, NOC/C-code/J-code pathway, OPPS pass-through, WAC/AWP/ASP basis, allowed amount.", "Expected drug payment and claim risk."],
            ["Commercial medical benefit", "Prior authorization, white-bagging, site-of-care rules, denial rate, appeal recovery, days to payment.", "Buy-and-bill viability."],
            ["Medicaid", "EPSDT relevance, FFS vs MCO, state PAD methodology, NDC capture, utilization reporting.", "Whether medical necessity translates into payment."],
            ["340B", "Eligibility, estimated acquisition cost, carve-in/out posture, duplicate-discount controls.", "Whether the center has a structural acquisition advantage."],
            ["Manufacturer support", "Acquisition support, bridge logic, appeal support, prompt-pay assumptions.", "Break-even support threshold."],
        ],
        [1.5, 3.35, 2.2],
        font_size=7.45,
    )


def add_formula(doc: Document) -> None:
    add_heading(doc, "4. Core Formula")
    add_para(doc, "The model should use one operating equation across all scenarios:")
    add_callout(
        doc,
        "Expected provider net position",
        "Drug payment + administration revenue + provider-facing support - drug acquisition cost - administration and handling cost - denial/rework cost - payment-delay carrying cost.",
        fill=PALE_GRAY,
    )
    add_table(
        doc,
        ["Component", "Plain-English meaning"],
        [
            ["Drug payment", "What the payer allows for the drug claim."],
            ["Administration revenue", "Payment for the procedure and related services."],
            ["Acquisition cost", "What the provider pays for the drug, or avoids paying under white-bagging."],
            ["Denial / rework cost", "Expected administrative burden and revenue leakage."],
            ["Carrying cost", "Financing cost while waiting for reimbursement."],
            ["Break-even support", "Support needed to prevent provider loss."],
        ],
        [2.05, 5.0],
    )


def add_coding(doc: Document) -> None:
    add_heading(doc, "5. Coding and Pass-Through Timing")
    add_para(doc, "Coding and pass-through should be modeled as separate but related adoption drivers.")
    add_table(
        doc,
        ["Timing step", "Fact / planning point", "Model implication"],
        [
            ["PDUFA timing", "August 2026 is a client-supplied planning assumption. [1]", "Build launch scenarios around early coding uncertainty."],
            ["HCPCS Level II cycle", "CMS lists drug and biological product application deadlines as the first business day of each quarter: January, April, July, and October. [2]", "Test when a product-specific J-code could become available."],
            ["HCPCS limitation", "CMS states HCPCS coding does not itself determine coverage or payment. [2]", "Coding reduces friction, but does not guarantee reimbursement."],
            ["C9399 bridge", "CMS created C9399 for certain new FDA-approved outpatient drugs without product-specific HCPCS codes, with payment at 95% of AWP as determined by the contractor. [3][4]", "Model hospital outpatient bridge payment separately."],
            ["Initial sales period", "CMS guidance states payment may be up to 103% of WAC when ASP is not yet available. [3]", "Use as a Medicare benchmark, not a commercial guarantee."],
            ["OPPS pass-through", "Pass-through applications are handled through CMS MEARIS; pass-through status is temporary and subject to CMS approval. [5][6]", "Treat January 1, 2027 as a planning scenario, not a certainty."],
        ],
        [1.35, 3.65, 2.05],
        font_size=7.25,
    )


def add_340b(doc: Document) -> None:
    add_heading(doc, "6. Why 340B Matters")
    add_para(
        doc,
        "340B should be modeled as an acquisition advantage, not as a guaranteed margin. HRSA describes the 340B ceiling price as based on AMP minus URA, with operational adjustments for package and case-pack pricing. [7][8]",
    )
    add_table(
        doc,
        ["Scenario", "Provider position"],
        [
            ["Non-340B commercial COE", "Purchases at or near WAC, depends heavily on payment rate and denial timing."],
            ["340B COE", "May acquire at a materially lower cost, improving net recovery potential."],
            ["Medicaid 340B", "Requires carve-in/out and duplicate-discount controls."],
            ["White-bagging", "Provider does not buy the drug, but also loses drug margin opportunity."],
        ],
        [2.05, 5.0],
    )
    add_callout(
        doc,
        "Key implication",
        "Adrabetadex launch may concentrate in 340B-eligible hospital systems because they are more likely to withstand acquisition cost, timing, and payer friction.",
    )


def add_strategy(doc: Document) -> None:
    add_heading(doc, "7. Strategic Value for Beren")
    add_table(
        doc,
        ["Decision need", "What the model shows", "Leadership use"],
        [
            ["Site prioritization", "Which COEs can treat without expected loss.", "Focus launch resources."],
            ["Provider value story", "Whether drug cost, administration, and administrative burden are recoverable.", "Support provider education."],
            ["Payer friction", "Effect of white-bagging, prior authorization, and site-of-care rules.", "Shape contracting and field strategy."],
            ["Support strategy", "Support needed to reach break-even.", "Target support instead of over-subsidizing."],
            ["MLR clarity", "Separation of public facts, internal assumptions, and model outputs.", "Reduce review risk."],
        ],
        [1.65, 3.3, 2.1],
    )


def add_build(doc: Document) -> None:
    add_heading(doc, "8. Recommended Build")
    add_para(doc, "The first build should be a disciplined Excel model with a printable executive output tab. A web calculator can follow later.")
    add_table(
        doc,
        ["Tab", "Purpose"],
        [
            ["Instructions", "Define audience, limitations, and scenario use."],
            ["Core Assumptions", "WAC, dose, payment basis, denial rate, labor, handling, days to payment."],
            ["Scenario Selector", "Commercial buy-and-bill, white-bagging, Medicaid, 340B COE, pass-through."],
            ["Calculation Engine", "Locked formulas and source notes."],
            ["Executive Output", "Per-dose margin, annual exposure, cash burden, break-even result."],
            ["Source Appendix", "Public source notes and assumption status."],
        ],
        [1.85, 5.2],
    )


def add_mlr(doc: Document) -> None:
    add_heading(doc, "9. MLR Review Frame")
    add_para(doc, "The model should be positioned as reimbursement education and launch planning, not as a promise of:")
    add_bullets(doc, ["Coverage", "Payment", "Coding approval", "Pass-through approval", "Provider margin", "340B eligibility"])
    add_table(
        doc,
        ["Review area", "Recommended treatment"],
        [
            ["Clinical value", "Treat the 71% mortality-risk reduction as client-supplied unless public substantiation is added."],
            ["Coding", "Describe CMS process; do not imply J-code approval."],
            ["Pass-through", "Present January 1, 2027 as a planning scenario only."],
            ["Commercial controls", "Use payer policies as analogues, not predictions."],
            ["Medicaid", "Separate medical necessity from payment execution."],
            ["340B", "Model acquisition scenarios, not guaranteed margin."],
        ],
        [1.85, 5.2],
    )
    add_heading(doc, "Final Positioning", level=1, before=9, after=4)
    add_para(doc, "The Provider Net Cost Recovery Model is valuable because it makes provider adoption risk visible before launch.")
    add_para(
        doc,
        "It does not weaken the clinical case for Adrabetadex. It shows the commercial operating conditions required for hospitals and centers of excellence to administer the product without taking unacceptable financial risk.",
    )


def add_sources(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Source Notes", level=1, before=0, after=4)
    add_para(
        doc,
        "Source notes distinguish public reimbursement authorities, public payer-policy analogues, and client-supplied planning assumptions. Public sources were reviewed on April 28, 2026.",
        size=8.3,
    )
    rows = [[name, use, source] for name, use, source in SOURCES]
    add_table(
        doc,
        ["Reference", "Use in brief", "Source"],
        rows,
        [2.15, 2.25, 2.65],
        font_size=6.15,
        header_size=6.4,
        after=0,
    )


def build_docx() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    configure_document(doc)
    add_title_page(doc)
    add_provider_model(doc)
    add_buy_bill(doc)
    add_model_calculation(doc)
    add_formula(doc)
    add_coding(doc)
    add_340b(doc)
    add_strategy(doc)
    add_build(doc)
    add_mlr(doc)
    add_sources(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Revised executive brief for Beren Therapeutics provider net cost model planning"
    doc.core_properties.author = "Navisync reimbursement and health policy team"
    doc.core_properties.keywords = "Adrabetadex, provider net cost recovery, buy-and-bill, HCPCS, pass-through, Medicaid, 340B"
    doc.save(DOCX_PATH)
    normalize_letterhead_anchor(DOCX_PATH)
    return DOCX_PATH


def extract_plaintext(docx_path: Path) -> str:
    doc = Document(str(docx_path))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            vals = [cell.text.strip() for cell in row.cells]
            if any(vals):
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def inspect_package(docx_path: Path) -> dict:
    with zipfile.ZipFile(docx_path) as zf:
        names = zf.namelist()
        media = [name for name in names if name.startswith("word/media/")]
    text = extract_plaintext(docx_path)
    forbidden = {
        "em_dash": "\u2014" in text,
        "placeholder": any(token in text.lower() for token in ["lorem ipsum", "dear client", "full name", "address line", "\ufffc"]),
        "wrong_title": TITLE not in text,
    }
    return {
        "path": str(docx_path),
        "pdf_path": str(PDF_PATH),
        "paragraphs": len(Document(str(docx_path)).paragraphs),
        "tables": len(Document(str(docx_path)).tables),
        "media_count": len(media),
        "content_width_inches": CONTENT_WIDTH,
        "forbidden_checks": forbidden,
        "source_files": [str(TEMPLATE), str(SOURCE_DOCX), str(SOURCE_PDF)],
    }


def main() -> None:
    path = build_docx()
    report = inspect_package(path)
    MANIFEST_PATH.write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
