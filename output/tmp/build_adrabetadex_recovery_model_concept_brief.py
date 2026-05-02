from __future__ import annotations

import json
import shutil
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = Path("/Users/josephstewart/Downloads/Navisync Corporate Template.docx")
SOURCE_DOCX = Path("/Users/josephstewart/Downloads/Adrabetadex_Navisync_Final_Combined.docx")
SOURCE_PDF = Path("/Users/josephstewart/Downloads/Adrabetadex Provider Net Cost Calculator Concept Brief 4_1_26.pdf")
OUT_DIR = ROOT / "output" / "exports" / "adrabetadex_provider_net_cost_recovery_model_concept_brief_2026-04-28"
DOCX_PATH = OUT_DIR / "Adrabetadex_Provider_Net_Cost_Recovery_Model_Concept_Brief.docx"
MANIFEST_PATH = OUT_DIR / "build_manifest.json"

TITLE = "Adrabetadex Provider Net Cost Recovery Model"
SUBTITLE = "Concept brief for Beren Therapeutics, PBC Medical, Legal, and Regulatory review"
DATE_LINE = "April 2026"

GREEN = "4E8A2D"
DARK_GREEN = "1F3D23"
NAVY = "1F2E2B"
TEXT = "3F4A4A"
PALE_GREEN = "EEF5EA"
LIGHT_GREEN = "DDEAD7"
PALE_GRAY = "F6F8F5"
LINE = "BFD6B4"
AMBER = "8A5A12"
RED = "8C3D2B"
WHITE = "FFFFFF"


SOURCES = [
    {
        "id": 1,
        "name": "Beren Therapeutics / Navisync planning assumptions",
        "note": "Client-supplied planning facts: August 2026 PDUFA timing, infantile-onset Niemann-Pick disease population, 71% mortality-risk reduction statement, $39,000 WAC per 900 mg vial, and every-two-week dosing.",
        "url": "Internal planning assumption, supplied by client.",
    },
    {
        "id": 2,
        "name": "CMS, HCPCS Level II Coding Procedures",
        "note": "Drug and biological code applications are due the first business day of each quarter; HCPCS codes do not themselves determine coverage or payment.",
        "url": "https://www.cms.gov/medicare/coding-billing/healthcare-common-procedure-system/level-ii-coding-process",
    },
    {
        "id": 3,
        "name": "CMS, Part B Drug Payment Limits Overview",
        "note": "Initial sales period payment limit may use up to 103% of WAC when ASP is not yet available; OPPS drugs without an assigned HCPCS code may use 95% of AWP.",
        "url": "https://www.cms.gov/files/document/part-b-drug-payment-limits-overview.pdf-0",
    },
    {
        "id": 4,
        "name": "CMS, OPPS Pass-Through Payment Status page",
        "note": "OPPS drug and biological pass-through applications are submitted through MEARIS.",
        "url": "https://www.cms.gov/medicare/payment/prospective-payment-systems/hospital-outpatient/pass-through-payment-status-new-technology-ambulatory-payment-classification-apc",
    },
    {
        "id": 5,
        "name": "CMS, Transitional Pass-Through Drug and Biological Eligibility Process",
        "note": "Complete applications received by the first business day in September can support a January 1 earliest effective date; pass-through payments last at least two years and not more than three years.",
        "url": "https://www.cms.gov/files/document/determine-eligibility-drugs-biologicals-transitional-pass-through-payment-under-hospital-outpatient.pdf",
    },
    {
        "id": 6,
        "name": "CMS, C9399 unclassified outpatient drug payment",
        "note": "Hospitals may use C9399 for new FDA-approved drugs without a product-specific HCPCS code; payment is set at 95% of AWP as determined by the contractor.",
        "url": "https://www.cms.gov/newsroom/press-releases/medicare-pay-unclassified-fda-approved-drugs-administered-outpatient-departments",
    },
    {
        "id": 7,
        "name": "Medicaid.gov, EPSDT",
        "note": "EPSDT provides comprehensive services for Medicaid-enrolled children under age 21 and requires medically necessary services to correct or ameliorate health conditions.",
        "url": "https://www.medicaid.gov/medicaid/benefits/early-and-periodic-screening-diagnostic-and-treatment",
    },
    {
        "id": 8,
        "name": "Medicaid.gov, State Drug Utilization Data and FAQ",
        "note": "State utilization reporting under the Medicaid Drug Rebate Program is NDC-based; HCPCS-to-NDC mapping occurs before states report utilization to CMS.",
        "url": "https://www.medicaid.gov/faq/regarding-state-drug-utilization-data-sdud-there-official-crosswalk-of-national-drug-code-ndc-and-healthcare-common-procedure-coding-system-hcpcs-codes-which-providers-can-use/index.html",
    },
    {
        "id": 9,
        "name": "HRSA, 340B ceiling price calculation",
        "note": "The 340B ceiling price equals AMP minus URA, calculated at the smallest unit of measure and then published by HRSA.",
        "url": "https://www.hrsa.gov/about/faqs/how-340b-ceiling-price-calculated",
    },
    {
        "id": 10,
        "name": "HRSA, 340B pricing formulas",
        "note": "HRSA pricing guidance applies package-size and case-pack adjustments to the AMP minus URA calculation for operational pricing.",
        "url": "https://340bpricingsubmissions.hrsa.gov/Help/Manufacturer/Pricing%20Formulas/Pricing%20Formulas.htm",
    },
    {
        "id": 11,
        "name": "UnitedHealthcare, medication sourcing protocol",
        "note": "For sourced drugs, the specialty pharmacy bills UnitedHealthcare for the drug and outpatient providers may bill only for administration.",
        "url": "https://www.uhcprovider.com/content/dam/provider/docs/public/resources/pharmacy/medication-expansion-sourcing-faq.pdf",
    },
    {
        "id": 12,
        "name": "UnitedHealthcare, Specialty Pharmacy - Medical Benefit Management",
        "note": "UnitedHealthcare defines buy-and-bill as provider purchase, administration, and claim submission for the drug and related services.",
        "url": "https://www.uhcprovider.com/en/resource-library/drug-lists-pharmacy/specialty-medical-injectable-drug-programs.html",
    },
    {
        "id": 13,
        "name": "UnitedHealthcare, Provider Administered Drugs - Site of Care policy",
        "note": "UnitedHealthcare applies site-of-care criteria to specified provider-administered specialty medications.",
        "url": "https://www.uhcprovider.com/content/dam/provider/docs/public/policies/index/commercial/provider-administered-drugs-soc-04012026.pdf",
    },
    {
        "id": 14,
        "name": "UnitedHealthcare, Aqneursa prior authorization criteria",
        "note": "Comparable NPC policy requires genetic confirmation, specialist involvement, and positive clinical response for reauthorization.",
        "url": "https://www.uhcprovider.com/content/dam/provider/docs/public/prior-auth/drugs-pharmacy/commercial/a-g/PA-Med-Nec-Aqneursa.pdf",
    },
    {
        "id": 15,
        "name": "UnitedHealthcare, Miplyffa prior authorization criteria",
        "note": "Comparable NPC policy requires genetic confirmation, combination therapy criteria, specialist involvement, and positive clinical response for reauthorization.",
        "url": "https://www.uhcprovider.com/content/dam/provider/docs/public/prior-auth/drugs-pharmacy/commercial/h-p/PA-Med-Nec-Miplyffa.pdf",
    },
]


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = LINE, size: str = "6") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_twips: int) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_twips))
    tbl_w.set(qn("w:type"), "dxa")


def set_column_widths(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None, italic: bool = False) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = rgb(color)


def add_para(doc, text: str = "", style: str | None = None, size: float = 9.4, color: str = TEXT, bold: bool = False, before: float = 0, after: float = 4, line_spacing: float = 1.06):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
    return p


def add_heading(doc, text: str, level: int = 1, before: float = 12, after: float = 5):
    style = f"Heading {level}" if level in (1, 2, 3) else None
    p = add_para(doc, text, style=style, size=14 if level == 1 else 11.5 if level == 2 else 10, color=GREEN, bold=True, before=before, after=after)
    p.paragraph_format.keep_with_next = True
    return p


def add_small_label(doc, text: str):
    p = add_para(doc, text.upper(), size=7.5, color=GREEN, bold=True, before=6, after=1)
    p.paragraph_format.keep_with_next = True
    return p


def add_bullets(doc, items: list[str], size: float = 9.1):
    for item in items:
        p = doc.add_paragraph(style="List Paragraph")
        p.paragraph_format.left_indent = Inches(0.23)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.line_spacing = 1.04
        r = p.add_run(item)
        set_run_font(r, size=size, color=TEXT)


def add_callout(doc, label: str, body: str, fill: str = PALE_GREEN, label_color: str = GREEN):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(t, 9300)
    cell = t.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_border(cell, color=LINE, size="8")
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(label)
    set_run_font(r, size=9.2, bold=True, color=label_color)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    p2.paragraph_format.line_spacing = 1.05
    r2 = p2.add_run(body)
    set_run_font(r2, size=9, color=TEXT)
    add_para(doc, "", size=1, after=3)
    return t


def add_table(
    doc,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float],
    header_fill: str = PALE_GREEN,
    first_col_bold: bool = True,
    font_size: float = 8.4,
    header_margin: int = 110,
    body_margin: int = 100,
    trailing_after: float = 4,
):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 9300)
    set_column_widths(table, widths)
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, header_fill)
        set_cell_border(cell, color=LINE, size="7")
        set_cell_margins(cell, top=header_margin, start=115, bottom=header_margin, end=115)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=font_size, bold=True, color=GREEN)
    for r_idx, row_values in enumerate(rows):
        row = table.add_row()
        for c_idx, value in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_shading(cell, WHITE if r_idx % 2 else PALE_GRAY)
            set_cell_border(cell, color=LINE, size="5")
            set_cell_margins(cell, top=body_margin, start=115, bottom=body_margin, end=115)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            r = p.add_run(value)
            set_run_font(r, size=font_size - 0.2, bold=(first_col_bold and c_idx == 0), color=NAVY if c_idx == 0 else TEXT)
    add_para(doc, "", size=1, after=trailing_after)
    return table


def add_flow(doc, steps: list[tuple[str, str]], caption: str):
    headers = []
    row = []
    for i, (title, body) in enumerate(steps):
        headers.append(title)
        row.append(body)
        if i < len(steps) - 1:
            headers.append("")
            row.append(">")
    widths = []
    for i in range(len(headers)):
        widths.append(1.45 if headers[i] else 0.28)
    table = doc.add_table(rows=2, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_width(table, 9300)
    set_column_widths(table, widths)
    for c_idx, title in enumerate(headers):
        for r_idx in (0, 1):
            cell = table.cell(r_idx, c_idx)
            set_cell_margins(cell, top=80, start=70, bottom=80, end=70)
            if title:
                set_cell_shading(cell, PALE_GREEN if c_idx < 3 else PALE_GRAY)
                set_cell_border(cell, color=LINE, size="6")
            else:
                set_cell_shading(cell, WHITE)
                set_cell_border(cell, color=WHITE, size="0")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = table.cell(0, c_idx).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(title)
        set_run_font(r, size=8.3, bold=True, color=NAVY)
        p2 = table.cell(1, c_idx).paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(row[c_idx])
        set_run_font(r2, size=7.5, bold=not bool(title), color=GREEN if not title else TEXT)
    p = add_para(doc, caption, size=7.7, color=TEXT, after=5)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def clear_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1.05)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0.15)
    styles = doc.styles
    for style_name in ("Normal", "Body Text", "List Paragraph"):
        if style_name in styles:
            s = styles[style_name]
            s.font.name = "Arial"
            s._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
            s._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
            s.font.size = Pt(9.4)
            s.font.color.rgb = rgb(TEXT)
            s.paragraph_format.space_after = Pt(4)
            s.paragraph_format.line_spacing = 1.06
    for style_name, size in (("Heading 1", 14), ("Heading 2", 11.5), ("Heading 3", 10)):
        if style_name in styles:
            s = styles[style_name]
            s.font.name = "Arial"
            s._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
            s._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
            s.font.size = Pt(size)
            s.font.color.rgb = rgb(GREEN)
            s.font.bold = True
            s.paragraph_format.space_before = Pt(12)
            s.paragraph_format.space_after = Pt(5)
            s.paragraph_format.keep_with_next = True


def add_cover(doc: Document) -> None:
    p = add_para(doc, DATE_LINE, size=9.5, color=TEXT, after=6)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(8)
    title.paragraph_format.space_after = Pt(6)
    r = title.add_run(TITLE)
    set_run_font(r, size=23, bold=True, color=GREEN)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(10)
    r = sub.add_run(SUBTITLE)
    set_run_font(r, size=11, color=TEXT)
    add_callout(
        doc,
        "Management read",
        "Adrabetadex may have a compelling clinical profile, but provider adoption will depend on whether centers of excellence can buy, bill, administer, document, and collect payment without taking an unsustainable loss. The model converts that launch risk into a provider-level financial answer.",
        fill=PALE_GREEN,
    )
    add_small_label(doc, "Audience and use")
    add_table(
        doc,
        ["Audience", "What the brief answers", "MLR guardrail"],
        [
            [
                "Beren MLR, medical, access, and launch leadership",
                "What the provider net cost model is, why it matters, how it works, and what format it should take before a full build.",
                "Educational planning brief. It does not assert coverage, payment, coding approval, or payer acceptance.",
            ],
            [
                "Provider and center-of-excellence lens",
                "Whether the treatment pathway is financially workable at the site that purchases or administers Adrabetadex.",
                "All product-specific economics are labeled as supplied or planning assumptions.",
            ],
        ],
        [1.55, 3.25, 2.05],
        font_size=8,
    )


def add_summary(doc: Document) -> None:
    add_heading(doc, "Executive summary", 1)
    add_para(
        doc,
        "Beren is preparing Adrabetadex for an August 2026 PDUFA date in infantile-onset Niemann-Pick disease. The supplied clinical narrative includes a 71% reduction in mortality risk, and the supplied price assumption is $39,000 per 900 mg vial administered every two weeks. At that dosing cadence, gross annual drug exposure is approximately $1,014,000 before administration, pharmacy handling, denial work, and financing cost. [1]",
    )
    add_para(
        doc,
        "The reimbursement issue is not whether the clinical story matters. It does. The issue is that buy-and-bill adoption happens at the provider level. A hospital or center of excellence must purchase or control the drug, use the right code, satisfy payer documentation rules, submit a clean claim, wait for payment, and absorb any denial, delay, or rework. If that sequence fails, strong clinical value does not make the provider financially whole.",
    )
    add_callout(
        doc,
        "Core point",
        "The Provider Net Cost Recovery Model should answer one operational question: after all payment, cost, denial, support, and timing assumptions are applied, does a specific provider make or lose money treating a specific patient scenario?",
        fill=PALE_GREEN,
    )


def add_market_reality(doc: Document) -> None:
    add_heading(doc, "1. Why a provider model is needed", 1)
    add_para(
        doc,
        "Adrabetadex has the risk profile of a high-cost, provider-administered, specialty therapy. That means uptake will be shaped by the center that handles treatment, not only by payer coverage language. Intrathecal administration also means operational readiness matters: trained personnel, scheduling, documentation, pharmacy handling, procedure billing, and site capability all sit around the drug claim.",
    )
    add_table(
        doc,
        ["Launch issue", "What the provider sees", "Why it matters for adoption"],
        [
            ["Acquisition exposure", "$39,000 per vial before payer payment.", "Creates working-capital and inventory risk, especially if payment is delayed."],
            ["Coding pathway", "Early claims may use temporary or not-otherwise-classified logic before a product-specific J-code.", "Manual billing increases documentation burden and denial risk. [2][6]"],
            ["Payment basis", "Initial Medicare-linked benchmarks can use WAC, AWP, ASP, or contractor pricing depending on site and timing.", "Small basis differences can move margin from positive to negative. [3][6]"],
            ["Payer controls", "Prior authorization, site-of-care limits, and specialty-pharmacy sourcing can apply even for high-need rare diseases.", "Clinical need does not remove administrative friction. [11][13][14][15]"],
            ["Medicaid and 340B", "EPSDT can support medical necessity, while Medicaid drug data and 340B rules still drive payment integrity.", "Coverage and payment mechanics must both work. [7][8][9][10]"],
        ],
        [1.55, 2.8, 2.5],
        font_size=8,
    )
    add_para(
        doc,
        "The model should therefore be framed as a reimbursement execution model, not as a price model. It should sit between payer strategy, provider activation, field reimbursement, and finance.",
    )


def add_buy_bill_landscape(doc: Document) -> None:
    add_heading(doc, "2. Buy-and-bill in a vertically integrated market", 1)
    add_para(
        doc,
        "In a clean buy-and-bill pathway, the provider purchases the drug, administers it, and submits a claim for the drug and related services. UnitedHealthcare describes buy-and-bill in those terms on its medical-benefit specialty pharmacy page. [12] In practice, that pathway can be narrowed or displaced by specialty-pharmacy sourcing, commonly referred to as white-bagging.",
    )
    add_table(
        doc,
        ["Pathway", "Provider economics", "Commercial launch implication"],
        [
            ["Buy-and-bill", "Provider buys Adrabetadex, bills payer, and keeps or loses the spread after acquisition, administration, and claim costs.", "Viable only if reimbursement, denial rates, and days to payment support positive net position."],
            ["White-bagging", "Specialty pharmacy bills the drug. Provider generally bills only administration and cannot bill the member or payer for the drug under the sourcing protocol. [11]", "Drug margin is removed. Provider willingness depends on procedure economics, operational burden, and scheduling control."],
            ["Site-of-care steering", "Payer may require a non-hospital or lower-cost setting unless criteria support hospital outpatient care. [13]", "COE strategy must account for which sites can administer intrathecal therapy and still be payable."],
            ["Prior authorization", "Coverage criteria can require genetic confirmation, specialist involvement, combination criteria, and positive clinical response for continuation in comparable NPC products. [14][15]", "Strong clinical data may still need documentation, reauthorization, and appeal support."],
        ],
        [1.45, 2.9, 2.5],
        font_size=8,
    )
    add_flow(
        doc,
        [
            ("Payer policy", "Medical necessity, site, sourcing"),
            ("Provider pathway", "Buy-and-bill or white-bagging"),
            ("Claim execution", "Code, NDC, documentation"),
            ("Net result", "Margin and cash exposure"),
        ],
        "The model should show how payer rules change the provider's financial result, not only whether Adrabetadex is clinically appropriate.",
    )


def add_model_definition(doc: Document) -> None:
    add_heading(doc, "3. What the net cost recovery model is", 1)
    add_para(
        doc,
        "The Provider Net Cost Recovery Model is a scenario-based calculator that translates reimbursement mechanics into provider economics. It should calculate the expected net provider position for each major channel, payer type, and site of care. It should also show the support level or contracting condition needed to keep the provider whole.",
    )
    add_table(
        doc,
        ["Model domain", "Inputs the model should capture", "Output for Beren and providers"],
        [
            ["Product and administration", "Dose, vial size, WAC, administrations per year, intrathecal procedure assumptions, pharmacy handling, nursing, physician oversight, and facility overhead.", "Per-dose and annual treatment economics by site."],
            ["Coding and payment", "HCPCS status, NOC or C-code bridge, J-code timing, OPPS pass-through status, WAC/AWP/ASP basis, and payer-specific allowed amount.", "Expected drug payment and claim-cleanliness risk."],
            ["Commercial medical benefit", "Prior authorization, reauthorization, site-of-care rule, white-bagging requirement, contracted payment, denial rate, appeal recovery, and days to payment.", "Buy-and-bill viability and white-bagging impact."],
            ["Medicaid", "EPSDT relevance, fee-for-service versus managed care, state PAD methodology, NDC capture, and claim-level data requirements.", "Where medical necessity may be strong but payment execution remains fragile."],
            ["340B", "Eligibility, estimated acquisition cost, Medicaid carve-in or carve-out posture, duplicate-discount control, and COE status.", "Which centers have a structural acquisition advantage."],
            ["Manufacturer support", "Prompt-pay discounts, provider-facing acquisition support, free-drug bridge logic, field reimbursement support, and appeal support.", "Break-even support threshold without masking the base reimbursement problem."],
        ],
        [1.55, 3.1, 2.2],
        font_size=7.9,
    )


def add_methodology(doc: Document) -> None:
    add_heading(doc, "4. How the model should work", 1)
    add_para(
        doc,
        "The model should use one operating equation across all scenarios. Different channels change the inputs, not the logic.",
    )
    add_callout(
        doc,
        "Core formula",
        "Expected provider net position = drug payment + administration revenue + provider-facing manufacturer support - drug acquisition cost - administration and handling cost - denial and rework cost - carrying cost from payment delay.",
        fill=PALE_GRAY,
        label_color=GREEN,
    )
    add_table(
        doc,
        ["Formula component", "Plain-English definition", "Model treatment"],
        [
            ["Drug payment", "What the payer allows for the drug claim.", "Varies by WAC, AWP, ASP, contractor pricing, contracted rate, Medicaid methodology, or specialty-pharmacy sourcing."],
            ["Administration revenue", "Payment for the procedure and associated professional or facility services.", "Modeled separately from the drug so white-bagging does not look falsely neutral."],
            ["Acquisition cost", "What the provider pays or avoids paying for the drug.", "WAC for non-340B buy-and-bill; estimated 340B acquisition for eligible COEs; zero provider purchase in white-bagging."],
            ["Denial and rework cost", "Administrative burden and lost payment probability from prior authorization, missing documentation, wrong code, or failed NDC detail.", "Probability-weighted cost and revenue leakage."],
            ["Carrying cost", "Financing cost while the provider waits for reimbursement.", "Driven by acquisition cost, days to payment, and provider cost of capital."],
            ["Break-even support", "Support needed for the provider to avoid a negative result.", "Reported as a per-dose and annual threshold, not as a promotional guarantee."],
        ],
        [1.55, 2.75, 2.55],
        font_size=7.9,
    )


def add_channel_method(doc: Document) -> None:
    add_heading(doc, "5. Channel methodology and assumptions", 1)
    add_heading(doc, "Commercial medical benefit", 2, before=5)
    add_para(
        doc,
        "The commercial section should compare buy-and-bill against specialty-pharmacy sourcing. UnitedHealthcare public materials show both mechanisms: buy-and-bill remains possible when sourcing is not required, while the sourcing protocol can route the drug claim to a specialty pharmacy and leave the outpatient provider with only administration reimbursement. [11][12]",
    )
    add_heading(doc, "Medicaid", 2, before=5)
    add_para(
        doc,
        "For infantile-onset disease, Medicaid should be modeled as a coverage tailwind but not a payment guarantee. EPSDT requires medically necessary services for Medicaid-enrolled children under 21 when needed to correct or ameliorate health conditions. [7] The provider model still needs state payment methodology, managed-care variation, NDC capture, and utilization reporting logic because Medicaid drug rebate and utilization data are NDC-based. [8]",
    )
    add_heading(doc, "340B", 2, before=5)
    add_para(
        doc,
        "The 340B module should separate acquisition advantage from reimbursement. HRSA describes the ceiling price as AMP minus URA, with package and case-pack adjustments for operational pricing. [9][10] For Beren planning, the model should show primary and pediatric 340B acquisition assumptions separately and avoid implying that a 340B price is a provider margin guarantee.",
    )
    add_table(
        doc,
        ["Scenario block", "Default planning use", "Why it matters"],
        [
            ["Non-340B commercial COE", "WAC acquisition, commercial allowed amount, denial and days-to-payment assumptions.", "Tests whether ordinary buy-and-bill is economically viable."],
            ["340B COE", "Estimated 340B acquisition, same reimbursement benchmark where appropriate, duplicate-discount controls.", "Shows why launch may concentrate in eligible hospital systems."],
            ["Commercial white-bagging", "No provider drug purchase or drug claim; provider bills administration only.", "Shows when drug spread disappears and operational burden remains."],
            ["Medicaid non-340B", "State PAD payment, EPSDT medical necessity lens, NDC capture, Medicaid MCO variation.", "Tests whether coverage translates into collectible payment."],
            ["Hospital outpatient pass-through", "C9399 or product code bridge, pass-through status, OPPS payment treatment, acquisition-cost evidence.", "Shows how temporary payment visibility affects early hospital adoption."],
        ],
        [1.55, 2.8, 2.5],
        font_size=7.9,
    )


def add_coding_timing(doc: Document) -> None:
    add_heading(doc, "6. Coding and pass-through timing", 1)
    add_para(
        doc,
        "The model should distinguish three different payment states: launch bridge coding, product-specific coding, and temporary pass-through. CMS states that drug and biological HCPCS Level II applications are due the first business day of each quarter, and that the existence of a HCPCS code does not itself determine coverage or payment. [2]",
    )
    add_table(
        doc,
        ["Timing step", "Public rule or planning assumption", "Model implication"],
        [
            ["PDUFA timing", "Beren-supplied planning assumption: August 2026 PDUFA. [1]", "Build launch scenarios for pre-J-code and early commercial access."],
            ["HCPCS Level II drug cycle", "Applications for drugs and biologicals are due the first business day of January, April, July, and October; effective dates follow quarterly cycles. [2]", "Test whether an early product-specific J-code is available by the first full launch year."],
            ["C9399 hospital bridge", "CMS created C9399 for new FDA-approved hospital outpatient drugs without product-specific HCPCS codes; payment is 95% of AWP as contractor-priced. [6]", "Model hospital outpatient bridge payment separately from physician-office NOC billing."],
            ["Initial sales period", "CMS Part B guidance states payment can be up to 103% of WAC when ASP is not yet available. [3]", "Use as a federal benchmark, not as a commercial guarantee."],
            ["OPPS pass-through", "Complete September applications can support January 1 earliest status, and pass-through lasts at least two years but not more than three years. [5]", "If approval occurs in August 2026 and submission is complete, January 1, 2027 is a planning scenario, not a certainty."],
        ],
        [1.45, 3.1, 2.3],
        font_size=7.7,
    )
    add_flow(
        doc,
        [
            ("FDA approval", "August 2026 planning date"),
            ("Bridge claim", "NOC or C9399, manual support"),
            ("Code status", "J-code or C-code pathway"),
            ("Stable billing", "Lower manual friction"),
        ],
        "The model should show that coding reduces friction but does not remove payer policy, site-of-care, or provider cash exposure.",
    )


def add_value_and_format(doc: Document) -> None:
    add_heading(doc, "7. What value the model provides", 1)
    add_table(
        doc,
        ["Decision need", "What the model will show", "How leadership uses it"],
        [
            ["Site prioritization", "Which COEs can treat without expected loss under commercial, Medicaid, and 340B assumptions.", "Focus launch resources on financially workable sites."],
            ["Provider value story", "Why the provider can or cannot recover drug cost, procedure cost, and administrative burden.", "Support provider-facing reimbursement education."],
            ["Payer and PBM friction", "Effect of white-bagging, site-of-care rules, prior authorization, and reauthorization requirements.", "Shape contracting and field reimbursement plans before launch."],
            ["Support strategy", "How much provider-facing support is needed to reach break-even by scenario.", "Avoid broad support when targeted support would address the true barrier."],
            ["MLR clarity", "Which statements are public facts, which are internal assumptions, and which are model outputs.", "Review educational materials without confusing them for promotional claims."],
        ],
        [1.55, 3.05, 2.25],
        font_size=8,
    )
    add_heading(doc, "Recommended format", 2, before=7)
    add_para(
        doc,
        "The strongest first build is a disciplined Excel model with a clean executive output tab and transparent assumption tabs. A lightweight web calculator can follow later, but Excel is better for MLR review because formulas, labels, source notes, and assumption status can be audited directly.",
    )
    add_table(
        doc,
        ["Tab or view", "Purpose", "MLR control"],
        [
            ["Instructions", "Define use, audience, scenario selection, and limitations.", "States that output is a planning estimate, not reimbursement advice."],
            ["Core assumptions", "WAC, dose, administrations, payment bases, denial, appeal, labor, handling, and days-to-payment assumptions.", "Labels every input as supplied, public, modeling, or validation-needed."],
            ["Scenario selector", "Commercial buy-and-bill, white-bagging, Medicaid, 340B COE, and pass-through cases.", "Prevents uncontrolled scenario drift."],
            ["Calculation engine", "Applies the net position formula and break-even support logic.", "Locked formulas and visible source notes."],
            ["Executive output", "Per-dose margin, annual exposure, cash burden, and narrative readout.", "Provider-facing summary without exposing unnecessary formula detail."],
            ["Appendix assumptions", "Public source notes, payer-policy excerpts, and source-date tracking.", "Supports MLR substantiation."],
        ],
        [1.55, 3.05, 2.25],
        font_size=7.8,
    )


def add_mlr_considerations(doc: Document) -> None:
    add_heading(doc, "8. MLR review considerations", 1)
    add_para(
        doc,
        "The brief and the model should be positioned as reimbursement education and launch planning. They should not promise payer coverage, provider margin, code assignment, pass-through approval, or 340B eligibility. The safest review posture is to keep sourced facts, client-supplied facts, and model assumptions visibly separated.",
    )
    add_table(
        doc,
        ["Review area", "Recommended treatment", "Reason"],
        [
            ["Clinical value", "State the supplied 71% mortality-risk reduction as a Beren planning fact unless a public source is added.", "Avoids unsupported promotional framing in a reimbursement document."],
            ["Coding", "Describe CMS process and timing; do not imply J-code approval.", "CMS makes coding decisions independent of coverage and payment. [2]"],
            ["Pass-through", "Treat January 1, 2027 as a planning scenario only if the September application is complete and CMS timing permits.", "CMS timing is earliest possible, not guaranteed. [5]"],
            ["Commercial controls", "Use public payer policies as analogues, not predictions for Adrabetadex.", "Analogues show friction type, not a payer-specific final outcome."],
            ["Medicaid", "Separate medical necessity from payment execution.", "EPSDT can support coverage while NDC and state payment mechanics still affect provider recovery. [7][8]"],
            ["340B", "Use acquisition scenarios, not guaranteed margin language.", "340B changes acquisition cost but does not decide reimbursement, duplicate discount controls, or claim payment. [9][10]"],
        ],
        [1.45, 3.0, 2.4],
        font_size=7.9,
    )
    add_callout(
        doc,
        "Final review frame",
        "The model is valuable because it makes provider adoption risk explicit before launch. It does not diminish the clinical case for Adrabetadex. It shows the operating conditions needed for centers of excellence to administer the product without becoming financially exposed.",
        fill=PALE_GREEN,
    )


def add_sources(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Sources and assumption notes", 1)
    add_para(
        doc,
        "Source notes below distinguish public reimbursement authorities, public payer-policy analogues, and client-supplied planning assumptions. Public sources were reviewed on April 28, 2026.",
        size=8.8,
    )
    rows = [[f"[{s['id']}] {s['name']}", s["note"], s["url"]] for s in SOURCES]
    add_table(
        doc,
        ["Source", "Use in brief", "Reference"],
        rows,
        [1.75, 2.55, 2.55],
        font_size=6.25,
        first_col_bold=True,
        header_margin=45,
        body_margin=35,
        trailing_after=0,
    )


def build_docx() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(TEMPLATE))
    clear_body(doc)
    configure_document(doc)
    add_cover(doc)
    add_summary(doc)
    add_market_reality(doc)
    add_buy_bill_landscape(doc)
    add_model_definition(doc)
    add_methodology(doc)
    add_channel_method(doc)
    add_coding_timing(doc)
    add_value_and_format(doc)
    add_mlr_considerations(doc)
    add_sources(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Concept brief for Adrabetadex provider net cost recovery model"
    doc.core_properties.author = "Navisync reimbursement and health policy team"
    doc.core_properties.keywords = "Adrabetadex, provider net cost, buy-and-bill, HCPCS, pass-through, Medicaid, 340B"
    doc.save(DOCX_PATH)
    return DOCX_PATH


def extract_plaintext(docx_path: Path) -> str:
    doc = Document(str(docx_path))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            vals = [cell.text.strip() for cell in row.cells]
            if any(vals):
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def inspect_package(docx_path: Path) -> dict:
    with zipfile.ZipFile(docx_path) as zf:
        names = zf.namelist()
        media = [n for n in names if n.startswith("word/media/")]
    text = extract_plaintext(docx_path)
    forbidden = {
        "em_dash": "\u2014" in text,
        "placeholder": any(token in text.lower() for token in ["lorem ipsum", "dear client", "full name", "address line"]),
        "wrong_title": TITLE not in text,
    }
    return {
        "path": str(docx_path),
        "paragraphs": len(Document(str(docx_path)).paragraphs),
        "tables": len(Document(str(docx_path)).tables),
        "media_count": len(media),
        "forbidden_checks": forbidden,
        "source_files": [str(TEMPLATE), str(SOURCE_DOCX), str(SOURCE_PDF)],
    }


def main() -> None:
    path = build_docx()
    report = inspect_package(path)
    MANIFEST_PATH.write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
